"""Exercise a real standalone runtime in a disposable home, without provider calls.

Usage: python tools/smoke_local_agent.py --package dist/HarnessAgent
Omit --package to exercise source with the current Python interpreter.
Only generated sample documents, process logs and a redacted report are retained.
"""
from __future__ import annotations

import argparse
import contextlib
import hashlib
import http.cookiejar
import json
import os
from pathlib import Path
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request

ROOT = Path(__file__).resolve().parents[1]


class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, request, fp, code, message, headers, newurl):
        return None


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def request(opener, url, payload=None):
    body = json.dumps(payload).encode("utf-8") if payload is not None else None
    headers = {"Content-Type": "application/json"} if body is not None else {}
    call = urllib.request.Request(url, data=body, headers=headers)
    try:
        with opener.open(call, timeout=15) as response:
            return response.status, response.headers, response.read()
    except urllib.error.HTTPError as response:
        return response.code, response.headers, response.read()


def config_values(home):
    values = {}
    for line in (home / ".env").read_text(encoding="utf-8-sig").splitlines():
        if line.strip() and not line.lstrip().startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            values[key.strip()] = value.strip()
    return values


def status_value(home):
    try:
        return json.loads((home / "runtime.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def wait_ready(process, home, timeout=100):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        state = status_value(home)
        # Windows venv python.exe is a redirector: the serving interpreter has
        # a different PID. Each smoke home is exclusive and restart follows a
        # verified stopped status, so there is no stale ready owner to accept.
        if state.get("state") == "ready" and state.get("pid"):
            return state
        if process.poll() is not None:
            raise RuntimeError(f"server exited before ready (exit {process.returncode})")
        if state.get("state") == "error":
            raise RuntimeError("server reported startup error; inspect retained process log")
        time.sleep(.15)
    raise TimeoutError("server did not become ready")


def stop(process, *, eof=False):
    if process.poll() is not None:
        return process.returncode
    if process.stdin is not None and not process.stdin.closed:
        if not eof:
            process.stdin.write(b"stop\n")
            process.stdin.flush()
        process.stdin.close()
    try:
        return process.wait(timeout=45)
    except subprocess.TimeoutExpired:
        process.kill()
        process.wait(timeout=10)
        raise AssertionError("managed child did not stop gracefully within 45 seconds")


DEPENDENCY_PROBE = r'''
import asyncio,json,os,subprocess,sys
from pathlib import Path
app,home,out=map(Path,sys.argv[1:4])
sys.path.insert(0,str(app))
import local_agent
local_agent.configure_environment(home,local_agent.read_config(home/'.env'))
os.chdir(home)
from backend.config import DATA_DIR,WORKSPACE_DIR
assert DATA_DIR==home/'data' and WORKSPACE_DIR==home/'data/workspaces'
from backend.runtime import builtin_tools
workspace=WORKSPACE_DIR/'standalone-smoke'
workspace.mkdir(parents=True,exist_ok=True)
(workspace/'valid.py').write_text('answer = 42'+chr(10),encoding='utf-8')
(workspace/'invalid.py').write_text('if:'+chr(10),encoding='utf-8')
ctx=builtin_tools.BuiltinToolContext(root=workspace)
good=json.loads(asyncio.run(builtin_tools._lsp({'action':'diagnostics','path':'valid.py'},ctx)))
bad=json.loads(asyncio.run(builtin_tools._lsp({'action':'diagnostics','path':'invalid.py'},ctx)))
assert good['ok'] and not bad['ok'], 'real Python diagnostics failed'
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject,DictionaryObject,NameObject
from PIL import Image
from backend.runtime import presentation_requirements as source
pdf=out/'sample.pdf'
writer=PdfWriter()
font=writer._add_object(DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')}))
page=writer.add_blank_page(width=300,height=400)
page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
content=DecodedStreamObject()
content.set_data(b'BT /F1 12 Tf 20 340 Td (Standalone source page: clarify objectives and verify results.) Tj ET')
page[NameObject('/Contents')]=writer._add_object(content)
with pdf.open('wb') as stream: writer.write(stream)
render=out/'sample.png'
source._source_worker({'action':'render','source_path':str(pdf),'source_page':1,'output':str(render)},30)
with Image.open(render) as image:
    assert image.width>0 and max(image.size)<=source.MAX_RENDER_EDGE+1
    pixels=list(image.size)
built=source._source_worker({'action':'build','objective':'中文逐页重制ppt','documents':[str(pdf)]},30)
assert len(built['requirements']['pages'])==1, 'PDF text source worker failed'
try:
    source._source_worker({'action':'render','source_path':str(pdf),'source_page':1,'output':str(out/'timeout.png')},.0001)
except subprocess.TimeoutExpired:
    timeout=True
else:
    raise AssertionError('PDF child timeout did not fire')
from docx import Document
document=Document()
document.add_paragraph('Standalone document validation')
document.save(out/'sample.docx')
assert Document(out/'sample.docx').paragraphs[0].text=='Standalone document validation'
from openpyxl import Workbook,load_workbook
workbook=Workbook()
workbook.active['A1']='Standalone worksheet validation'
workbook.save(out/'sample.xlsx')
reopened=load_workbook(out/'sample.xlsx')
assert reopened.active['A1'].value=='Standalone worksheet validation'
reopened.close()
from backend.capabilities.presentations import create_presentation
presentation=create_presentation({'title':'Standalone validation','output_name':'sample','slides':[{'title':'原生文字测试','body':'已验证独立运行依赖。'}]},out)
assert presentation['slide_count']==1 and presentation['editable_text']
from backend.secret_store import encrypt_secret,decrypt_secret
assert decrypt_secret(encrypt_secret('disposable-smoke-value'))=='disposable-smoke-value'
print(json.dumps({'python_diagnostics':True,'pdf_render_pixels':pixels,'pdf_source_build':True,'pdf_timeout':timeout,'docx':True,'xlsx':True,'pptx':True,'credential_encryption':True,'python_executable':sys.executable},ensure_ascii=False))
'''


def run_smoke(args, report):
    package = args.package.resolve() if args.package else None
    app = package / "app" if package else ROOT
    python = package / "runtime/python.exe" if package else Path(sys.executable)
    require(python.is_file() and (app / "local_agent.py").is_file(), "package Python or app entry point is missing")
    report.update(package=str(package) if package else None, python=str(python), app=str(app))
    with tempfile.TemporaryDirectory(prefix="harness-standalone-smoke-") as temporary:
        root = Path(temporary)
        home = root / "独立 Agent 数据"
        cwd = root / "unrelated-working-directory"
        cwd.mkdir()
        alien = root / "ambient-must-remain-absent"
        environment = dict(os.environ)
        for name in ("PYTHONHOME", "PYTHONPATH", "VIRTUAL_ENV"):
            environment.pop(name, None)
        environment.update({
            "PATH": str(Path(os.environ.get("SystemRoot", "C:/Windows")) / "System32") if os.name == "nt" else "/usr/bin:/bin",
            "PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1",
            "APP_ENV_FILE": str(alien / "wrong.env"),
            "APP_DATA_DIR": str(alien / "data"),
            "AGENT_WORKSPACE_ROOT": str(alien / "workspaces"),
            "DATABASE_URL": "sqlite:///" + (alien / "wrong.db").as_posix(),
            "JWT_SECRET": "ambient-invalid-jwt", "SECRET_MASTER_KEY": "ambient-invalid-key",
            "ROOT_PASSWORD": "ambient-invalid-password", "APP_HOST": "0.0.0.0",
            "ALLOW_INSECURE_DEFAULTS": "true", "AUTH_COOKIE_SECURE": "true",
            "JOB_WORKER_ENABLED": "false", "CRON_SCHEDULER_ENABLED": "false",
            "CORS_ORIGINS": "*", "TRUST_PROXY_HEADERS": "true",
        })
        occupied = socket.socket()
        occupied.bind(("127.0.0.1", 0))
        occupied.listen(1)
        requested_port = occupied.getsockname()[1]
        require(requested_port < 65535, "OS selected final port; retry smoke")
        command = [str(python), "-B", "-X", "utf8", str(app / "local_agent.py"),
                   "serve", "--home", str(home), "--port", str(requested_port), "--managed"]
        flags = {"creationflags": subprocess.CREATE_NO_WINDOW} if os.name == "nt" else {}
        process = None
        stream = (args.output / "server-first.log").open("wb")
        try:
            process = subprocess.Popen(command, cwd=cwd, env=environment, stdin=subprocess.PIPE,
                                       stdout=stream, stderr=subprocess.STDOUT, **flags)
            state = wait_ready(process, home)
            url = state["url"]
            actual_port = int(url.rstrip("/").rsplit(":", 1)[1])
            require(actual_port != requested_port, "occupied port was not bypassed")
            report["checks"]["occupied_port_fallback"] = True
            anonymous = urllib.request.build_opener(urllib.request.ProxyHandler({}), NoRedirect())
            code, headers, body = request(anonymous, url)
            require(code == 303 and headers.get("Location") == "/login", "anonymous root did not redirect to login")
            code, _, body = request(anonymous, url + "healthz")
            health = json.loads(body)
            require(code == 200 and health["migration"]["ok"] and health["workers"]["ok"] and health["scheduler"]["ok"], "health not ready")
            require(health["deployment"]["security_ok"], "local security configuration is invalid")
            report["health"] = health
            report["checks"]["health_migration_workers_scheduler"] = True
            for resource, marker in [("login", b"<html"), ("static/app.js", b""), ("branding/favicon.svg", b"<svg")]:
                code, _, body = request(anonymous, url + resource)
                require(code == 200 and len(body) > 20 and marker in body.lower(), f"missing packaged resource: {resource}")
            report["checks"]["frontend_and_branding"] = True
            values = config_values(home)
            require(values["JWT_SECRET"] != "ambient-invalid-jwt" and values["SECRET_MASTER_KEY"] != "ambient-invalid-key", "ambient secrets were inherited")
            saved_config = hashlib.sha256((home / ".env").read_bytes()).hexdigest()
            jar = http.cookiejar.CookieJar()
            authenticated = urllib.request.build_opener(urllib.request.ProxyHandler({}), urllib.request.HTTPCookieProcessor(jar), NoRedirect())
            login = {"username": values["ROOT_USERNAME"], "password": values["ROOT_PASSWORD"]}
            code, _, body = request(authenticated, url + "api/v1/auth/login", login)
            require(code == 200 and json.loads(body)["role"] == "root", "initial login failed")
            code, _, body = request(authenticated, url)
            require(code == 200 and b"<html" in body.lower(), "authenticated app did not open")
            report["checks"]["initial_login"] = True
            require(not alien.exists(), "ambient storage path was created")
            report["checks"]["isolated_home_and_hostile_environment"] = True
            before = (home / "runtime.json").read_bytes()
            duplicate = subprocess.run(command, cwd=cwd, env=environment, input=b"", capture_output=True, timeout=20, **flags)
            (args.output / "duplicate.log").write_bytes(duplicate.stdout + duplicate.stderr)
            require(duplicate.returncode != 0, "duplicate launch unexpectedly succeeded")
            require((home / "runtime.json").read_bytes() == before, "duplicate launch overwrote live status")
            require(process.poll() is None, "duplicate launch stopped original process")
            report["checks"]["duplicate_launch_preserves_owner"] = True
            artifacts = args.output / "artifacts"
            artifacts.mkdir()
            dependency = subprocess.run([str(python), "-B", "-X", "utf8", "-c", DEPENDENCY_PROBE,
                                         str(app), str(home), str(artifacts)], cwd=cwd, env=environment,
                                        capture_output=True, timeout=90, **flags)
            (args.output / "dependencies.log").write_bytes(dependency.stdout + dependency.stderr)
            require(dependency.returncode == 0, "runtime dependency probe failed; inspect dependencies.log")
            report["dependencies"] = json.loads(dependency.stdout.decode("utf-8").strip().splitlines()[-1])
            report["checks"]["real_python_pdf_and_document_dependencies"] = True
            require(stop(process) == 0, "stdin stop did not exit successfully")
            require(status_value(home).get("state") == "stopped", "stop status was not persisted")
            report["checks"]["stdin_graceful_stop"] = True
        finally:
            occupied.close()
            if process is not None and process.poll() is None:
                with contextlib.suppress(Exception):
                    stop(process)
            stream.close()
        with contextlib.closing(sqlite3.connect(home / "data/app.db")) as database:
            revision = database.execute("SELECT version_num FROM alembic_version").fetchone()[0]
            password_hash = database.execute("SELECT password_hash FROM users WHERE username=?", (values["ROOT_USERNAME"],)).fetchone()[0]
            require(database.execute("PRAGMA quick_check").fetchone()[0] == "ok", "database integrity check failed")
            require(database.execute("SELECT count(*) FROM agents").fetchone()[0] >= 1, "default Agent is missing")
        report["schema_revision"] = revision
        command[command.index("--port") + 1] = "0"
        process = None
        with (args.output / "server-restart.log").open("wb") as stream:
            try:
                process = subprocess.Popen(command, cwd=cwd, env=environment, stdin=subprocess.PIPE,
                                           stdout=stream, stderr=subprocess.STDOUT, **flags)
                state = wait_ready(process, home)
                require(hashlib.sha256((home / ".env").read_bytes()).hexdigest() == saved_config, "restart changed persisted secrets")
                code, _, body = request(authenticated, state["url"] + "api/v1/auth/login", login)
                require(code == 200 and json.loads(body)["role"] == "root", "login credentials did not survive restart")
                with contextlib.closing(sqlite3.connect(home / "data/app.db")) as database:
                    require(database.execute("SELECT password_hash FROM users WHERE username=?", (values["ROOT_USERNAME"],)).fetchone()[0] == password_hash, "restart replaced account password")
                report["checks"]["restart_preserves_keys_accounts_and_login"] = True
                require(stop(process, eof=True) == 0, "managed EOF did not stop child successfully")
                require(status_value(home).get("state") == "stopped", "managed EOF did not persist stopped status")
                report["checks"]["managed_eof_graceful_stop"] = True
                require(not alien.exists(), "ambient storage path was touched on restart")
            finally:
                if process is not None and process.poll() is None:
                    with contextlib.suppress(Exception):
                        stop(process)
        # No credentials or disposable home are copied into evidence.
        for name in ("server-first.log", "server-restart.log", "duplicate.log", "dependencies.log"):
            content = (args.output / name).read_text(encoding="utf-8", errors="replace")
            for key in ("JWT_SECRET", "SECRET_MASTER_KEY", "ROOT_PASSWORD"):
                require(values[key] not in content, f"credential found in {name}")
        report["checks"]["retained_logs_exclude_generated_secrets"] = True


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--package", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)
    run_name = time.strftime("%Y%m%d-%H%M%S") + ("-package-" if args.package else "-source-") + str(os.getpid())
    args.output = (args.output or ROOT / "output/standalone-smoke" / run_name).resolve()
    args.output.mkdir(parents=True, exist_ok=False)
    report = {"ok": False, "checks": {}, "provider_calls": 0,
              "limitations": ["Native launcher window interaction and external system browser are separate checks."]}
    started = time.monotonic()
    try:
        run_smoke(args, report)
        report["ok"] = True
    except Exception as error:
        report["error"] = {"class": type(error).__name__, "message": str(error)}
    finally:
        report["duration_seconds"] = round(time.monotonic() - started, 2)
        report_path = args.output / "report.json"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"ok": report["ok"], "report": str(report_path)}, ensure_ascii=False))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
