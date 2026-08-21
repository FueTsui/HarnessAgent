"""内置 Harness 能力、安全边界与持久调度回归测试。"""
import asyncio
import base64
import datetime
import io
import json
import tempfile
import unittest
import uuid
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from PIL import Image

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from backend import jobs, scheduler
from backend.database import Base
from backend.models import Agent, Job, ScheduledTask, User
from backend.runtime import task_store
from backend.runtime import builtin_tools
from backend.runtime import run_harness


class _BuiltinCallingLlm:
    context_tokens = 8192

    def __init__(self, path):
        self.path = path
        self.calls = 0
        self.messages = []

    async def chat_with_tools(self, messages, tools=None, temperature=1):
        self.calls += 1
        self.messages = list(messages)
        if self.calls == 1:
            names = {row["function"]["name"] for row in tools or []}
            if "read" not in names:
                raise AssertionError("内置 read 未进入 Harness 工具目录")
            return {
                "role": "assistant", "content": "", "tool_calls": [{
                    "id": "read-1", "type": "function",
                    "function": {
                        "name": "read",
                        "arguments": json.dumps({"path": self.path}),
                    },
                }],
            }
        return {"role": "assistant", "content": "已读取并验证", "tool_calls": []}


class _ImageFailureLlm:
    context_tokens = 8192

    def __init__(self):
        self.calls = 0

    async def chat_with_tools(self, messages, tools=None, temperature=1):
        self.calls += 1
        if self.calls == 1:
            return {
                "role": "assistant",
                "content": "",
                "tool_calls": [{
                    "id": "image-1",
                    "type": "function",
                    "function": {
                        "name": "image_generate",
                        "arguments": json.dumps({
                            "prompt": "测试图片",
                            "size": "1024x1024",
                            "confirm": True,
                        }, ensure_ascii=False),
                    },
                }],
            }
        return {"role": "assistant", "content": "图片能力未配置。", "tool_calls": []}


class _DocumentCallingLlm:
    context_tokens = 8192

    def __init__(self):
        self.calls = 0

    async def chat_with_tools(self, messages, tools=None, temperature=1):
        self.calls += 1
        if self.calls == 1:
            names = {row["function"]["name"] for row in tools or []}
            if "document_create" not in names:
                raise AssertionError("document_create 未进入 Harness 工具目录")
            return {
                "role": "assistant",
                "content": "",
                "tool_calls": [{
                    "id": "doc-1",
                    "type": "function",
                    "function": {
                        "name": "document_create",
                        "arguments": json.dumps({
                            "title": "图片识别结果",
                            "tables": [{
                                "headers": ["序号", "原形", "过去式"],
                                "rows": [["1", "be", "was/were"]],
                            }],
                            "output_name": "图片识别结果.docx",
                            "confirm": True,
                        }, ensure_ascii=False),
                    },
                }],
            }
        return {"role": "assistant", "content": "已生成可下载的 Word 文档。", "tool_calls": []}


def _png_bytes() -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (8, 6), (20, 40, 60)).save(output, format="PNG")
    return output.getvalue()


def _editorial_panel_bytes() -> bytes:
    ivory = (243, 240, 232)
    image = Image.new("RGB", (600, 400), ivory)
    mask = Image.new("L", image.size, 0)
    from PIL import ImageDraw
    ImageDraw.Draw(mask).ellipse((120, 45, 500, 330), fill=255)
    gradient = Image.new("RGB", image.size)
    pixels = gradient.load()
    for y in range(image.height):
        ratio = y / max(1, image.height - 1)
        color = (
            round(190 - 70 * ratio),
            round(215 - 70 * ratio),
            round(230 - 65 * ratio),
        )
        for x in range(image.width):
            pixels[x, y] = color
    image.paste(gradient, (0, 0), mask)
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _image_pixels(image: Image.Image):
    flattened = getattr(image, "get_flattened_data", None)
    return flattened() if callable(flattened) else image.getdata()


class BuiltinCapabilityTests(unittest.TestCase):
    def setUp(self):
        # Every test owns a real, already-created workspace. Do not depend on the
        # repository's ignored tmp/ directory existing on a fresh CI checkout.
        self._workspace = tempfile.TemporaryDirectory()
        self.root = Path(self._workspace.name)
        self.prefix = "case_" + uuid.uuid4().hex
        self.ctx = builtin_tools.BuiltinToolContext(root=self.root, user_id=1, agent_id=1)

    def tearDown(self):
        self._workspace.cleanup()

    def run_tool(self, name, args):
        return json.loads(asyncio.run(builtin_tools.execute(name, args, self.ctx)))

    def test_catalog_contains_complete_capability_groups(self):
        names = {row["name"] for row in builtin_tools.capability_catalog()}
        expected = {
            "ls", "glob", "grep", "read", "read_many", "write", "edit",
            "multi_edit", "apply_patch", "shell", "git_status", "git_diff",
            "git_commit", "lsp", "html_generate", "image_generate", "image_render",
            "browser_open", "browser_snapshot", "browser_click", "browser_type",
            "CronCreate", "CronDelete", "CronList", "spawn_agent", "resume_agent",
            "wait_agent", "list_agents", "close_agent", "interrupt_agent",
            "web_search", "web_fetch",
            "document_inspect", "document_create", "document_format",
        }
        self.assertFalse(expected - names)

    def test_free_search_parsers_and_rank_fusion(self):
        ddg = builtin_tools._parse_duckduckgo(
            '<a class="result__a" href="https://example.com/a?utm_source=x">A</a>'
            '<a class="result__snippet">Alpha result</a>'
        )
        bing = builtin_tools._parse_bing(
            '<?xml version="1.0"?><rss><channel><item><title>A mirror</title>'
            '<link>https://example.com/a</link><description>Beta result</description>'
            '</item><item><title>B</title><link>https://example.org/b</link>'
            '<description>Second</description></item></channel></rss>'
        )
        merged = builtin_tools._merge_search_results([ddg, bing], 5)
        self.assertEqual(merged[0]["url"], "https://example.com/a")
        self.assertEqual(merged[0]["source"], "duckduckgo+bing")
        self.assertEqual(len(merged), 2)

    def test_news_queries_merge_news_rss_with_direct_fallback_results(self):
        news_row = {
            "title": "科技新闻",
            "url": "https://news.google.com/rss/articles/wrapper",
            "snippet": "摘要",
            "source": "google_news",
        }
        direct_row = {
            "title": "科技新闻原文",
            "url": "https://publisher.example/article",
            "snippet": "正文摘要",
            "source": "bing",
        }
        builtin_tools._SEARCH_CACHE.clear()
        with patch.object(
            builtin_tools,
            "_google_news_search",
            new=AsyncMock(return_value=[news_row]),
        ) as news, patch.object(
            builtin_tools,
            "_fallback_search",
            new=AsyncMock(return_value=[direct_row]),
        ) as fallback:
            result = asyncio.run(builtin_tools._web_search(
                {"query": "最新科技新闻", "count": 20, "language": "zh-CN"},
                builtin_tools.BuiltinToolContext(),
            ))
        self.assertEqual(result["provider"], "google_news+bing")
        self.assertEqual(result["results"][0]["url"], direct_row["url"])
        self.assertIn(news_row["url"], [row["url"] for row in result["results"]])
        news.assert_awaited_once()
        fallback.assert_awaited_once()

    def test_web_fetch_rejects_empty_html_instead_of_claiming_evidence(self):
        response = SimpleNamespace(
            headers={"content-type": "text/html; charset=utf-8"},
            url="https://news.google.com/rss/articles/wrapper",
        )
        with patch.object(
            builtin_tools,
            "_safe_web_get",
            new=AsyncMock(return_value=(response, "<html><title>Google News</title></html>")),
        ):
            with self.assertRaisesRegex(builtin_tools.ToolError, "未提取到可用正文"):
                asyncio.run(builtin_tools._web_fetch(
                    {"url": str(response.url)}, builtin_tools.BuiltinToolContext()
                ))

    def test_web_fetch_formats_json_and_reports_real_truncation(self):
        response = SimpleNamespace(
            headers={"content-type": "application/json; charset=utf-8"},
            url="https://api.example.test/data",
        )
        with patch.object(
            builtin_tools,
            "_safe_web_get",
            new=AsyncMock(return_value=(response, '{"message":"联网正常","items":[1,2]}')),
        ):
            result = asyncio.run(builtin_tools._web_fetch(
                {"url": str(response.url)}, builtin_tools.BuiltinToolContext()
            ))
        self.assertIn("联网正常", result["text"])
        self.assertFalse(result["truncated"])

    def test_root_global_switch_and_agent_binding_are_both_required(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        db = sessionmaker(bind=engine)()
        try:
            agent = Agent(
                name="能力授权测试",
                builtin_tools=json.dumps(["read", "write"], ensure_ascii=False),
            )
            db.add(agent)
            db.commit()
            self.assertEqual(
                builtin_tools.effective_tool_names(db, agent),
                {"read", "write"},
            )
            builtin_tools.set_global_enabled(db, "write", False)
            db.commit()
            self.assertEqual(
                builtin_tools.effective_tool_names(db, agent),
                {"read"},
            )
            states = {
                row["name"]: row["enabled"]
                for row in builtin_tools.capability_catalog(db, agent)
            }
            self.assertTrue(states["read"])
            self.assertFalse(states["write"])
        finally:
            db.close()
            engine.dispose()

    def test_execute_rejects_tools_outside_runtime_authorization(self):
        context = builtin_tools.BuiltinToolContext(
            root=self.root,
            enabled_tools={"read"},
        )
        result = json.loads(asyncio.run(
            builtin_tools.execute(
                "write",
                {
                    "path": self.prefix + "_denied.txt",
                    "content": "forbidden",
                    "confirm": True,
                },
                context,
            )
        ))
        self.assertFalse(result["ok"])
        self.assertIn("未授权", result["error"])
        self.assertFalse((self.root / (self.prefix + "_denied.txt")).exists())

    def test_filesystem_tools_are_jailed_and_edits_are_exact(self):
        written = self.run_tool("write", {
            "path": self.prefix + "_a.txt", "content": "alpha\nbeta\n", "confirm": True,
        })
        self.assertTrue(written["ok"])
        read = self.run_tool("read", {"path": self.prefix + "_a.txt"})
        self.assertIn("1: alpha", read["content"])
        edited = self.run_tool("edit", {
            "path": self.prefix + "_a.txt", "old": "beta", "new": "gamma", "confirm": True,
        })
        self.assertTrue(edited["ok"])
        grep = self.run_tool("grep", {"path": ".", "query": "gamma"})
        self.assertEqual(grep["matches"][0]["path"], self.prefix + "_a.txt")
        escaped = self.run_tool("read", {"path": "../outside.txt"})
        self.assertFalse(escaped["ok"])
        self.assertIn("路径越界", escaped["error"])

    def test_multi_edit_is_atomic_on_validation_failure(self):
        path = self.root / (self.prefix + "_a.txt")
        path.write_text("one two", encoding="utf-8")
        result = self.run_tool("multi_edit", {
            "confirm": True,
            "edits": [
                {"path": path.name, "old": "one", "new": "1"},
                {"path": path.name, "old": "missing", "new": "x"},
            ],
        })
        self.assertFalse(result["ok"])
        self.assertEqual(path.read_text(encoding="utf-8"), "one two")

    def test_structured_patch_adds_and_updates_atomically(self):
        existing = self.root / (self.prefix + "_existing.txt")
        added = self.root / (self.prefix + "_added.txt")
        existing.write_text("before", encoding="utf-8")
        result = self.run_tool("apply_patch", {
            "confirm": True,
            "files": [
                {
                    "path": existing.name,
                    "action": "update",
                    "operations": [{"old": "before", "new": "after"}],
                },
                {"path": added.name, "action": "add", "content": "created"},
            ],
        })
        self.assertTrue(result["ok"])
        self.assertEqual(existing.read_text(encoding="utf-8"), "after")
        self.assertEqual(added.read_text(encoding="utf-8"), "created")

    def test_shell_rejects_chaining_and_inline_code(self):
        chained = self.run_tool("shell", {
            "shell": "powershell", "script": "Get-ChildItem; Get-Content .env",
            "confirm": True,
        })
        self.assertFalse(chained["ok"])
        inline = self.run_tool("shell", {
            "shell": "powershell", "script": "python -c print(1)", "confirm": True,
        })
        self.assertFalse(inline["ok"])

    def test_lsp_python_diagnostics(self):
        valid = self.root / (self.prefix + "_valid.py")
        invalid = self.root / (self.prefix + "_invalid.py")
        valid.write_text("value = 1\n", encoding="utf-8")
        result = self.run_tool("lsp", {
            "action": "diagnostics", "path": valid.name,
        })
        self.assertTrue(result["ok"])
        invalid.write_text("if:\n", encoding="utf-8")
        result = self.run_tool("lsp", {
            "action": "diagnostics", "path": invalid.name,
        })
        self.assertFalse(result["ok"])

    def test_html_and_svg_artifacts_are_generated_safely(self):
        with patch.object(builtin_tools, "EXPORT_DIR", self.root):
            html_result = self.run_tool("html_generate", {
                "title": self.prefix,
                "html": "<main>report</main>",
                "confirm": True,
            })
            self.assertTrue(html_result["ok"])
            self.assertTrue((self.root / html_result["file"]).is_file())
            svg_result = self.run_tool("image_render", {
                "title": self.prefix,
                "svg": '<svg xmlns="http://www.w3.org/2000/svg"><text>ok</text></svg>',
                "confirm": True,
            })
            self.assertTrue(svg_result["ok"])
            unsafe = self.run_tool("image_render", {
                "title": self.prefix,
                "svg": '<svg><script>alert(1)</script></svg>',
                "confirm": True,
            })
            self.assertFalse(unsafe["ok"])

    def test_image_preflight_reports_capability_before_approval(self):
        context = builtin_tools.BuiltinToolContext(
            root=self.root,
            run_id="approval-bound-run",
        )
        with patch.object(builtin_tools.settings, "IMAGE_API_BASE_URL", ""), \
             patch.object(builtin_tools.settings, "IMAGE_API_KEY", ""), \
             patch.object(builtin_tools.settings, "IMAGE_MODEL", ""):
            result = json.loads(asyncio.run(builtin_tools.execute(
                "image_generate",
                {"prompt": "test", "size": "1024x1024", "confirm": False},
                context,
            )))
        self.assertFalse(result["ok"])
        self.assertEqual(result["code"], "image_provider_not_configured")
        self.assertIn("图片服务未配置", result["error"])
        issues = builtin_tools.completion_artifact_issues(context)
        self.assertEqual(len(issues), 1)
        self.assertIn("image_provider_not_configured", issues[0])
        self.assertIn("图片服务未配置", issues[0])

    def test_photo_editorial_skill_uses_model_layout_local_renderer(self):
        class EditorialVisionLlm:
            wire_api = "responses"
            model_id = "gpt-5.6-sol"

            async def vision(self, system, user, image_paths, temperature=0.2):
                self.system = system
                self.user = user
                self.image_paths = image_paths
                return json.dumps({
                    "title": "Shells Against Blue",
                    "subtitle": "Architecture held beneath open sky",
                    "motif": "arches",
                    "focus_x": 0.66,
                    "density": 4,
                })

        llm = EditorialVisionLlm()
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "source.png"
            source.write_bytes(_png_bytes())
            with patch("backend.artifacts.EXPORT_DIR", Path(temp)), \
                 patch.object(builtin_tools.settings, "IMAGE_API_BASE_URL", ""), \
                 patch.object(builtin_tools.settings, "IMAGE_API_KEY", ""), \
                 patch.object(builtin_tools.settings, "IMAGE_MODEL", ""):
                context = builtin_tools.BuiltinToolContext(
                    root=self.root,
                    llm=llm,
                    attachment_images=[source],
                    active_skill_names={"photo-abstract-editorial"},
                )
                result = json.loads(asyncio.run(builtin_tools.execute(
                    "image_generate",
                    {
                        "prompt": "Create a faithful abstract editorial diptych",
                        "size": "1024x1024",
                        "action": "edit",
                        "confirm": True,
                    },
                    context,
                )))
                completion_issues = builtin_tools.completion_artifact_issues(context)

            self.assertTrue(result["ok"])
            self.assertEqual(result["source"], "provider_vision_local_editorial")
            self.assertEqual(result["model"], "gpt-5.6-sol")
            self.assertEqual(result["action"], "compose")
            self.assertEqual(result["design"]["title"], "Shells Against Blue")
            self.assertEqual(result["renderer_version"], "photo_editorial_v3")
            self.assertEqual(result["design_source"], "vision_model")
            self.assertTrue(all(result["quality_checks"].values()))
            self.assertEqual(llm.image_paths, [source])
            target = Path(temp) / result["file"]
            self.assertTrue(target.is_file())
            with Image.open(target) as rendered:
                # 实体拱形质量应在母题区占据可观面积，避免退化成几条信息图轮廓线。
                photo_height = round(6 * rendered.width / 8)
                panel_height = rendered.height - photo_height
                motif = rendered.crop((
                    0,
                    photo_height + round(panel_height * .14),
                    rendered.width,
                    photo_height + round(panel_height * .58),
                ))
                ivory = (243, 240, 232)
                non_background = sum(
                    pixel != ivory
                    for pixel in _image_pixels(motif.convert("RGB"))
                )
                self.assertGreater(non_background, motif.width * motif.height * .04)
                self.assertGreaterEqual(result["design"]["motif_width"], .60)
                self.assertTrue(result["quality_checks"]["motif_scale"])
                self.assertTrue(result["quality_checks"]["tonal_depth"])
            self.assertEqual(completion_issues, [])

    def test_photo_editorial_rejects_empty_vision_design(self):
        class EmptyVisionLlm:
            model_id = "multimodal-model"

            async def vision(self, *_args, **_kwargs):
                return ""

        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "source.png"
            source.write_bytes(_png_bytes())
            with patch("backend.artifacts.EXPORT_DIR", Path(temp)), \
                 patch.object(builtin_tools.settings, "IMAGE_API_BASE_URL", ""), \
                 patch.object(builtin_tools.settings, "IMAGE_API_KEY", ""), \
                 patch.object(builtin_tools.settings, "IMAGE_MODEL", ""):
                context = builtin_tools.BuiltinToolContext(
                    root=self.root,
                    llm=EmptyVisionLlm(),
                    attachment_images=[source],
                    active_skill_names={"photo-abstract-editorial"},
                )
                result = json.loads(asyncio.run(builtin_tools.execute(
                    "image_generate",
                    {"prompt": "editorial", "action": "edit", "confirm": True},
                    context,
                )))
                issues = builtin_tools.completion_artifact_issues(context)

            self.assertFalse(result["ok"])
            self.assertEqual(result["code"], "editorial_render_failed")
            self.assertIn("有效的编辑设计 JSON", result["error"])
            self.assertEqual(context.artifacts, [])
            self.assertTrue(any("editorial_render_failed" in issue for issue in issues))

    def test_photo_editorial_uses_explicit_image_edit_panel_and_preserves_photo(self):
        class EditorialVisionLlm:
            model_id = "gpt-5.6-sol"

            async def vision(self, *_args, **_kwargs):
                return json.dumps({
                    "title": "Arcs Under Blue",
                    "subtitle": "",
                    "motif": "arches",
                    "focus_x": .58,
                    "density": 4,
                    "title_layout": "center",
                    "motif_width": .62,
                    "motif_height": .34,
                    "asymmetry": .06,
                    "mass_count": 4,
                    "horizon_count": 2,
                })

        panel_call = AsyncMock(return_value={
            "data": _editorial_panel_bytes(),
            "model": "gpt-image-test",
        })
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "source.png"
            source.write_bytes(_png_bytes())
            with patch("backend.artifacts.EXPORT_DIR", Path(temp)), \
                 patch.object(builtin_tools.settings, "IMAGE_EDIT_API_BASE_URL", "https://image.example/v1"), \
                 patch.object(builtin_tools.settings, "IMAGE_EDIT_API_KEY", "secret"), \
                 patch.object(builtin_tools.settings, "IMAGE_EDIT_MODEL", "gpt-image-test"), \
                 patch.object(builtin_tools, "_generate_editorial_panel", panel_call):
                context = builtin_tools.BuiltinToolContext(
                    root=self.root,
                    llm=EditorialVisionLlm(),
                    attachment_images=[source],
                    active_skill_names={"photo-abstract-editorial"},
                )
                result = json.loads(asyncio.run(builtin_tools.execute(
                    "image_generate",
                    {"prompt": "editorial", "action": "edit", "confirm": True},
                    context,
                )))

            self.assertTrue(result["ok"])
            self.assertEqual(result["source"], "environment_image_edit_composite")
            self.assertEqual(result["panel_renderer"], "image_edit_model")
            self.assertEqual(result["panel_model"], "gpt-image-test")
            self.assertEqual(result["renderer_version"], "photo_editorial_v3")
            self.assertTrue(all(result["quality_checks"].values()))
            panel_call.assert_awaited_once()
            with Image.open(Path(temp) / result["file"]) as rendered:
                expected_photo = Image.open(source).convert("RGB").resize(
                    (rendered.width, round(6 * rendered.width / 8)),
                    Image.Resampling.LANCZOS,
                )
                actual_photo = rendered.crop((0, 0, rendered.width, expected_photo.height))
                self.assertEqual(
                    list(_image_pixels(actual_photo)),
                    list(_image_pixels(expected_photo)),
                )

    def test_photo_editorial_text_only_client_fails_preflight(self):
        from backend.llm.client import LLMClient

        client = LLMClient(
            base_url="https://api.example.com/v1",
            api_key="secret",
            model_id="text-only-model",
            model_input=["text"],
            wire_api="responses",
        )
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "source.png"
            source.write_bytes(_png_bytes())
            context = builtin_tools.BuiltinToolContext(
                root=self.root,
                llm=client,
                attachment_images=[source],
                active_skill_names={"photo-abstract-editorial"},
            )
            with patch.object(builtin_tools.settings, "IMAGE_API_BASE_URL", ""), \
                 patch.object(builtin_tools.settings, "IMAGE_API_KEY", ""), \
                 patch.object(builtin_tools.settings, "IMAGE_MODEL", ""):
                result = json.loads(asyncio.run(builtin_tools.execute(
                    "image_generate",
                    {"prompt": "editorial", "action": "edit", "confirm": True},
                    context,
                )))

        self.assertFalse(result["ok"])
        self.assertEqual(result["code"], "image_provider_not_configured")

    def test_provider_native_image_configuration_is_not_accepted(self):
        with self.assertRaises(TypeError):
            builtin_tools.BuiltinToolContext(
                root=self.root,
                image_generation_mode="responses",
            )

    def test_general_multimodal_model_uses_same_model_for_image_input(self):
        from backend.llm.client import LLMClient

        client = LLMClient(
            base_url="https://api.example.com/v1",
            api_key="secret",
            model_id="multimodal-model",
            model_input=["text", "image"],
            wire_api="responses",
        )
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "source.png"
            source.write_bytes(_png_bytes())
            response = {
                "output": [{"type": "message", "content": [{
                    "type": "output_text", "text": "recognized",
                }]}],
            }
            with patch.object(
                client, "_request_json", new=AsyncMock(return_value=response)
            ) as request:
                result = asyncio.run(client.vision(
                    "recognize", "describe", [source],
                ))
            self.assertEqual(result, "recognized")
            payload = request.await_args.args[1]
            self.assertEqual(payload["model"], "multimodal-model")
            content = payload["input"][0]["content"]
            self.assertEqual(content[1]["type"], "input_image")
            self.assertTrue(content[1]["image_url"].startswith("data:image/png;base64,"))

    def test_word_formatting_preserves_content_and_exports_docx(self):
        from docx import Document

        source = self.root / (self.prefix + "_source.docx")
        document = Document()
        document.add_heading("项目标题", level=1)
        document.add_paragraph("正文内容保持不变。")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "字段"
        table.cell(0, 1).text = "值"
        document.save(source)
        before, _ = builtin_tools._docx_text_signature(source)

        with patch.object(builtin_tools, "EXPORT_DIR", self.root):
            inspected = self.run_tool("document_inspect", {"path": source.name})
            self.assertTrue(inspected["ok"])
            self.assertEqual(inspected["content_sha256"], before)
            formatted = self.run_tool("document_format", {
                "path": source.name,
                "output_name": "规范结果.docx",
                "east_asia_font": "宋体",
                "body_size_pt": 12,
                "line_spacing": 1.5,
                "confirm": True,
            })
            self.assertTrue(formatted["ok"])
            self.assertTrue(formatted["content_unchanged"])
            output = self.root / formatted["file"]
            self.assertTrue(output.is_file())
            after, _ = builtin_tools._docx_text_signature(output)
            self.assertEqual(after, before)

    def test_word_creation_exports_verified_text_and_table(self):
        from docx import Document

        with patch.object(builtin_tools, "EXPORT_DIR", self.root):
            created = self.run_tool("document_create", {
                "title": "101 Irregular Past Tense Verbs",
                "content": "# Tips\n- **be → was/were**：过去式根据主语变化。\n**© Woodward English**",
                "tables": [{
                    "headers": ["序号", "原形", "过去式"],
                    "rows": [["1", "be", "was/were"], ["2", "become", "became"]],
                }],
                "output_name": "图片识别结果.docx",
                "confirm": True,
            })
            self.assertTrue(created["ok"])
            self.assertEqual(created["tables"], 1)
            self.assertEqual(created["table_rows"], 2)
            output = self.root / created["file"]
            self.assertTrue(output.is_file())
            document = Document(output)
            paragraph_text = "\n".join(p.text for p in document.paragraphs)
            self.assertIn("be → was/were", paragraph_text)
            self.assertNotIn("**", paragraph_text)
            emphasized = [
                run
                for paragraph in document.paragraphs
                for run in paragraph.runs
                if run.text in {"be → was/were", "© Woodward English"}
            ]
            self.assertEqual(len(emphasized), 2)
            self.assertTrue(all(run.bold for run in emphasized))
            self.assertEqual(document.tables[0].cell(2, 2).text, "became")
            from docx.oxml.ns import qn
            table = document.tables[0]
            width = table._tbl.tblPr.find(qn("w:tblW"))
            layout = table._tbl.tblPr.find(qn("w:tblLayout"))
            grid = [int(node.get(qn("w:w"))) for node in table._tbl.tblGrid.gridCol_lst]
            header = table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader"))
            self.assertEqual(width.get(qn("w:type")), "dxa")
            self.assertEqual(sum(grid), int(width.get(qn("w:w"))))
            self.assertEqual(layout.get(qn("w:type")), "fixed")
            self.assertIsNotNone(header)

            ctx = builtin_tools.BuiltinToolContext(
                artifacts=[created["file"]], required_artifact_kinds={"document"}
            )
            self.assertEqual(builtin_tools.completion_artifact_issues(ctx), [])

    def test_missing_requested_word_artifact_is_a_completion_issue(self):
        ctx = builtin_tools.BuiltinToolContext(required_artifact_kinds={"document"})
        self.assertEqual(
            builtin_tools.completion_artifact_issues(ctx),
            ["缺少任务所需 Word 产物：没有生成可下载且有效的 DOCX Artifact"],
        )

    def test_cron_parser_and_persistence(self):
        first = scheduler.next_run(
            "*/15 9-10 * * 1-5",
            "Asia/Shanghai",
            datetime.datetime(2026, 7, 30, 1, 1, tzinfo=datetime.timezone.utc),
        )
        self.assertEqual(first.isoformat(), "2026-07-30T01:15:00")
        # 经典 Cron 在“日”和“星期”都受限时按 OR 匹配；7 也表示周日。
        by_weekday = scheduler.next_run(
            "0 9 1 * 1",
            "UTC",
            datetime.datetime(2026, 8, 2, 9, 0, tzinfo=datetime.timezone.utc),
        )
        self.assertEqual(by_weekday.isoformat(), "2026-08-03T09:00:00")
        sunday = scheduler.next_run(
            "0 9 * * 7",
            "UTC",
            datetime.datetime(2026, 8, 7, 9, 0, tzinfo=datetime.timezone.utc),
        )
        self.assertEqual(sunday.isoformat(), "2026-08-09T09:00:00")
        leap_day = scheduler.next_run(
            "0 0 29 2 *",
            "UTC",
            datetime.datetime(2026, 3, 1, 0, 0, tzinfo=datetime.timezone.utc),
        )
        self.assertEqual(leap_day.isoformat(), "2028-02-29T00:00:00")

        db_path = self.root / (self.prefix + "_schedule.db")
        engine = create_engine(f"sqlite:///{db_path.as_posix()}")
        Base.metadata.create_all(engine)
        factory = sessionmaker(bind=engine)
        db = factory()
        user = User(username="u", password_hash="x")
        db.add(user)
        db.flush()
        agent = Agent(name="a", enabled=True, created_by=user.id)
        db.add(agent)
        db.flush()
        thread = task_store.ensure_thread(
            db, thread_id="current-thread", owner_id=user.id, agent_id=agent.id,
        )
        turn = task_store.create_turn(
            db, thread=thread, input_text="初始消息", payload={}
        )
        task_store.finish_turn(db, turn, answer="初始回答")
        db.commit()
        with patch.object(scheduler, "SessionLocal", factory), \
             patch.object(jobs, "SessionLocal", factory):
            task = scheduler.create(user.id, agent.id, "daily", "0 9 * * *",
                                    "Asia/Shanghai", "生成日报")
            self.assertEqual(len(scheduler.list_tasks(user.id)), 1)
            row = db.get(ScheduledTask, task["id"])
            row.next_run_at = datetime.datetime(2026, 7, 30, 0, 0)
            db.commit()
            self.assertEqual(
                scheduler.dispatch_due(datetime.datetime(2026, 7, 30, 0, 1)), 1
            )
            created_job = db.query(Job).one()
            payload = json.loads(created_job.payload)
            self.assertEqual(payload["source"], "cron")
            self.assertEqual(payload["session_id"], "current-thread")
            self.assertEqual(row.session_id, "current-thread")
            self.assertTrue(scheduler.delete(user.id, task["id"]))
            self.assertEqual(scheduler.list_tasks(user.id), [])
            bad = scheduler.create(
                user.id, agent.id, "bad", "0 9 * * *", "Asia/Shanghai", "失败任务"
            )
            good = scheduler.create(
                user.id, agent.id, "good", "0 9 * * *", "Asia/Shanghai", "正常任务"
            )
            for task_id in (bad["id"], good["id"]):
                scheduled = db.get(ScheduledTask, task_id)
                scheduled.next_run_at = datetime.datetime(2026, 7, 30, 0, 0)
            db.commit()
            original_enqueue = jobs.enqueue_in_session

            def selective_enqueue(session, owner_id, agent_id, kind, payload, **kwargs):
                if payload.get("scheduled_task_id") == bad["id"]:
                    raise RuntimeError("poison schedule")
                return original_enqueue(
                    session, owner_id, agent_id, kind, payload, **kwargs
                )

            with patch.object(jobs, "enqueue_in_session", side_effect=selective_enqueue):
                self.assertEqual(
                    scheduler.dispatch_due(datetime.datetime(2026, 7, 30, 0, 1)), 1
                )
            db.expire_all()
            bad_row = db.get(ScheduledTask, bad["id"])
            good_row = db.get(ScheduledTask, good["id"])
            self.assertEqual(bad_row.last_dispatch_status, "retrying")
            self.assertEqual(bad_row.consecutive_failures, 1)
            self.assertIn("poison schedule", bad_row.last_error)
            self.assertTrue(good_row.last_job_id)
        db.close()
        engine.dispose()

    def test_subagent_lifecycle_uses_persistent_jobs(self):
        db_path = self.root / (self.prefix + "_agents.db")
        engine = create_engine(f"sqlite:///{db_path.as_posix()}")
        Base.metadata.create_all(engine)
        factory = sessionmaker(bind=engine)
        db = factory()
        user = User(username="u2", password_hash="x")
        db.add(user)
        db.flush()
        agent = Agent(name="worker", enabled=True, created_by=user.id)
        db.add(agent)
        db.commit()
        self.ctx.user_id = user.id
        self.ctx.agent_id = agent.id
        with patch("backend.database.SessionLocal", factory), \
             patch.object(jobs, "SessionLocal", factory):
            created = self.run_tool("spawn_agent", {"query": "检查项目"})
            self.assertTrue(created["ok"])
            task_id = created["task_id"]
            check_db = factory()
            try:
                created_row = check_db.get(Job, task_id)
                created_payload = json.loads(created_row.payload)
            finally:
                check_db.close()
            self.assertIn("execution_snapshot", created_payload)
            self.assertTrue(created_payload["delegation_authorized"])
            unrelated = jobs.enqueue(user.id, agent.id, "chat", {
                "session_id": "other-child",
                "source": "subagent",
                "parent_run_id": "other-parent",
            })
            listed = self.run_tool("list_agents", {})
            self.assertEqual(listed["tasks"][0]["task_id"], task_id)
            self.assertNotIn(unrelated, {item["task_id"] for item in listed["tasks"]})
            denied = self.run_tool("interrupt_agent", {
                "task_id": unrelated, "confirm": True,
            })
            self.assertFalse(denied["ok"])
            stopped = self.run_tool("interrupt_agent", {
                "task_id": task_id, "confirm": True,
            })
            self.assertTrue(stopped["interrupted"])
        db.close()
        engine.dispose()

    def test_builtin_tools_are_executed_by_harness_loop(self):
        path = self.root / (self.prefix + "_runtime.txt")
        path.write_text("runtime evidence", encoding="utf-8")
        llm = _BuiltinCallingLlm(path.name)
        answer, _ = asyncio.run(run_harness(
            llm,
            "可靠执行",
            "读取文件",
            builtin_context=self.ctx,
            tool_policy={
                "router": {"enabled": False},
                "max_iterations": 2,
                "max_successful_calls": 1,
            },
            verification_policy={"required": False},
        ))
        self.assertEqual(answer, "已读取并验证")
        tool_messages = [row for row in llm.messages if row.get("role") == "tool"]
        self.assertIn("runtime evidence", tool_messages[0]["content"])

    def test_image_tool_without_artifact_cannot_complete_turn(self):
        llm = _ImageFailureLlm()
        context = builtin_tools.BuiltinToolContext(root=self.root)
        with patch.object(builtin_tools.settings, "IMAGE_API_BASE_URL", ""), \
             patch.object(builtin_tools.settings, "IMAGE_API_KEY", ""), \
             patch.object(builtin_tools.settings, "IMAGE_MODEL", ""):
            with self.assertRaisesRegex(RuntimeError, "缺少任务所需图片产物"):
                asyncio.run(run_harness(
                    llm,
                    "可靠执行",
                    "生成一张图片",
                    builtin_context=context,
                    tool_policy={"router": {"enabled": False}, "max_iterations": 2},
                    verification_policy={"required": True, "max_revisions": 0},
                ))

    def test_image_ocr_to_word_completes_only_with_real_docx_artifact(self):
        llm = _DocumentCallingLlm()
        context = builtin_tools.BuiltinToolContext(
            root=self.root,
            enabled_tools={"document_create"},
            required_artifact_kinds={"document"},
        )
        with patch.object(builtin_tools, "EXPORT_DIR", self.root):
            answer, _ = asyncio.run(run_harness(
                llm,
                "可靠执行",
                "图片识别内容作为word",
                attachment_text=(
                    "### 图片附件识别\n101 Irregular Past Tense Verbs\n"
                    "1 | be | was/were"
                ),
                builtin_context=context,
                tool_policy={"router": {"enabled": False}, "max_iterations": 2},
                verification_policy={"required": True, "max_revisions": 0},
            ))
            self.assertEqual(answer, "已生成可下载的 Word 文档。")
            self.assertEqual(len(context.artifacts), 1)
            self.assertTrue((self.root / context.artifacts[0]).is_file())


if __name__ == "__main__":
    unittest.main()
