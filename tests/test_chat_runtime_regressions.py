"""对话运行时与命令选择的回归测试。"""
import asyncio
import inspect
import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import httpx
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from starlette.datastructures import FormData, UploadFile

from backend import attachments, jobs, worker
from backend.api import chat as chat_api
from backend.api.chat import (
    _end_line, _event_line, _validate_invocations, active_chat_jobs, agent_chat_statuses, chat_models,
    select_chat_provider,
)
from backend.api.providers import _upsert_provider, discover_models
from backend.capabilities import templates as template_capability
from backend.database import Base
from backend.llm.chatgpt_client import ChatGPTClient
from backend.llm.client import LLMClient, PROVIDER_PRESETS, client_for_provider
from backend.llm.codex_models import (
    available_chatgpt_codex_models,
    normalize_chatgpt_codex_model,
    preferred_chatgpt_codex_model,
)
from backend.logging_utils import redact_log_value
from backend.models import Agent, Attachment, Job, McpServer, ModelProvider, Skill, Template, Thread, Turn, User
from backend.schemas import ProviderProbe


class _StreamResponse:
    status_code = 200
    headers = {"content-type": "text/event-stream"}

    async def __aenter__(self):
        return self

    async def __aexit__(self, *_args):
        return False

    def raise_for_status(self):
        return None

    async def aiter_lines(self):
        yield 'data: {"choices":[{"delta":{"content":"你"}}]}'
        yield 'data: {"choices":[{"delta":{"content":"好"}}]}'
        yield "data: [DONE]"


class _FakeHttpClient:
    def stream(self, *_args, **_kwargs):
        return _StreamResponse()


class _AnthropicStreamResponse(_StreamResponse):
    async def aiter_lines(self):
        yield 'event: content_block_delta'
        yield 'data: {"type":"content_block_delta","delta":{"type":"text_delta","text":"你"}}'
        yield 'data: {"type":"content_block_delta","delta":{"type":"text_delta","text":"好"}}'
        yield 'data: {"type":"message_stop"}'


class _AnthropicHttpClient:
    def stream(self, *_args, **_kwargs):
        return _AnthropicStreamResponse()


class _ReasoningOnlyStreamResponse(_StreamResponse):
    async def aiter_lines(self):
        yield (
            'data: {"choices":[{"delta":{"reasoning_content":"内部分析，不应展示"},'
            '"finish_reason":null}]}'
        )
        yield 'data: {"choices":[{"delta":{},"finish_reason":"length"}]}'
        yield "data: [DONE]"


class _WhitespaceOnlyStreamResponse(_StreamResponse):
    async def aiter_lines(self):
        yield 'data: {"choices":[{"delta":{"content":"   "},"finish_reason":null}]}'
        yield 'data: {"choices":[{"delta":{},"finish_reason":"stop"}]}'
        yield "data: [DONE]"


class _ReasoningRecoveryHttpClient:
    def __init__(self):
        self.stream_payloads = []
        self.post_payloads = []

    def stream(self, *_args, **kwargs):
        self.stream_payloads.append(dict(kwargs["json"]))
        return _ReasoningOnlyStreamResponse()

    async def post(self, url, **kwargs):
        self.post_payloads.append(dict(kwargs["json"]))
        request = httpx.Request("POST", url)
        if (
            kwargs["json"].get("stream") is False
            and "stream_options" in kwargs["json"]
        ):
            return httpx.Response(
                400,
                request=request,
                json={"error": {"message": (
                    "Stream options can only be defined when stream=True."
                )}},
            )
        return httpx.Response(
            200,
            request=request,
            json={"choices": [{
                "message": {
                    "role": "assistant",
                    "content": "这是恢复后的最终答复",
                    "reasoning_content": "隐藏推理",
                },
                "finish_reason": "stop",
            }]},
        )


class _WhitespaceRecoveryHttpClient(_ReasoningRecoveryHttpClient):
    def stream(self, *_args, **kwargs):
        self.stream_payloads.append(dict(kwargs["json"]))
        return _WhitespaceOnlyStreamResponse()


class _CapturingStreamHttpClient:
    def __init__(self):
        self.payloads = []

    def stream(self, *_args, **kwargs):
        self.payloads.append(dict(kwargs["json"]))
        return _StreamResponse()


class _ReasoningOnlyNonstreamHttpClient:
    def __init__(self):
        self.payloads = []

    async def post(self, url, **kwargs):
        self.payloads.append(dict(kwargs["json"]))
        request = httpx.Request("POST", url)
        if len(self.payloads) == 1:
            body = {"choices": [{"message": {
                "role": "assistant",
                "content": None,
                "reasoning_content": "只有内部推理",
            }}]}
        else:
            body = {"choices": [{"message": {
                "role": "assistant",
                "content": "恢复后的非流式答案",
            }}]}
        return httpx.Response(200, request=request, json=body)


class _DiscoveryClient:
    async def list_models(self):
        return ["model-a", "model-b"]


class _FormRequest:
    def __init__(self, entries):
        self._form = FormData(entries)
        self.headers = {}

    async def form(self):
        return self._form


class _TokenCompatibilityHttpClient:
    def __init__(self):
        self.payloads = []

    async def post(self, url, **kwargs):
        self.payloads.append(dict(kwargs["json"]))
        request = httpx.Request("POST", url)
        if len(self.payloads) == 1:
            return httpx.Response(
                400,
                request=request,
                json={"error": {"message": (
                    "Unsupported parameter: 'max_completion_tokens' is not supported "
                    "with this model. Use 'max_tokens' instead."
                )}},
            )
        return httpx.Response(
            200,
            request=request,
            json={"choices": [{"message": {"role": "assistant", "content": "pong"}}]},
        )


class _HarmonyHeaderRetryHttpClient:
    def __init__(self):
        self.payloads = []

    async def post(self, url, **kwargs):
        self.payloads.append(dict(kwargs["json"]))
        request = httpx.Request("POST", url)
        if len(self.payloads) == 1:
            return httpx.Response(
                200,
                request=request,
                json={"message": (
                    "unexpected tokens remaining in message header: "
                    'Some("to=functions.web_search")'
                )},
            )
        return httpx.Response(
            200,
            request=request,
            json={"choices": [{"message": {
                "role": "assistant",
                "content": "",
                "tool_calls": [{
                    "id": "call_1",
                    "type": "function",
                    "function": {
                        "name": "web_search",
                        "arguments": '{"query":"美股资金流向"}',
                    },
                }],
            }}]},
        )


class ChatRuntimeRegressionTests(unittest.TestCase):
    def test_full_access_policy_requires_agent_manager_and_rejects_unknown_values(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        db = sessionmaker(bind=engine)()
        owner = User(username="policy-owner", password_hash="x", role="admin")
        visitor = User(username="policy-visitor", password_hash="x", role="user")
        db.add_all([owner, visitor])
        db.flush()
        agent = Agent(
            name="public-policy-agent",
            enabled=True,
            is_public=True,
            created_by=owner.id,
        )
        db.add(agent)
        db.commit()

        with patch.object(chat_api, "resolve_agent", return_value=agent), \
             patch.object(chat_api, "enforce"):
            with self.assertRaises(HTTPException) as forbidden:
                asyncio.run(chat_api.chat(
                    _FormRequest([
                        ("agent_id", str(agent.id)),
                        ("query", "执行任务"),
                        ("approval_policy", "full_access"),
                    ]),
                    visitor,
                    db,
                ))
            self.assertEqual(forbidden.exception.status_code, 403)

            with self.assertRaises(HTTPException) as invalid:
                asyncio.run(chat_api.chat(
                    _FormRequest([
                        ("agent_id", str(agent.id)),
                        ("query", "执行任务"),
                        ("approval_policy", "unknown"),
                    ]),
                    owner,
                    db,
                ))
            self.assertEqual(invalid.exception.status_code, 400)

        db.close()
        engine.dispose()

    def test_chat_turn_freezes_root_selected_full_access_policy(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        factory = sessionmaker(bind=engine)
        db = factory()
        owner = User(username="policy-root", password_hash="x", role="root")
        db.add(owner)
        db.flush()
        agent = Agent(
            name="private-policy-agent",
            enabled=True,
            active_version=1,
            created_by=owner.id,
        )
        db.add(agent)
        db.commit()

        snapshot = {"provider": {}, "builtin_tools": [], "approval_policy": "full_access"}
        with patch.object(jobs, "SessionLocal", factory), \
             patch.object(chat_api, "resolve_agent", return_value=agent), \
             patch.object(chat_api, "enforce"), \
             patch.object(chat_api, "_validate_invocations"), \
             patch.object(
                 chat_api, "select_chat_provider", return_value=SimpleNamespace(id=1)
             ), \
             patch.object(
                 chat_api, "build_execution_snapshot", return_value=snapshot
             ) as build_snapshot:
            result = asyncio.run(chat_api.chat(
                _FormRequest([
                    ("agent_id", str(agent.id)),
                    ("query", "执行已授权任务"),
                    ("approval_policy", "full_access"),
                ]),
                owner,
                db,
            ))

        payload = json.loads(db.get(Job, result["turn_id"]).payload)
        self.assertEqual(result["approval_policy"], "full_access")
        self.assertEqual(payload["approval_policy"], "full_access")
        self.assertEqual(payload["execution_snapshot"]["approval_policy"], "full_access")
        self.assertEqual(build_snapshot.call_args.kwargs["approval_policy"], "full_access")
        db.close()
        engine.dispose()

    def test_codex_chatgpt_legacy_model_uses_account_catalog_default(self):
        with tempfile.TemporaryDirectory() as temp:
            codex_dir = Path(temp)
            (codex_dir / "models_cache.json").write_text(json.dumps({
                "models": [
                    {"slug": "gpt-current", "supported_in_api": True},
                    {"slug": "gpt-internal-wm", "supported_in_api": False},
                    {"slug": "codex-auto-review", "supported_in_api": True},
                ],
            }), encoding="utf-8")
            (codex_dir / "config.toml").write_text(
                'model = "gpt-current"\n', encoding="utf-8"
            )

            self.assertEqual(available_chatgpt_codex_models(codex_dir), ["gpt-current"])
            self.assertEqual(preferred_chatgpt_codex_model(codex_dir), "gpt-current")
            self.assertEqual(
                normalize_chatgpt_codex_model("gpt-5", codex_dir),
                "gpt-current",
            )

    def test_chatgpt_provider_preset_and_runtime_use_codex_responses_models(self):
        preset = PROVIDER_PRESETS["chatgpt"]
        self.assertEqual(preset["wire_api"], "responses")
        self.assertNotEqual(preset["model_id"], "gpt-5")
        self.assertNotIn("gpt-5", preset["models"])

        provider = SimpleNamespace(
            id=None,
            provider_type="openai",  # 旧版表单曾错误改成 openai
            base_url="https://chatgpt.com/backend-api/codex",
            model_id="gpt-5",
            model_input='["text"]',
            auth_extra="{}",
            api_key="token",
            context_window=0,
        )
        client = client_for_provider(provider)
        self.assertIsInstance(client, ChatGPTClient)
        self.assertNotEqual(client.text_model, "gpt-5")

    def test_configured_text_only_provider_does_not_hide_a_vision_fallback(self):
        provider = SimpleNamespace(
            id=9,
            provider_type="openai",
            base_url="https://api.example.com/v1",
            api_key="secret",
            model_id="text-only-model",
            model_input='["text"]',
            model_reasoning=False,
        )

        client = client_for_provider(provider)

        self.assertEqual(client.vision_model, "")
        self.assertIsNone(client.vision_fallback)

    def test_codex_reimport_reuses_legacy_provider_by_base_url(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        db = sessionmaker(bind=engine)()
        owner = User(username="codex-owner", password_hash="x", role="root")
        db.add(owner)
        db.flush()
        legacy = ModelProvider(
            name="ChatGPT legacy Codex",
            provider_type="openai",
            base_url="https://chatgpt.com/backend-api/codex",
            model_id="gpt-5",
            wire_api="chat_completions",
            created_by=owner.id,
        )
        db.add(legacy)
        db.flush()

        repaired = _upsert_provider(
            db,
            owner,
            "ChatGPT 订阅 (Codex)",
            "chatgpt",
            PROVIDER_PRESETS["chatgpt"],
        )
        self.assertEqual(repaired.id, legacy.id)
        self.assertEqual(repaired.name, "ChatGPT 订阅 (Codex)")
        self.assertEqual(repaired.provider_type, "chatgpt")
        self.assertEqual(repaired.model_id, "gpt-5.6-sol")
        self.assertEqual(repaired.wire_api, "responses")
        db.close()
        engine.dispose()

    def test_openai_compatible_client_streams_message_deltas(self):
        client = LLMClient(
            base_url="https://example.com/v1",
            api_key="secret",
            model_id="model",
        )
        deltas = []
        fake = _CapturingStreamHttpClient()

        async def run():
            with patch("backend.llm.client.get_http_client", return_value=fake):
                return await client.chat_messages_stream(
                    [{"role": "user", "content": "hello"}],
                    on_delta=deltas.append,
                )

        self.assertEqual(asyncio.run(run()), "你好")
        self.assertEqual(deltas, ["你", "好"])
        self.assertEqual(fake.payloads[0]["stream_options"], {"include_usage": True})

    def test_stream_http_error_body_is_read_before_parameter_fallback(self):
        client = LLMClient(
            base_url="https://example.com/v1",
            api_key="secret",
            model_id="model",
            wire_api="chat_completions",
            stream_max_retries=0,
        )
        payloads = []

        async def handler(request):
            payload = json.loads(request.content)
            payloads.append(payload)
            if "stream_options" in payload:
                return httpx.Response(
                    400,
                    json={"error": {"message": (
                        "Unsupported parameter: 'stream_options'."
                    )}},
                )
            return httpx.Response(
                200,
                headers={"content-type": "text/event-stream"},
                content=(
                    'data: {"choices":[{"delta":{"content":"恢复成功"}}]}\n\n'
                    "data: [DONE]\n\n"
                ).encode(),
            )

        async def run():
            fake = httpx.AsyncClient(transport=httpx.MockTransport(handler))
            try:
                with patch("backend.llm.client.get_http_client", return_value=fake):
                    return await client.chat_messages_stream(
                        [{"role": "user", "content": "hello"}],
                    )
            finally:
                await fake.aclose()

        self.assertEqual(asyncio.run(run()), "恢复成功")
        self.assertEqual(len(payloads), 2)
        self.assertIn("stream_options", payloads[0])
        self.assertNotIn("stream_options", payloads[1])

    def test_responses_stream_reads_error_body_before_temperature_fallback(self):
        client = LLMClient(
            base_url="https://example.com/v1",
            api_key="secret",
            model_id="gpt-test",
            wire_api="responses",
            supports_temperature=True,
            stream_max_retries=0,
        )
        payloads = []

        async def handler(request):
            payload = json.loads(request.content)
            payloads.append(payload)
            if "temperature" in payload:
                return httpx.Response(
                    400,
                    json={"error": {"message": (
                        "Unsupported parameter: 'temperature' is not supported with this model."
                    )}},
                )
            return httpx.Response(
                200,
                headers={"content-type": "text/event-stream"},
                content=(
                    'data: {"type":"response.output_text.delta","delta":"恢复成功"}\n\n'
                    "data: [DONE]\n\n"
                ).encode(),
            )

        async def run():
            fake = httpx.AsyncClient(transport=httpx.MockTransport(handler))
            try:
                with patch("backend.llm.client.get_http_client", return_value=fake):
                    return await client.chat_messages_stream(
                        [{"role": "user", "content": "hello"}],
                    )
            finally:
                await fake.aclose()

        self.assertEqual(asyncio.run(run()), "恢复成功")
        self.assertEqual(len(payloads), 2)
        self.assertIn("temperature", payloads[0])
        self.assertNotIn("temperature", payloads[1])

    def test_stream_http_error_preserves_upstream_detail(self):
        client = LLMClient(
            base_url="https://example.com/v1",
            api_key="secret",
            model_id="model",
            wire_api="responses",
            stream_max_retries=0,
        )

        async def handler(_request):
            return httpx.Response(
                429,
                json={"error": {
                    "code": "rate_limit_exceeded",
                    "message": "request capacity reached",
                }},
            )

        async def run():
            fake = httpx.AsyncClient(transport=httpx.MockTransport(handler))
            try:
                with patch("backend.llm.client.get_http_client", return_value=fake):
                    return await client.chat_messages_stream(
                        [{"role": "user", "content": "hello"}],
                    )
            finally:
                await fake.aclose()

        with self.assertRaisesRegex(RuntimeError, "request capacity reached") as caught:
            asyncio.run(run())
        self.assertNotIn("ResponseNotRead", str(caught.exception))

    def test_anthropic_messages_use_top_level_system_and_x_api_key(self):
        client = LLMClient(
            base_url="https://api.example.com/v1",
            api_key="secret",
            model_id="claude-test",
            provider_type="anthropic",
            wire_api="messages",
            auth_type="x_api_key",
            api_version="2023-06-01",
            api_version_mode="header",
        )
        payload = client._anthropic_payload([
            {"role": "system", "content": "system prompt"},
            {"role": "user", "content": "hello"},
        ])
        self.assertEqual(client.headers["x-api-key"], "secret")
        self.assertEqual(client.headers["anthropic-version"], "2023-06-01")
        self.assertEqual(payload["system"], "system prompt")
        self.assertNotIn("system", [message["role"] for message in payload["messages"]])
        self.assertEqual(payload["max_tokens"], 8192)

    def test_anthropic_streams_text_delta_events(self):
        client = LLMClient(
            base_url="https://api.example.com/v1",
            api_key="secret",
            model_id="claude-test",
            wire_api="messages",
            auth_type="x_api_key",
        )
        deltas = []

        async def run():
            with patch("backend.llm.client.get_http_client", return_value=_AnthropicHttpClient()):
                return await client.chat_messages_stream(
                    [{"role": "user", "content": "hello"}],
                    on_delta=deltas.append,
                )

        self.assertEqual(asyncio.run(run()), "你好")
        self.assertEqual(deltas, ["你", "好"])

    def test_reasoning_only_stream_uses_single_nonstream_recovery(self):
        client = LLMClient(
            base_url="https://integrate.api.nvidia.com/v1",
            api_key="secret",
            model_id="openai/gpt-oss-20b",
            stream_max_retries=3,
        )
        fake = _ReasoningRecoveryHttpClient()
        deltas = []

        async def run():
            with patch("backend.llm.client.get_http_client", return_value=fake):
                return await client.chat_messages_stream(
                    [{"role": "user", "content": "给出最终答复"}],
                    on_delta=deltas.append,
                )

        self.assertEqual(asyncio.run(run()), "这是恢复后的最终答复")
        self.assertEqual(deltas, ["这是恢复后的最终答复"])
        self.assertEqual(len(fake.stream_payloads), 1)
        self.assertEqual(len(fake.post_payloads), 1)
        self.assertFalse(fake.post_payloads[0]["stream"])
        self.assertNotIn("stream_options", fake.post_payloads[0])
        self.assertEqual(fake.post_payloads[0]["reasoning_effort"], "low")

    def test_harmony_channel_suffix_is_removed_from_tool_name(self):
        data = {"choices": [{"message": {
            "role": "assistant",
            "content": "",
            "tool_calls": [{
                "id": "call_1",
                "type": "function",
                "function": {
                    "name": "web_fetch<|channel|>commentary",
                    "arguments": '{"url":"https://example.test"}',
                },
            }],
        }}]}

        parsed = LLMClient._first_message(data)

        self.assertEqual(
            parsed["tool_calls"][0]["function"]["name"],
            "web_fetch",
        )
        self.assertEqual(
            data["choices"][0]["message"]["tool_calls"][0]["function"]["name"],
            "web_fetch<|channel|>commentary",
        )

    def test_whitespace_only_stream_is_not_treated_as_final_text(self):
        client = LLMClient(
            base_url="https://integrate.api.nvidia.com/v1",
            api_key="secret",
            model_id="openai/gpt-oss-20b",
            stream_max_retries=3,
        )
        fake = _WhitespaceRecoveryHttpClient()
        deltas = []

        async def run():
            with patch("backend.llm.client.get_http_client", return_value=fake):
                return await client.chat_messages_stream(
                    [{"role": "user", "content": "生成最终答复"}],
                    on_delta=deltas.append,
                )

        self.assertEqual(asyncio.run(run()), "这是恢复后的最终答复")
        self.assertEqual("".join(deltas).strip(), "这是恢复后的最终答复")
        self.assertEqual(len(fake.stream_payloads), 1)
        self.assertEqual(len(fake.post_payloads), 1)

    def test_reasoning_only_tool_response_recovers_final_answer(self):
        client = LLMClient(
            base_url="https://integrate.api.nvidia.com/v1",
            api_key="secret",
            model_id="openai/gpt-oss-20b",
        )
        fake = _ReasoningOnlyNonstreamHttpClient()

        async def run():
            with patch("backend.llm.client.get_http_client", return_value=fake):
                return await client.chat_with_tools(
                    [{"role": "user", "content": "总结已有工具结果"}],
                    tools=[{"type": "function", "function": {
                        "name": "search",
                        "description": "检索",
                        "parameters": {"type": "object"},
                    }}],
                )

        result = asyncio.run(run())
        self.assertEqual(result["content"], "恢复后的非流式答案")
        self.assertEqual(len(fake.payloads), 2)
        self.assertIn("tools", fake.payloads[0])
        self.assertNotIn("tools", fake.payloads[1])
        self.assertEqual(fake.payloads[1]["reasoning_effort"], "low")
        self.assertIn("不得返回空内容", fake.payloads[1]["messages"][-1]["content"])

    def test_gpt_oss_retries_harmony_header_error_with_stable_tool_temperature(self):
        client = LLMClient(
            base_url="https://integrate.api.nvidia.com/v1",
            api_key="secret",
            model_id="openai/gpt-oss-20b",
            max_retries=1,
        )
        fake = _HarmonyHeaderRetryHttpClient()

        async def run():
            with (
                patch("backend.llm.client.get_http_client", return_value=fake),
                patch("backend.llm.client.asyncio.sleep", new_callable=AsyncMock),
            ):
                return await client.chat_with_tools(
                    [{"role": "user", "content": "今日美股资金流向复盘"}],
                    tools=[{"type": "function", "function": {
                        "name": "web_search",
                        "description": "联网检索",
                        "parameters": {"type": "object"},
                    }}],
                )

        result = asyncio.run(run())
        self.assertEqual(result["tool_calls"][0]["function"]["name"], "web_search")
        self.assertEqual(len(fake.payloads), 2)
        self.assertEqual([item["temperature"] for item in fake.payloads], [0.2, 0.2])

    def test_nemotron_final_answer_disables_thinking(self):
        client = LLMClient(
            base_url="https://integrate.api.nvidia.com/v1",
            api_key="secret",
            model_id="nvidia/nemotron-3-super-120b-a12b",
        )
        fake = _CapturingStreamHttpClient()

        async def run():
            with patch("backend.llm.client.get_http_client", return_value=fake):
                return await client.chat_messages_stream(
                    [{"role": "user", "content": "生成最终答复"}],
                )

        self.assertEqual(asyncio.run(run()), "你好")
        self.assertFalse(
            fake.payloads[0]["chat_template_kwargs"]["enable_thinking"]
        )
        self.assertNotIn("reasoning_budget", fake.payloads[0])
        self.assertNotIn("reasoning_effort", fake.payloads[0])
        self.assertFalse(client._has_assistant_output({
            "content": "<think>只有内部推理</think>",
            "tool_calls": [],
        }))

    def test_nemotron_empty_tool_turn_recovers_without_thinking(self):
        client = LLMClient(
            base_url="https://integrate.api.nvidia.com/v1",
            api_key="secret",
            model_id="nvidia/nemotron-3-super-120b-a12b",
        )
        fake = _ReasoningOnlyNonstreamHttpClient()

        async def run():
            with patch("backend.llm.client.get_http_client", return_value=fake):
                return await client.chat_with_tools(
                    [{"role": "user", "content": "总结工具结果"}],
                    tools=[{"type": "function", "function": {
                        "name": "search",
                        "description": "检索",
                        "parameters": {"type": "object"},
                    }}],
                )

        self.assertEqual(asyncio.run(run())["content"], "恢复后的非流式答案")
        recovery = fake.payloads[1]
        self.assertFalse(recovery["chat_template_kwargs"]["enable_thinking"])
        self.assertNotIn("reasoning_effort", recovery)

    def test_responses_api_converts_tools_and_parses_function_calls(self):
        client = LLMClient(
            base_url="https://api.example.com/v1",
            api_key="secret",
            model_id="gpt-test",
            wire_api="responses",
            reasoning_effort="medium",
        )
        payload = client._responses_payload(
            [{"role": "system", "content": "be concise"}, {"role": "user", "content": "search"}],
            [{"type": "function", "function": {
                "name": "search", "description": "Search", "parameters": {"type": "object"},
            }}],
        )
        self.assertEqual(payload["instructions"], "be concise")
        self.assertEqual(payload["reasoning"], {"effort": "medium"})
        self.assertEqual(payload["tools"][0]["name"], "search")
        parsed = client._parse_responses({"output": [{
            "type": "function_call", "call_id": "call_1", "name": "search",
            "arguments": '{"q":"test"}',
        }]})
        self.assertEqual(parsed["tool_calls"][0]["function"]["name"], "search")

    def test_unsaved_provider_can_discover_models(self):
        body = ProviderProbe(
            provider_type="openai",
            base_url="https://api.example.com/v1",
            api_key="secret",
            wire_api="responses",
            model_id="",
            custom_headers={"X-Route": "qa"},
        )

        async def fake_list_models(client):
            self.assertIsNone(client.provider_id)
            return ["model-a", "model-b"]

        async def run():
            with patch.object(LLMClient, "list_models", fake_list_models):
                return await discover_models(body, None)

        self.assertEqual(asyncio.run(run()), {"models": ["model-a", "model-b"], "source": "live"})

    def test_chat_completion_auto_token_field_falls_back_for_legacy_gateway(self):
        client = LLMClient(
            base_url="https://api.example.com/v1",
            api_key="secret",
            model_id="model",
            wire_api="chat_completions",
            max_tokens_param="auto",
            max_retries=0,
        )
        fake = _TokenCompatibilityHttpClient()

        async def run():
            with patch("backend.llm.client.get_http_client", return_value=fake):
                return await client.chat_with_tools(
                    [{"role": "user", "content": "ping"}], temperature=0.2
                )

        self.assertEqual(asyncio.run(run())["content"], "pong")
        self.assertIn("max_completion_tokens", fake.payloads[0])
        self.assertNotIn("max_tokens", fake.payloads[0])
        self.assertIn("max_tokens", fake.payloads[1])
        self.assertNotIn("max_completion_tokens", fake.payloads[1])

    def test_exact_max_tokens_error_switches_to_max_completion_tokens(self):
        client = LLMClient(
            base_url="https://api.example.com/v1",
            api_key="secret",
            model_id="model",
            wire_api="chat_completions",
            max_tokens_param="auto",
        )
        request = httpx.Request("POST", "https://api.example.com/v1/chat/completions")
        response = httpx.Response(
            400,
            request=request,
            json={"error": {"message": (
                "Unsupported parameter: 'max_tokens' is not supported with this model. "
                "Use 'max_completion_tokens' instead."
            )}},
        )
        error = httpx.HTTPStatusError("bad request", request=request, response=response)
        payload = {"model": "model", "max_tokens": 4096}

        self.assertTrue(client._apply_parameter_compatibility(payload, error))
        self.assertNotIn("max_tokens", payload)
        self.assertEqual(payload["max_completion_tokens"], 4096)

    def test_bound_private_capability_can_be_invoked_but_unbound_one_cannot(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        db = sessionmaker(bind=engine)()
        owner = User(username="owner", password_hash="x", role="root")
        user = User(username="user", password_hash="x", role="user")
        db.add_all([owner, user])
        db.flush()
        bound = Skill(name="bound", created_by=owner.id, enabled=True, is_public=False)
        private = Skill(name="private", created_by=owner.id, enabled=True, is_public=False)
        mcp = McpServer(name="bound-mcp", created_by=owner.id, enabled=True, is_public=False)
        db.add_all([bound, private, mcp])
        db.flush()
        agent = Agent(
            name="agent",
            enabled=True,
            is_default=True,
            skill_ids=json.dumps([bound.id]),
            mcp_ids=json.dumps([mcp.id]),
            created_by=owner.id,
        )
        db.add(agent)
        db.flush()

        _validate_invocations(db, user, agent, [bound.id], [mcp.id], [])
        with self.assertRaises(HTTPException):
            _validate_invocations(db, user, agent, [private.id], [], [])
        db.close()
        engine.dispose()

    def test_chat_model_picker_only_exposes_public_models(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        db = sessionmaker(bind=engine)()
        owner = User(username="owner-model", password_hash="x", role="root")
        user = User(username="chat-user", password_hash="x", role="user")
        db.add_all([owner, user])
        db.flush()
        private = ModelProvider(
            name="private-model", base_url="https://private.example/v1",
            model_id="private-llm", enabled=True, is_public=False, created_by=owner.id,
        )
        public = ModelProvider(
            name="public-model", base_url="https://public.example/v1",
            model_id="public-llm", enabled=True, is_public=True, created_by=owner.id,
        )
        db.add_all([private, public])
        db.flush()
        agent = Agent(
            name="model-agent", enabled=True, is_default=True,
            provider_id=private.id, created_by=owner.id,
        )
        db.add(agent)
        db.flush()

        self.assertEqual(select_chat_provider(db, user, agent, None), private)
        self.assertEqual(select_chat_provider(db, user, agent, public.id), public)
        with self.assertRaisesRegex(ValueError, "未开放"):
            select_chat_provider(db, user, agent, private.id)
        catalog = chat_models(agent.id, user, db)
        self.assertEqual(catalog["default"]["model"], "private-llm")
        self.assertEqual(
            [item["model"] for item in catalog["items"]], ["public-llm"]
        )
        db.close()
        engine.dispose()

    def test_sample_word_template_replaces_original_body(self):
        from docx import Document

        workspace = Path(__file__).resolve().parents[1]
        source = workspace / "data" / "templates" / "_test_sample_replace.docx"
        output = None
        try:
            document = Document()
            document.add_paragraph("旧示例题：1 + 1 = 2")
            document.save(source)

            self.assertIn(
                "旧示例题",
                template_capability.extract_reference_text(source, "word"),
            )
            output = template_capability.render(
                source,
                "word",
                {},
                title="替换测试",
                append_body="# 新试题\n1. 新生成的问题\n答案：A",
                replace_body=True,
            )
            text = "\n".join(p.text for p in Document(output).paragraphs)
            self.assertNotIn("旧示例题", text)
            self.assertIn("新生成的问题", text)
            self.assertIn("答案：A", text)
        finally:
            source.unlink(missing_ok=True)
            if output is not None:
                output.unlink(missing_ok=True)

    def test_sample_ppt_template_keeps_pages_and_writes_paginated_answer(self):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "sample.pptx"
            presentation = Presentation()
            blank = presentation.slide_layouts[6]

            cover = presentation.slides.add_slide(blank)
            cover_title = cover.shapes.add_textbox(
                Inches(1), Inches(2), Inches(11), Inches(1)
            )
            cover_title.text_frame.paragraphs[0].add_run().font.size = Pt(30)
            cover_title.text_frame.paragraphs[0].runs[0].text = "旧封面标题"
            cover.shapes.add_textbox(
                Inches(1), Inches(3.2), Inches(11), Inches(1)
            ).text = "旧封面信息"

            for number in (2, 3):
                slide = presentation.slides.add_slide(blank)
                slide.shapes.add_textbox(
                    Inches(0.6), Inches(0.2), Inches(12), Inches(0.6)
                ).text = "请 输 入 您 的 标 题"
                slide.shapes.add_textbox(
                    Inches(0.6), Inches(1.1), Inches(12), Inches(4.8)
                ).text = f"旧示例正文{number}"

            closing = presentation.slides.add_slide(blank)
            closing.shapes.add_textbox(
                Inches(1), Inches(2), Inches(11), Inches(1)
            ).text = "汇报结束，谢谢！"
            presentation.save(source)

            answer = """## 第1页｜封面
**新培训标题**
新副标题

## 第2页｜第一部分
**第一部分标题**
- 第一条内容
- 第二条内容

## 第3页｜第二部分
**第二部分标题**
这是第二部分正文。

## 第4页｜总结
**培训总结**
完成全部培训内容。
"""
            with patch.object(template_capability, "EXPORT_DIR", root):
                output = template_capability.render(
                    source,
                    "ppt",
                    {},
                    title="PPT套版回归",
                    append_body=answer,
                    replace_body=True,
                )

            rendered = Presentation(output)
            self.assertEqual(len(rendered.slides), 4)
            slide_two_text = "\n".join(
                shape.text
                for shape in template_capability._iter_pptx_shapes(
                    rendered.slides[1].shapes
                )
                if getattr(shape, "has_text_frame", False)
            )
            slide_three_text = "\n".join(
                shape.text
                for shape in template_capability._iter_pptx_shapes(
                    rendered.slides[2].shapes
                )
                if getattr(shape, "has_text_frame", False)
            )
            self.assertIn("第一部分标题", slide_two_text)
            self.assertIn("第一条内容", slide_two_text)
            self.assertIn("第二部分标题", slide_three_text)
            self.assertIn("这是第二部分正文", slide_three_text)
            self.assertNotIn("请 输 入 您 的 标 题", slide_two_text)

    def test_template_ppt_export_replaces_false_unavailable_delivery(self):
        answer, artifacts = chat_api._prefer_template_artifacts(
            "当前会话未提供 `.pptx` 文件创建与导出能力，因此暂时无法生成可下载的PPT文件。\n\n## 第1页｜封面",
            ["通用中间文件.pptx", "检查报告.txt"],
            ["汇报PPT_修复后.pptx"],
        )

        self.assertNotIn("无法生成", answer)
        self.assertIn("汇报PPT_修复后.pptx", answer)
        self.assertIn("下载套用模板后的PPT", answer)
        self.assertNotIn("通用中间文件.pptx", artifacts)
        self.assertIn("检查报告.txt", artifacts)
        self.assertIn("汇报PPT_修复后.pptx", artifacts)

    def test_presentation_postprocessor_completes_deferred_plan_steps(self):
        from backend.runtime.builtin_tools import BuiltinToolContext

        events = []

        async def record(event_type, payload):
            events.append((event_type, payload))

        context = BuiltinToolContext(
            plan_revision=3,
            plan_steps=[
                {"id": "step_1", "step": "梳理培训内容与页纲", "status": "completed"},
                {"id": "step_2", "step": "套用模板生成演示稿", "status": "blocked"},
                {"id": "step_3", "step": "逐页校验并交付", "status": "blocked"},
            ],
        )

        asyncio.run(chat_api._complete_presentation_plan_steps(context, record))

        self.assertEqual(context.plan_revision, 4)
        self.assertEqual(
            [item["status"] for item in context.plan_steps],
            ["completed", "completed", "completed"],
        )
        updated = next(payload for name, payload in events if name == "plan.updated")
        self.assertEqual(updated["revision"], 4)
        self.assertEqual(
            [name for name, _payload in events].count("step.completed"),
            2,
        )

    def test_ppt_parser_removes_markdown_markers_split_across_lines(self):
        specs = template_capability._parse_pptx_body(
            "## 第1页｜结束页\n**规范使用平台  \n提升协同效率**"
        )

        self.assertEqual(specs[0]["title"], "规范使用平台")
        self.assertEqual(specs[0]["body"], ["提升协同效率"])

    def test_ppt_section_title_is_moved_inside_the_slide_edge(self):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = slide.shapes.add_textbox(
            -Inches(0.2), Inches(4.2), Inches(4.5), Inches(0.8)
        )
        run = title.text_frame.paragraphs[0].add_run()
        run.text = "旧章节标题"
        run.font.size = Pt(36)
        detail = slide.shapes.add_textbox(
            Inches(1.2), Inches(5.2), Inches(3), Inches(0.5)
        )
        detail.text = "旧章节说明"

        template_capability._fill_pptx_section_slide(
            slide,
            {"title": "业务办理", "body": ["启动 aTrust VPN"]},
        )

        self.assertGreaterEqual(title.left, Inches(0.55))
        self.assertEqual(title.text, "业务办理")
        self.assertEqual(detail.text, "启动 aTrust VPN")

    def test_word_sample_template_uses_attached_docx_as_body(self):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Pt

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            template_path = root / "template.docx"
            source_path = root / "source.docx"

            template = Document()
            template.add_table(rows=1, cols=1).cell(0, 0).text = "固定版头构件"
            template.add_paragraph("固定文号区")
            title = template.add_paragraph()
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title.add_run("关于××事项的请示").font.size = Pt(22)
            template.add_paragraph("××单位（主送单位）：")
            body_sample = template.add_paragraph("模板示例正文，不得残留。")
            body_sample.paragraph_format.first_line_indent = Pt(32)
            heading_sample = template.add_paragraph("一、模板示例标题")
            heading_sample.runs[0].font.name = "仿宋"
            template.add_paragraph("妥否，请示。")
            template.add_paragraph()
            template.add_paragraph()
            closing = template.add_paragraph("模板固定落款")
            closing.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            closing_date = template.add_paragraph("202X年X月X日")
            closing_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            template.add_paragraph("（联系人及电话：××）")
            template.save(template_path)

            source = Document()
            source.add_paragraph("关于测试事项的请示")
            source.add_paragraph("上级单位：")
            source.add_paragraph("这是必须原样保留的引言。")
            source.add_paragraph("一、测试标题")
            source.add_paragraph("这是必须原样保留的正文。")
            source.add_paragraph("妥否，请批示。")
            source.save(source_path)

            with patch.object(template_capability, "EXPORT_DIR", root):
                output = template_capability.render_word_template_from_document(
                    template_path, source_path, title="套版验证"
                )

            rendered = Document(output)
            text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
            self.assertIn("关于测试事项的请示", text)
            self.assertIn("上级单位：", text)
            self.assertIn("这是必须原样保留的引言。", text)
            self.assertIn("一、测试标题", text)
            self.assertIn("妥否，请批示。", text)
            self.assertNotIn("模板示例正文", text)
            self.assertNotIn("模板示例标题", text)
            self.assertIn("模板固定落款", text)
            self.assertEqual(rendered.tables[0].cell(0, 0).text, "固定版头构件")
            heading = next(p for p in rendered.paragraphs if p.text == "一、测试标题")
            heading_fonts = heading.runs[0]._r.rPr.rFonts
            self.assertEqual(heading_fonts.get(qn("w:ascii")), "仿宋")

    def test_word_sample_template_uses_generated_answer_without_source_docx(self):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
        from zipfile import ZipFile

        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        db = sessionmaker(bind=engine)()
        template_row = Template(
            name="请示模板", kind="word", ext=".docx", placeholders="[]", enabled=True,
        )
        db.add(template_row)
        db.commit()

        try:
            with tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                template = Document()
                template.add_table(rows=1, cols=1).cell(0, 0).text = "固定版头构件"
                title = template.add_paragraph()
                title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title.add_run("关于××事项的请示").font.size = Pt(22)
                template.add_paragraph("××单位（主送单位）：")
                body_sample = template.add_paragraph("模板示例正文，不得残留。")
                body_sample.paragraph_format.first_line_indent = Pt(32)
                heading_sample = template.add_paragraph("一、模板示例标题")
                heading_sample.paragraph_format.space_before = Pt(9)
                template.add_paragraph("妥否，请示。")
                closing = template.add_paragraph("模板固定落款")
                closing.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                template.sections[0].footer.paragraphs[0].text = "固定模板页脚"
                template.save(root / f"{template_row.id}.docx")

                generated_answer = """# 关于申请使用CA证书的请示

公司领导：

为保障售电业务顺利开展，现申请使用CA证书。

## 一、基本情况

CA证书用于办理电力市场交易业务。

## 二、请示事项

申请同意在授权范围内使用CA证书。

妥否，请示。

售电业务经办部门
2026年8月11日

[下载通用文档](sandbox:/api/v1/exports/错误文件.docx)
"""
                with patch.object(chat_api, "TEMPLATES_DIR", root), patch.object(
                    template_capability, "EXPORT_DIR", root
                ):
                    outputs = asyncio.run(chat_api.render_chat_templates(
                        db,
                        AsyncMock(),
                        [template_row.id],
                        generated_answer,
                        "生成一个申请使用CA证书的请示",
                    ))

                self.assertEqual(len(outputs), 1)
                rendered = Document(root / outputs[0])
                text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
                self.assertIn("关于申请使用CA证书的请示", text)
                self.assertIn("公司领导：", text)
                self.assertIn("CA证书用于办理电力市场交易业务。", text)
                self.assertIn("二、请示事项", text)
                self.assertNotIn("关于××事项", text)
                self.assertNotIn("模板示例正文", text)
                self.assertNotIn("错误文件.docx", text)
                self.assertNotIn("售电业务经办部门", text)
                self.assertNotIn("2026年8月11日", text)
                self.assertIn("模板固定落款", text)
                self.assertEqual(rendered.tables[0].cell(0, 0).text, "固定版头构件")
                self.assertEqual(
                    rendered.sections[0].footer.paragraphs[0].text,
                    "固定模板页脚",
                )
                generated_heading = next(
                    paragraph for paragraph in rendered.paragraphs
                    if paragraph.text == "一、基本情况"
                )
                self.assertEqual(generated_heading.paragraph_format.space_before, Pt(9))
                with ZipFile(root / f"{template_row.id}.docx") as source_zip, ZipFile(
                    root / outputs[0]
                ) as output_zip:
                    self.assertEqual(source_zip.namelist(), output_zip.namelist())
                    for member in source_zip.namelist():
                        if member != "word/document.xml":
                            self.assertEqual(
                                source_zip.read(member),
                                output_zip.read(member),
                                member,
                            )
        finally:
            db.close()
            engine.dispose()

    def test_chat_template_render_never_writes_assistant_delivery_text_into_word(self):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        db = sessionmaker(bind=engine)()
        template_row = Template(
            name="请示模板", kind="word", ext=".docx", placeholders="[]", enabled=True,
        )
        db.add(template_row)
        db.commit()

        try:
            with tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                template = Document()
                title = template.add_paragraph()
                title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title.add_run("关于××事项的请示").font.size = Pt(22)
                template.add_paragraph("××单位（主送单位）：")
                template.add_paragraph("模板示例正文")
                suffix = template.add_paragraph("模板固定落款")
                suffix.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                template.save(root / f"{template_row.id}.docx")

                source_path = root / "原始请示.docx"
                source = Document()
                source.add_paragraph("关于真实事项的请示")
                source.add_paragraph("上级单位：")
                source.add_paragraph("真实正文。")
                source.save(source_path)

                assistant_delivery = (
                    "已经处理完成。[下载文件](sandbox:/api/v1/exports/错误文件.docx)"
                )
                with patch.object(chat_api, "TEMPLATES_DIR", root), patch.object(
                    template_capability, "EXPORT_DIR", root
                ):
                    outputs = asyncio.run(chat_api.render_chat_templates(
                        db,
                        AsyncMock(),
                        [template_row.id],
                        assistant_delivery,
                        "请套用模板",
                        source_documents=[source_path],
                    ))

                self.assertEqual(len(outputs), 1)
                text = "\n".join(
                    paragraph.text for paragraph in Document(root / outputs[0]).paragraphs
                )
                self.assertIn("关于真实事项的请示", text)
                self.assertIn("真实正文。", text)
                self.assertNotIn("已经处理完成", text)
                self.assertNotIn("错误文件.docx", text)

                answer, artifacts = chat_api._prefer_template_artifacts(
                    assistant_delivery,
                    ["通用排版中间件.docx", "检查报告.txt"],
                    outputs,
                )
                self.assertNotIn("通用排版中间件.docx", artifacts)
                self.assertIn("检查报告.txt", artifacts)
                self.assertIn(outputs[0], artifacts)
                self.assertIn(outputs[0], answer)
                self.assertNotIn("错误文件.docx", answer)
        finally:
            db.close()
            engine.dispose()

    def test_active_chat_jobs_returns_recoverable_running_topics_only(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        db = sessionmaker(bind=engine)()
        user = User(username="owner", password_hash="x", role="user")
        db.add(user)
        db.flush()
        db.add_all([
            Job(
                id="running-job", owner_id=user.id, agent_id=7, kind="chat",
                status="running", progress="正在调用工具",
                payload=json.dumps({
                    "inputs": {"query": "重构 agent"}, "session_id": "session-running",
                }, ensure_ascii=False),
            ),
            Job(
                id="done-job", owner_id=user.id, agent_id=7, kind="chat",
                status="done", payload=json.dumps({
                    "query": "已完成", "session_id": "session-done",
                }, ensure_ascii=False),
            ),
        ])
        db.commit()

        result = active_chat_jobs(user, db)

        self.assertEqual(len(result["items"]), 1)
        self.assertEqual(result["items"][0]["turn_id"], "running-job")
        self.assertEqual(result["items"][0]["session_id"], "session-running")
        self.assertEqual(result["items"][0]["query"], "重构 agent")
        db.close()
        engine.dispose()

    def test_agent_chat_statuses_aggregate_active_and_latest_terminal_state(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        db = sessionmaker(bind=engine)()
        user = User(username="light-owner", password_hash="x", role="user")
        other = User(username="other-owner", password_hash="x", role="user")
        agents = [Agent(name="thinking"), Agent(name="attention"), Agent(name="done")]
        db.add_all([user, other, *agents])
        db.flush()
        thread = Thread(id="light-thread", owner_id=user.id, agent_id=agents[2].id)
        turn = Turn(
            id="terminal-turn", thread_id=thread.id, owner_id=user.id,
            agent_id=agents[2].id, sequence=1, status="completed", input="完成",
        )
        db.add_all([
            thread,
            turn,
            Job(
                id="running-light", owner_id=user.id, agent_id=agents[0].id,
                kind="chat", status="running", payload=json.dumps({"session_id": "run-thread"}),
            ),
            Job(
                id="approval-light", owner_id=user.id, agent_id=agents[1].id,
                kind="chat", status="awaiting_approval",
                payload=json.dumps({"session_id": "approval-thread"}),
            ),
            Job(
                id="foreign-light", owner_id=other.id, agent_id=agents[0].id,
                kind="chat", status="awaiting_approval", payload="{}",
            ),
        ])
        db.commit()

        result = agent_chat_statuses(user, db)
        by_agent = {item["agent_id"]: item for item in result["items"]}

        self.assertEqual(by_agent[agents[0].id]["active_status"], "running")
        self.assertEqual(by_agent[agents[1].id]["active_status"], "awaiting_approval")
        self.assertEqual(by_agent[agents[2].id]["terminal_status"], "completed")
        self.assertEqual(by_agent[agents[2].id]["terminal_turn_id"], "terminal-turn")
        self.assertEqual(by_agent[agents[2].id]["terminal_session_id"], "light-thread")
        db.close()
        engine.dispose()

    def test_text_followup_inherits_attachment_and_selected_skill_context(self):
        # 测试全程在同一线程运行；使用 SQLite 的线程本地内存池，确保 dispose
        # 能确定关闭唯一连接，避免 StaticPool 在嵌套 Session 后延迟到解释器退出。
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        factory = sessionmaker(bind=engine)
        db = factory()
        user = User(username="attachment-owner", password_hash="x", role="user")
        agent = Agent(name="attachment-agent", active_version=1)
        db.add_all([user, agent])
        db.commit()
        with tempfile.TemporaryDirectory() as temp, patch.object(
            chat_api, "UPLOAD_DIR", Path(temp)
        ), patch.object(
            attachments, "UPLOAD_DIR", Path(temp)
        ), patch.object(
            jobs, "SessionLocal", factory
        ), patch.object(
            chat_api, "resolve_agent", return_value=agent
        ), patch.object(
            chat_api, "enforce"
        ), patch.object(
            chat_api, "_validate_invocations"
        ), patch.object(
            chat_api, "select_chat_provider", return_value=SimpleNamespace(id=1)
        ), patch.object(
            chat_api,
            "build_execution_snapshot",
            side_effect=lambda *_args, skill_ids=None, **_kwargs: {
                "skills": [{"id": value, "name": f"skill-{value}"} for value in skill_ids or []],
                "builtin_tools": [],
                "provider": {},
            },
        ):
            first = asyncio.run(chat_api.chat(
                _FormRequest([
                    ("agent_id", str(agent.id)),
                    ("query", "处理这个文件"),
                    ("skill_ids", "[15]"),
                    ("attachments", UploadFile(
                        file=BytesIO("附件正文".encode("utf-8")),
                        filename="项目资料.txt",
                    )),
                ]),
                user,
                db,
            ))
            second = asyncio.run(chat_api.chat(
                _FormRequest([
                    ("agent_id", str(agent.id)),
                    ("query", "按原规范处理，仅改格式"),
                    ("session_id", first["session_id"]),
                ]),
                user,
                db,
            ))

            db.expire_all()
            payload = json.loads(db.get(Job, second["turn_id"]).payload)
            self.assertEqual(payload["continuation_of_turn_id"], first["turn_id"])
            self.assertEqual(payload["skill_ids"], [15])
            self.assertEqual(len(payload["attachment_docs"]), 1)
            self.assertTrue(payload["attachments_inherited"])
            self.assertTrue(payload["attachment_context"][0]["inherited"])
            self.assertEqual(payload["attachment_context"][0]["name"], "项目资料.txt")
            self.assertEqual(db.query(Attachment).count(), 1)
            self.assertEqual(second["continuation_of_turn_id"], first["turn_id"])
        db.close()
        engine.dispose()

    def test_idempotency_race_returns_canonical_job_and_removes_losing_upload(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        factory = sessionmaker(bind=engine)
        db = factory()
        user = User(username="race-owner", password_hash="x", role="user")
        agent = Agent(name="race-agent", active_version=1)
        db.add_all([user, agent])
        db.flush()
        existing = Job(
            id="canonical-turn",
            owner_id=user.id,
            agent_id=agent.id,
            kind="chat",
            status=jobs.PENDING,
            idempotency_key="race-key",
            payload=json.dumps({
                "session_id": "canonical-session",
                "attachment_context": [],
            }),
        )
        db.add(existing)
        db.commit()

        request = _FormRequest([
            ("agent_id", str(agent.id)),
            ("query", "重复提交"),
            ("documents", UploadFile(
                file=BytesIO("只属于失败方的文件".encode("utf-8")),
                filename="race-input.txt",
            )),
        ])
        request.headers["Idempotency-Key"] = "race-key"
        with tempfile.TemporaryDirectory() as temp, patch.object(
            chat_api, "UPLOAD_DIR", Path(temp)
        ), patch.object(
            jobs, "UPLOAD_DIR", Path(temp)
        ), patch.object(
            jobs, "SessionLocal", factory
        ), patch.object(
            chat_api, "resolve_agent", return_value=agent
        ), patch.object(
            chat_api, "enforce"
        ), patch.object(
            chat_api, "_validate_invocations"
        ), patch.object(
            chat_api, "select_chat_provider", return_value=SimpleNamespace(id=1)
        ), patch.object(
            chat_api, "build_execution_snapshot", return_value={"provider": {}}
        ), patch.object(
            jobs, "view_by_idempotency", return_value=None
        ):
            result = asyncio.run(chat_api.chat(request, user, db))
            self.assertEqual(result["turn_id"], "canonical-turn")
            self.assertEqual(result["session_id"], "canonical-session")
            self.assertEqual(list(Path(temp).iterdir()), [])
            self.assertEqual(db.query(Job).count(), 1)

        db.close()
        engine.dispose()

    def test_empty_done_result_is_exposed_as_failure(self):
        event = json.loads(_end_line(SimpleNamespace(
            status="done",
            result={"answer": "", "turn_id": "turn-9"},
            error="",
        )))
        self.assertEqual(event["status"], "failed")
        self.assertIn("未返回可展示结果", event["error"])
        self.assertNotIn("turn_id", event)

        raw_fallback = json.loads(_end_line(SimpleNamespace(
            status="done",
            result={"answer": (
                "工具已返回结果，但模型未能生成最终答复：\n"
                "[结构化观察]\n{\"raw\":\"tool output\"}"
            )},
            error="",
        )))
        self.assertEqual(raw_fallback["status"], "failed")
        self.assertNotIn("answer", raw_fallback)

        safe = json.loads(_end_line(SimpleNamespace(
            status="done",
            result={
                "answer": "公开答复",
                "turn_id": "turn-safe",
                "reasoning": "内部推理不得公开",
                "unexpected_private_field": "不得透传",
            },
            error="",
        )))
        self.assertEqual(safe["answer"], "公开答复")
        self.assertNotIn("reasoning", safe)
        self.assertNotIn("unexpected_private_field", safe)

    def test_conversation_title_is_nonblocking_and_has_fallback(self):
        self.assertEqual(
            worker.conversation_title_fallback("  复盘今日韩国股市？！  "),
            "复盘今日韩国股市",
        )

        async def scenario():
            gate = asyncio.Event()

            async def slow_title(*_args):
                await gate.wait()

            with patch.object(worker, "_persist_conversation_title", slow_title):
                task = worker._schedule_conversation_title(
                    "thread-title", 1, "问题", "答案", "问题"
                )
                await asyncio.sleep(0)
                self.assertFalse(task.done())
                task.cancel()
                await asyncio.gather(task, return_exceptions=True)

        asyncio.run(scenario())
        source = inspect.getsource(worker._chat_handler)
        self.assertNotIn("await make_conversation_title", source)
        self.assertLess(
            source.index("db.commit()"),
            source.index("_schedule_conversation_title"),
        )

    def test_stream_event_envelope_contains_required_identity(self):
        event = json.loads(_event_line(
            "task-7",
            "assistant.status",
            {"type": "progress", "text": "正在规划"},
            {
                "event_id": "event-9",
                "timestamp": "2026-08-03T08:00:00+00:00",
                "revision": 9,
            },
        ))
        self.assertEqual(event["task_id"], "task-7")
        self.assertEqual(event["event_id"], "event-9")
        self.assertEqual(event["revision"], 9)
        self.assertEqual(event["event_type"], "assistant.status")
        self.assertIn("timestamp", event)

    def test_secret_query_parameters_are_redacted(self):
        value = "POST https://example.test/mcp?tavilyApiKey=secret-value&mode=fast"
        redacted = redact_log_value(value)
        self.assertNotIn("secret-value", redacted)
        self.assertIn("tavilyApiKey=<redacted>", redacted)


if __name__ == "__main__":
    unittest.main()
