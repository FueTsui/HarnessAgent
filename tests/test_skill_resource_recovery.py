import asyncio
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

from pptx import Presentation

from backend import artifacts
from backend.runtime import builtin_tools, run_harness
from backend.runtime.control import analyze_observation
from backend.runtime.orchestrator import _skill_prompt
from backend.runtime.skill_resources import SkillResourceIndex


class SkillResourceRecoveryTests(unittest.TestCase):
    def test_virtual_entry_uses_loaded_instructions_and_accepts_safe_aliases(self):
        index = SkillResourceIndex([{
            "name": "pptx", "instructions": "当前技能说明：遇到失败请修正。",
            "resources": [{"name": "skill.md", "content": "过期的附件说明"}],
        }])
        for name in ("SKILL", "SKILL.md", "./SKILL.md", "skill.md", "SKILL/./SKILL.md", ".\\Skill.md"):
            with self.subTest(name=name):
                result = json.loads(index.read("pptx", name))
                self.assertTrue(result["ok"])
                self.assertEqual(result["file"], "SKILL.md")
                self.assertEqual(result["content"], "当前技能说明：遇到失败请修正。")
                self.assertTrue(analyze_observation(result).ok)

    def test_missing_and_binary_resources_are_failures_with_readable_names(self):
        index = SkillResourceIndex([{
            "name": "pptx", "instructions": "已载入技能说明",
            "resources": [
                {"name": "references/guide.md", "content": "RESOURCE_BODY_SENTINEL"},
                {"name": "assets/template.pptx", "binary": True, "content": "BINARY_SENTINEL"},
                {"name": "assets/logo.png"},
            ],
        }])
        for name, code in (("missing.md", "skill_resource_not_found"),
                           ("assets/template.pptx", "skill_resource_not_text"),
                           ("assets/logo.png", "skill_resource_not_text")):
            with self.subTest(name=name):
                result = index.read("pptx", name)
                payload = json.loads(result)
                self.assertFalse(payload["ok"])
                self.assertEqual(payload["available_resources"], ["SKILL.md", "references/guide.md"])
                self.assertNotIn("SENTINEL", result)
                observation = analyze_observation(result)
                self.assertFalse(observation.ok)
                self.assertEqual(observation.error_type, "skill_resource_error")
                self.assertEqual(observation.error_code, code)
                self.assertFalse(analyze_observation(observation).ok)

    def test_paths_and_skill_names_cannot_escape_authorized_snapshot(self):
        index = SkillResourceIndex([{
            "name": "pptx", "instructions": "ENTRY_SENTINEL",
            "resources": [
                {"name": "../other/SKILL.md", "content": "UNSAFE_SENTINEL"},
                {"name": "references/Guide.md", "content": "参考说明"},
            ],
        }])
        self.assertEqual(index.readable_names("pptx"), ["SKILL.md", "references/Guide.md"])
        with patch("builtins.open", side_effect=AssertionError("资源读取不得访问磁盘")):
            for path in ("../SKILL.md", "references/../../SKILL.md", "/SKILL.md", "C:\\private\\SKILL.md",
                         "\\\\server\\share\\SKILL.md", "https://example.test/SKILL.md", "SKILL.md\x00", "", None):
                with self.subTest(path=path):
                    payload = json.loads(index.read("pptx", path))
                    self.assertEqual(payload["error"]["code"], "skill_resource_invalid_path")
                    self.assertNotIn("SENTINEL", json.dumps(payload))
            unavailable = json.loads(index.read("another_skill", "SKILL.md"))
            self.assertEqual(unavailable["error"]["code"], "skill_not_available")
            self.assertEqual(unavailable["available_skills"], ["pptx"])
            self.assertNotIn("content", unavailable)
            good = json.loads(index.read("pptx", "./references\\Guide.md"))
            self.assertTrue(good["ok"])
            self.assertFalse(json.loads(index.read("pptx", "references/guide.md"))["ok"])

    def test_instruction_only_skill_exposes_entry_in_prompt_and_tool(self):
        skills = [{"name": "simple", "instructions": "无需附加资源也能读取入口。"}]
        prompt = _skill_prompt(skills)
        self.assertIn("可按需读取文本资源：SKILL.md", prompt)
        self.assertIn("先调用 read_skill_resource", prompt)
        self.assertNotIn(skills[0]["instructions"], prompt)
        test = self

        class Model:
            context_tokens = 32000
            calls = 0

            async def chat_with_tools(self, messages, tools=None, **kwargs):
                self.calls += 1
                if self.calls == 1:
                    spec = next(item["function"] for item in tools if item["function"]["name"] == "read_skill_resource")
                    test.assertIn("SKILL.md", spec["parameters"]["properties"]["file"]["description"])
                    test.assertIn("simple", spec["parameters"]["properties"]["skill"]["description"])
                    return _tool_call("read_skill_resource", {"skill": "simple", "file": "SKILL.md"}, "entry")
                test.assertIn(skills[0]["instructions"], messages[-1]["content"])
                return {"role": "assistant", "content": "已读取技能入口。", "tool_calls": []}

        answer, _ = asyncio.run(run_harness(Model(), "", "读取技能", skills=skills))
        self.assertEqual(answer, "已读取技能入口。")

    def test_historical_missing_then_skill_aliases_reach_valid_native_pptx(self):
        # 只使用合成材料重放缺失资源 -> SKILL -> SKILL.md 的实际失败顺序。
        # 最后的 PPTX 由真实能力生成并重新打开，不能仅凭模型文本判定恢复成功。
        test = self
        events = []
        paths = ["missing.md", "SKILL", "SKILL.md"]

        class Model:
            context_tokens = 32000
            calls = 0

            async def chat_with_tools(self, messages, tools=None, **kwargs):
                self.calls += 1
                if self.calls in {1, 6}:
                    return _tool_call("update_plan", {"plan": [
                        {"step": "核对已载入材料", "status": "completed"},
                        {"step": "生成中文演示稿", "status": "completed" if self.calls == 6 else "in_progress"},
                    ]}, f"plan{self.calls}")
                if self.calls in {2, 3, 4}:
                    if self.calls == 3:
                        test.assertIn("状态：失败", messages[-1]["content"])
                        test.assertIn("SKILL.md", messages[-1]["content"])
                    return _tool_call("read_skill_resource", {"skill": "pptx", "file": paths[self.calls - 2]}, f"read{self.calls}")
                if self.calls == 5:
                    test.assertIn("状态：成功", messages[-1]["content"])
                    test.assertIn("presentation_create", {item["function"]["name"] for item in tools})
                    return _tool_call("presentation_create", {"slides": [
                        {"title": "技能读取恢复验证", "body": "读取入口说明\n生成并验证可编辑演示稿"},
                    ], "confirm": True}, "artifact")
                return {"role": "assistant", "content": "已生成中文可编辑演示稿。", "tool_calls": []}

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            context = builtin_tools.BuiltinToolContext(
                root=root, enabled_tools={"presentation_create"}, approval_policy="auto",
            )
            with patch.object(builtin_tools, "EXPORT_DIR", root), patch.object(artifacts, "EXPORT_DIR", root), \
                 patch.object(builtin_tools, "enforce_content", new=AsyncMock()), \
                 patch("backend.runtime.orchestrator.guard_model_client", side_effect=lambda model, **kwargs: model):
                asyncio.run(run_harness(
                    Model(), "", "重试上面任务",
                    history=[{"role": "user", "content": "使用技能中文重新创建这个ppt"}],
                    attachment_text="合成材料：读取入口说明；生成并验证可编辑演示稿。",
                    skills=[{"name": "pptx", "instructions": "根据已载入材料调用 presentation_create 创建原生文本 PPTX。"}],
                    builtin_context=context,
                    runtime_event=lambda name, payload: events.append((name, payload)),
                    tool_policy={"profile": "small_model", "max_iterations": 8},
                ))
                completed = [payload for name, payload in events if name == "tool.completed"]
                self.assertEqual([item["ok"] for item in completed], [False, True, True, True])
                self.assertEqual(completed[0]["error_code"], "skill_resource_not_found")
                missing_iteration = next(payload for name, payload in events
                                         if name == "loop.iteration.completed" and payload["iteration"] == 2)
                self.assertEqual(missing_iteration["successful_tools"], 0)
                self.assertFalse(any(payload.get("reason") == "repeated_tool_call" for _, payload in events))
                self.assertEqual(len(context.artifacts), 1)
                self.assertEqual(builtin_tools.completion_artifact_issues(context), [])
                presentation = Presentation(str(root / context.artifacts[0]))
                self.assertEqual(len(presentation.slides), 1)
                self.assertIn("生成并验证可编辑演示稿", presentation.slides[0].shapes[1].text)


def _tool_call(name, arguments, call_id):
    return {"role": "assistant", "content": "", "tool_calls": [{
        "id": call_id, "type": "function",
        "function": {"name": name, "arguments": json.dumps(arguments, ensure_ascii=False)},
    }]}


if __name__ == "__main__":
    unittest.main()
