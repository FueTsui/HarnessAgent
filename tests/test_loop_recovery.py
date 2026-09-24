"""End-to-end deterministic model replays for the failed Harness tool loops."""
import asyncio
import copy
import json
from contextlib import ExitStack
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

from pptx import Presentation

from backend import artifacts
from backend.runtime import builtin_tools, run_harness
from backend.runtime.orchestrator import CompletionVerificationError


def tool_call(name, arguments):
    return {"role": "assistant", "content": "", "tool_calls": [{
        "id": "test_call", "type": "function", "function": {
            "name": name,
            "arguments": json.dumps(arguments, ensure_ascii=False) if isinstance(arguments, dict) else arguments,
        },
    }]}


def plan(status="in_progress"):
    return {"plan": [
        {"step": "确认已载入材料", "status": "completed"},
        {"step": "完成用户交付目标", "status": status},
    ]}


FINAL = {"role": "assistant", "content": "已完成用户要求的全部工作。", "tool_calls": []}
PPT = {"slides": [{"title": "团队工作流程", "body": "准备上下文\n验证交付物"}], "confirm": True}


class ScriptedModel:
    context_tokens = 32_000

    def __init__(self, responses):
        self.responses = responses
        self.calls = 0
        self.requests = []

    async def chat_with_tools(self, messages, tools=None, **kwargs):
        self.requests.append(copy.deepcopy({"messages": messages, "tools": tools}))
        index = self.calls
        self.calls += 1
        response = copy.deepcopy(self.responses[index] if index < len(self.responses) else FINAL)
        for call in response.get("tool_calls") or []:
            call["id"] = f"call_{self.calls}"
        return response


class LoopRecoveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.stack.enter_context(patch.object(builtin_tools, "EXPORT_DIR", self.root))
        self.stack.enter_context(patch.object(artifacts, "EXPORT_DIR", self.root))
        self.stack.enter_context(patch.object(builtin_tools, "enforce_content", new=AsyncMock()))
        self.stack.enter_context(patch("backend.runtime.orchestrator.enforce_content", new=AsyncMock()))
        self.stack.enter_context(patch("backend.runtime.orchestrator.guard_model_client", side_effect=lambda model, **kwargs: model))
        self.execution = self.stack.enter_context(patch.object(builtin_tools, "execute", wraps=builtin_tools.execute))
        self.ctx = builtin_tools.BuiltinToolContext(
            root=self.root, enabled_tools={"presentation_create"}, approval_policy="auto",
        )
        self.events = []
        self.metadata = {}

    def run_model(self, responses, query="根据已载入材料制作 PPT", **overrides):
        self.model = ScriptedModel(responses)
        options = {
            "skills": [{"name": "pptx", "description": "演示文稿", "instructions": "使用已载入正文调用 presentation_create。"}],
            "attachment_text": "团队工作流程：准备上下文；验证交付物。",
            "builtin_context": self.ctx,
            "runtime_event": lambda event, data: self.events.append((event, data)),
            "completion_metadata": self.metadata,
            "tool_policy": {"max_iterations": 8, "max_successful_calls": 8, "router": {"enabled": False}},
            "verification_policy": {"max_revisions": 0},
        }
        options.update(overrides)
        return asyncio.run(run_harness(self.model, "", query, **options))

    def event_data(self, name):
        return [data for event, data in self.events if event == name]

    def assert_valid_ppt(self):
        self.assertEqual(len(self.ctx.artifacts), 1)
        filename = self.ctx.artifacts[0]
        self.assertTrue(artifacts.inspect_presentation_artifact(filename)["valid"])
        presentation = Presentation(str(self.root / filename))
        self.assertEqual(len(presentation.slides), 1)
        self.assertEqual(presentation.slides[0].shapes[1].text, PPT["slides"][0]["body"])
        self.assertEqual(self.metadata["completion_status"], "completed")

    def test_real_parameter_bearing_name_reads_skill_and_produces_valid_ppt(self):
        self.ctx.required_artifact_kinds.add("presentation")
        self.run_model([
            tool_call("update_plan", plan()),
            tool_call('read_skill_resource(file="SKILL", skill="pptx")\n</parameter', ""),
            tool_call("presentation_create", PPT),
            tool_call("update_plan", plan("completed")), FINAL,
        ])
        self.assert_valid_ppt()
        read_calls = [event for event in self.event_data("tool.called") if event["tool"] == "read_skill_resource"]
        self.assertEqual(len(read_calls), 1)
        self.assertIn("工具调用格式已规范化", read_calls[0]["repairs"])
        self.assertTrue(any(event["tool"] == "read_skill_resource" and event["ok"] for event in self.event_data("tool.completed")))
        self.assertFalse(self.event_data("loop.recovery"))

    def test_one_cached_plan_can_be_followed_by_changed_completed_state(self):
        self.run_model([
            tool_call("update_plan", plan()), tool_call("update_plan", plan()),
            tool_call("update_plan", plan("completed")), FINAL,
        ], query="核对输入并给出总结")
        self.assertEqual(self.ctx.plan_revision, 2)
        self.assertTrue(all(item["status"] == "completed" for item in self.ctx.plan_steps))
        self.assertEqual(self.metadata["completion_status"], "completed")
        self.assertEqual(len(self.event_data("loop.recovery")), 1)
        self.assertFalse(self.event_data("loop.stopped"))

    def test_repeated_missing_resource_can_switch_to_presentation_tool(self):
        self.ctx.required_artifact_kinds.add("presentation")
        missing = tool_call("read_skill_resource", {"skill": "pptx", "file": "missing.txt"})
        self.run_model([
            tool_call("update_plan", plan()), missing, missing,
            tool_call("presentation_create", PPT), tool_call("update_plan", plan("completed")), FINAL,
        ])
        self.assert_valid_ppt()
        resources = [data for data in self.event_data("tool.completed") if data["tool"] == "read_skill_resource"]
        self.assertEqual([data["ok"] for data in resources], [False, False])
        self.assertEqual([data["repeated"] for data in resources], [False, True])
        self.assertEqual(len(self.event_data("loop.recovery")), 1)
        self.assertFalse(self.event_data("loop.stopped"))

    def test_cached_mutating_presentation_is_not_executed_twice(self):
        self.run_model([
            tool_call("update_plan", plan()), tool_call("presentation_create", PPT),
            tool_call("presentation_create", PPT), tool_call("update_plan", plan("completed")), FINAL,
        ])
        self.assert_valid_ppt()
        self.assertEqual(self.execution.call_count, 1)
        self.assertEqual(len(list(self.root.glob("*.pptx"))), 1)
        self.assertEqual(len(self.event_data("loop.recovery")), 1)

    def test_persistent_repetition_is_bounded_and_missing_ppt_cannot_be_claimed_complete(self):
        self.ctx.required_artifact_kinds.add("presentation")
        missing = tool_call("read_skill_resource", {"skill": "pptx", "file": "missing.txt"})
        with self.assertRaises(CompletionVerificationError):
            self.run_model([
                tool_call("update_plan", plan()), missing, missing, missing,
                FINAL, tool_call("update_plan", plan("completed")), FINAL,
            ])
        self.assertEqual(len(self.event_data("loop.recovery")), 1)
        self.assertEqual([event["reason"] for event in self.event_data("loop.stopped")], ["repeated_tool_call"])
        self.assertLessEqual(self.model.calls, 7)
        self.assertFalse(self.ctx.artifacts)
        self.assertFalse(self.event_data("loop.completed"))
        self.assertTrue(self.event_data("loop.blocked"))

    def test_persistent_plan_repetition_does_not_fabricate_completion(self):
        with self.assertRaises(CompletionVerificationError):
            self.run_model([
                tool_call("update_plan", plan()), tool_call("update_plan", plan()),
                tool_call("update_plan", plan()), FINAL, FINAL, FINAL,
            ], query="核对输入并给出总结")
        self.assertEqual(self.ctx.plan_revision, 1)
        self.assertEqual(len(self.event_data("loop.recovery")), 1)
        self.assertFalse(self.event_data("loop.completed"))
        self.assertLessEqual(self.model.calls, 6)

    def test_protocol_rejection_diagnostics_do_not_publish_raw_parameters(self):
        secret = "SENSITIVE_PROTOCOL_ARGUMENT"
        self.run_model([
            tool_call(f'read_skill_resource(file=unsafe("{secret}"), skill="pptx")', ""),
            FINAL,
        ], query="检查工具调用格式")
        self.assertEqual(self.execution.call_count, 0)
        rejected = self.event_data("tool.rejected")
        self.assertEqual(len(rejected), 1)
        self.assertEqual(rejected[0]["reason"], "invalid_tool_call")
        self.assertNotIn(secret, json.dumps(self.events, ensure_ascii=False))
        self.assertNotIn("unsafe", json.dumps(self.events, ensure_ascii=False))
        tool_feedback = [m for m in self.model.requests[-1]["messages"] if m["role"] == "tool"]
        self.assertTrue(tool_feedback)
        self.assertNotIn(secret, json.dumps(tool_feedback))

    def test_budget_closeout_recovers_parameter_bearing_update_plan(self):
        completed = repr(plan("completed")["plan"])
        self.run_model([
            tool_call("update_plan", plan()), FINAL,
            tool_call(f"update_plan(plan={completed})\n</parameter", ""), FINAL,
        ], query="核对输入并给出总结", tool_policy={"max_iterations": 1, "router": {"enabled": False}})
        self.assertEqual(self.metadata["completion_status"], "completed")
        self.assertEqual(self.ctx.plan_revision, 2)
        self.assertTrue(self.event_data("plan.closeout.completed")[0]["applied"])

    def test_nondict_function_is_rejected_without_crashing_regular_loop(self):
        for raw_function in ("malformed", ["malformed"]):
            with self.subTest(raw_function=raw_function):
                self.events.clear()
                response = tool_call("unused", {})
                response["tool_calls"][0]["function"] = raw_function
                self.run_model([response, FINAL], query="检查工具调用格式")
                self.assertEqual(self.event_data("tool.rejected")[0]["reason"], "invalid_tool_call")
                self.assertEqual(self.execution.call_count, 0)

    def test_nondict_function_during_closeout_keeps_plan_unresolved(self):
        for raw_function in ("malformed", ["malformed"]):
            with self.subTest(raw_function=raw_function):
                self.events.clear()
                self.ctx.plan_steps = []
                self.ctx.plan_revision = 0
                response = tool_call("unused", {})
                response["tool_calls"][0]["function"] = raw_function
                with self.assertRaises(CompletionVerificationError):
                    self.run_model([
                        tool_call("update_plan", plan()), FINAL, response, FINAL,
                    ], query="核对输入并给出总结", tool_policy={"max_iterations": 1})
                self.assertFalse(self.event_data("plan.closeout.completed")[0]["applied"])
                self.assertFalse(self.event_data("loop.completed"))


if __name__ == "__main__":
    unittest.main()
