import asyncio
import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from starlette.datastructures import UploadFile

from backend import artifacts
from backend.api import skills as skills_api
from backend.api import templates as templates_api
from backend.database import Base
from backend.models import Skill, Template, User
from backend.runtime.control import requires_presentation_artifact


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "QA title"
    slide.placeholders[1].text = "QA subtitle"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _artifact_skill_zip() -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "artifact-template/SKILL.md",
            "---\nname: audit-ppt\ndescription: PPT template\n---\n\nCreate a PPT.",
        )
        archive.writestr(
            "artifact-template/artifact-template.json",
            json.dumps({
                "schemaVersion": 1,
                "kind": "presentation",
                "reference": "assets/reference.pptx",
                "preview": "assets/preview.png",
            }),
        )
        archive.writestr(
            "artifact-template/assets/reference.pptx",
            _pptx_bytes(),
        )
        archive.writestr(
            "artifact-template/assets/preview.png",
            b"\x89PNG\r\n\x1a\n\x00binary-preview",
        )
        archive.writestr(
            "artifact-template/references/notes.md",
            "# Notes\nKeep the package path.",
        )
    return output.getvalue()


class SkillArtifactTemplateTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine)()
        self.admin = User(username="root", password_hash="x", role="root")
        self.db.add(self.admin)
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def test_zip_import_preserves_binary_assets_and_registers_ppt_template(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            skill_root = root / "skills"
            template_root = root / "templates"
            skill_root.mkdir()
            template_root.mkdir()
            package_bytes = _artifact_skill_zip()
            upload = UploadFile(
                filename="audit-ppt.zip",
                file=io.BytesIO(package_bytes),
            )
            with patch.object(skills_api, "SKILLS_DIR", skill_root), patch.object(
                skills_api, "TEMPLATES_DIR", template_root
            ), patch.object(templates_api, "TEMPLATES_DIR", template_root):
                result = asyncio.run(skills_api.import_skills(upload, self.admin, self.db))

                self.assertEqual(result["imported"], 1)
                skill = self.db.query(Skill).one()
                resources = json.loads(skill.resources)
                by_path = {item["package_path"]: item for item in resources}
                self.assertIn("assets/reference.pptx", by_path)
                self.assertTrue(by_path["assets/reference.pptx"]["binary"])
                self.assertNotIn("content", by_path["assets/reference.pptx"])
                with zipfile.ZipFile(io.BytesIO(package_bytes)) as package:
                    original_pptx = package.read(
                        "artifact-template/assets/reference.pptx"
                    )
                self.assertEqual(
                    (skill_root / str(skill.id) / "assets/reference.pptx").read_bytes(),
                    original_pptx,
                )
                self.assertEqual(
                    by_path["references/notes.md"]["content"],
                    "# Notes\nKeep the package path.",
                )

                template = self.db.query(Template).one()
                manifest = by_path["artifact-template.json"]
                self.assertEqual(manifest["template_id"], template.id)
                self.assertTrue(template.enabled)
                self.assertEqual(template.kind, "ppt")
                self.assertTrue((template_root / f"{template.id}.pptx").is_file())
                self.assertEqual(
                    skills_api.artifact_template_ids(self.db, [skill.id]),
                    [template.id],
                )

    def test_zip_import_rejects_package_path_traversal(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("skill/SKILL.md", "---\nname: unsafe\n---\n")
            archive.writestr("skill/../escape.pptx", b"bad")
        with self.assertRaisesRegex(ValueError, "非法资源路径"):
            skills_api._read_zip_skill(output.getvalue())

    def test_presentation_artifact_requires_a_real_pptx_package(self):
        with tempfile.TemporaryDirectory() as temp, patch.object(
            artifacts, "EXPORT_DIR", Path(temp)
        ):
            (Path(temp) / "valid.pptx").write_bytes(_pptx_bytes())
            (Path(temp) / "fake.pptx").write_text("not a ppt", encoding="utf-8")
            self.assertTrue(artifacts.valid_presentation_artifact("valid.pptx"))
            self.assertFalse(artifacts.valid_presentation_artifact("fake.pptx"))
            self.assertFalse(artifacts.valid_presentation_artifact("missing.pptx"))
            report = artifacts.inspect_presentation_artifact("valid.pptx")
            self.assertTrue(report["valid"])
            self.assertEqual(report["slide_count"], 1)

    def test_presentation_quality_gate_rejects_visible_markdown(self):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as temp, patch.object(
            artifacts, "EXPORT_DIR", Path(temp)
        ):
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            slide.shapes.title.text = "**Unclean title"
            slide.placeholders[1].text = "Body"
            presentation.save(Path(temp) / "markup.pptx")

            report = artifacts.inspect_presentation_artifact("markup.pptx")

            self.assertFalse(report["valid"])
            self.assertIn("第1页：残留 Markdown 标记", report["issues"])

    def test_presentation_create_intent_is_detected(self):
        self.assertTrue(requires_presentation_artifact("使用 skill 生成PPT"))
        self.assertTrue(requires_presentation_artifact("请制作一份 PowerPoint 演示文稿"))
        self.assertFalse(requires_presentation_artifact("解释一下什么是 PPT"))


if __name__ == "__main__":
    unittest.main()
