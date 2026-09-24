import asyncio
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from pptx import Presentation

from backend import artifacts
from backend.capabilities.presentations import create_presentation
from backend.runtime import builtin_tools
from backend.runtime.control import resolve_task_objective
from backend.runtime.presentation_requirements import (
    build_presentation_requirements, inspect_presentation_requirements,
)


def _pdf(path: Path, count: int):
    writer = PdfWriter()
    font = writer._add_object(DictionaryObject({
        NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    }))
    for number in range(1, count + 1):
        page = writer.add_blank_page(width=720, height=405)
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 18 Tf 40 340 Td (Source page {number}: clarify objectives and verify results.) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(content)
    with path.open("wb") as output:
        writer.write(output)


def _slides(count: int, *, chinese: bool = True):
    return [{
        "title": f"第 {i} 页工作流程" if chinese else f"Workflow {i}",
        "body": "明确本页目标与验证方法\n依据原材料重制正文" if chinese else "Clarify objectives and verify the results.",
        "source_pages": [i],
    } for i in range(1, count + 1)]


class PresentationCoverageTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source.pdf"
        _pdf(self.source, 44)
        self.requirements = build_presentation_requirements("使用技能中文重新创建这个ppt", [self.source])

    def test_pdf_pages_define_contract_and_context_even_for_retry(self):
        objective = resolve_task_objective("重试上面任务", [{"role": "user", "content": "使用技能中文重新创建这个ppt"}])
        requirements = build_presentation_requirements(objective, [self.source])
        self.assertEqual(requirements.page_count, 44)
        self.assertTrue(requirements.require_chinese)
        context = requirements.context()
        for number in (1, 21, 44):
            self.assertIn(f"[来源页 {number}/44", context)
            self.assertIn(f"Source page {number}:", context)
        self.assertIn("不能只做封面或摘要", context)

    def test_summary_without_recreation_and_non_presentation_goals_stay_unrestricted(self):
        for query in ("根据附件总结一份中文PPT", "只做一份中文摘要PPT", "重新创建中文Word", "重新创建摘要版PPT"):
            with self.subTest(query=query):
                self.assertIsNone(build_presentation_requirements(query, [self.source]))
        self.assertIsNone(build_presentation_requirements("中文重新创建这个ppt", [self.root / "note.txt"]))
        output = create_presentation({"slides": _slides(1, chinese=False)}, self.root)
        self.assertEqual(output["slide_count"], 1)

    def test_source_mapping_and_language_cannot_be_weakened_by_tool_arguments(self):
        args = {"slides": _slides(1, chinese=False), "source_page_count": 1,
                "require_chinese": False, "requirements": {"page_count": 1}}
        with self.assertRaisesRegex(ValueError, "PDF 共 44 页"):
            create_presentation(args, self.root, requirements=self.requirements)
        args["slides"][0]["source_pages"] = list(range(1, 45))
        with self.assertRaisesRegex(ValueError, "不得合并"):
            create_presentation(args, self.root, requirements=self.requirements)
        args["slides"] = _slides(44, chinese=False)
        with self.assertRaisesRegex(ValueError, "中文重制"):
            create_presentation(args, self.root, requirements=self.requirements)
        args["slides"] = _slides(44)
        args["slides"][43]["source_pages"] = [1]
        with self.assertRaisesRegex(ValueError, "第 44 页 source_pages"):
            create_presentation(args, self.root, requirements=self.requirements)
        self.assertEqual(list(self.root.glob("*.pptx")), [])

    def test_all_source_pages_generate_native_chinese_slides_with_provenance(self):
        result = create_presentation({"slides": _slides(44)}, self.root, requirements=self.requirements)
        report = inspect_presentation_requirements(result["file"], self.requirements, self.root)
        self.assertTrue(report["valid"], report)
        self.assertEqual(report["source_page_count"], 44)
        presentation = Presentation(str(self.root / result["file"]))
        self.assertEqual(len(presentation.slides), 44)
        self.assertIn("第 44 页工作流程", presentation.slides[43].shapes[0].text)
        self.assertIn("来源页 44/44", presentation.slides[43].notes_slide.notes_text_frame.text)

    def test_runtime_handler_uses_server_contract_and_final_gate_reopens_artifact(self):
        context = builtin_tools.BuiltinToolContext(
            root=self.root, enabled_tools={"presentation_create"}, approval_policy="auto",
            presentation_requirements=self.requirements,
        )
        with patch.object(builtin_tools, "EXPORT_DIR", self.root), patch.object(artifacts, "EXPORT_DIR", self.root), \
             patch.object(builtin_tools, "enforce_content", new=AsyncMock()):
            result = json.loads(asyncio.run(builtin_tools.execute("presentation_create", {
                "slides": _slides(1, chinese=False), "confirm": True,
            }, context)))
            self.assertFalse(result["ok"])
            self.assertEqual(context.artifacts, [])
            for count in (1, 44):
                with self.subTest(count=count):
                    exported = create_presentation({"slides": _slides(count, chinese=False)}, self.root)
                    self.assertTrue(artifacts.inspect_presentation_artifact(exported["file"])["valid"])
                    context.artifacts = [exported["file"]]
                    context.required_artifact_kinds.add("presentation")
                    context.artifact_metadata[exported["file"]] = {"source_page_count": 44, "require_chinese": True, "valid": True}
                    issues = builtin_tools.completion_artifact_issues(context)
                    self.assertTrue(issues)
                    self.assertTrue(any("中文重制" in issue for issue in issues))

    def test_source_parse_error_remains_a_hard_requirement(self):
        broken = self.root / "broken.pdf"
        broken.write_bytes(b"invalid")
        requirements = build_presentation_requirements("中文重新创建这个ppt", [broken])
        self.assertIsNotNone(requirements)
        self.assertTrue(requirements.source_errors)
        with self.assertRaisesRegex(ValueError, "来源.*无法读取"):
            create_presentation({"slides": _slides(1)}, self.root, requirements=requirements)

    def test_chat_postprocessing_rejects_structurally_valid_english_cover(self):
        from sqlalchemy import create_engine
        from sqlalchemy.orm import Session
        from backend.api import chat
        from backend.database import Base
        from backend.models import Agent, User
        from backend.runtime.contracts import TaskInput

        engine = create_engine("sqlite://")
        Base.metadata.create_all(engine)
        self.addCleanup(engine.dispose)
        snapshot = {"provider": {}, "agent": {"memory_enabled": False}, "harness": {},
                    "skills": [], "mcp_servers": [], "sub_agents": [], "builtin_tools": []}
        test = self

        async def rogue_cover(*args, **kwargs):
            context = kwargs["builtin_context"]
            test.assertEqual(context.presentation_requirements.page_count, 44)
            test.assertIn("[来源页 44/44", kwargs["attachment_text"])
            cover = create_presentation({"slides": _slides(1, chinese=False)}, self.root)
            context.artifacts.append(cover["file"])
            return "已完成整份中文重制。", ""

        with Session(engine) as db:
            db.add_all([User(id=1, username="test", password_hash="x", role="root"),
                        Agent(id=1, name="test", created_by=1, is_public=True)])
            db.commit()
            with patch.object(chat, "_client_from_execution_snapshot", return_value=object()), \
                 patch("backend.guardrail_policies.guard_model_client", side_effect=lambda model, **kwargs: model), \
                 patch.object(builtin_tools, "workspace_for_run", return_value=self.root / "workspace"), \
                 patch.object(builtin_tools, "globally_enabled_names", return_value={"presentation_create"}), \
                 patch.object(builtin_tools, "EXPORT_DIR", self.root), patch.object(artifacts, "EXPORT_DIR", self.root), \
                 patch.object(chat, "run_harness", side_effect=rogue_cover), \
                 patch.object(chat, "render_chat_templates", new=AsyncMock(return_value=[])):
                with self.assertRaisesRegex(chat.CompletionVerificationError, "来源覆盖不完整"):
                    asyncio.run(chat.execute_chat(db, db.get(Agent, 1),
                        TaskInput(query="使用技能中文重新创建这个ppt", documents=[self.source]),
                        user_id=1, execution_snapshot=snapshot, attachment_docs=[self.source]))


if __name__ == "__main__":
    unittest.main()
