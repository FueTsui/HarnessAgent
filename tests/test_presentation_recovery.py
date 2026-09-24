import asyncio
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

from pptx import Presentation
from backend import artifacts
from backend.capabilities.presentations import create_presentation
from backend.runtime import builtin_tools, run_harness
from backend.runtime.control import resolve_task_objective, requires_presentation_artifact, route_tools


class PresentationRecoveryTests(unittest.TestCase):
    def test_recreation_intent_requires_artifact_but_page_review_does_not(self):
        for query in ("中文逐页重制这个ppt", "把PPT重建为中文", "重做演示文稿", "rebuild these slides"):
            with self.subTest(query=query):
                self.assertTrue(requires_presentation_artifact(query))
        self.assertFalse(requires_presentation_artifact("逐页检查这个PPT并指出错误"))

    def test_retry_inherits_latest_user_goal_without_reviving_cancelled_goal(self):
        history = [{"role": "user", "content": "使用技能中文重新创建这个ppt"},
                   {"role": "assistant", "content": "无法创建"},
                   {"role": "user", "content": "重试上面任务"}]
        goal = resolve_task_objective("重试上面任务", history)
        self.assertTrue(requires_presentation_artifact(goal))
        self.assertEqual(resolve_task_objective("只总结正文", history), "只总结正文")
        history.append({"role": "user", "content": "只总结正文"})
        self.assertFalse(requires_presentation_artifact(resolve_task_objective("继续", history)))

    def test_router_retains_authorized_presentation_capability(self):
        tools = builtin_tools.tool_specs(enabled_names=builtin_tools.PUBLIC_AGENT_SAFE_TOOLS)
        goal = resolve_task_objective("重试上面任务", [{"role": "user", "content": "使用技能中文重新创建这个ppt"}])
        selected = route_tools(tools, goal, threshold=8, limit=7)
        self.assertIn("presentation_create", [x["function"]["name"] for x in selected])
        self.assertLessEqual(len(selected), 7)
        self.assertNotIn("shell", [x["function"]["name"] for x in selected])
        restricted = route_tools([x for x in tools if x["function"]["name"] != "presentation_create"], goal, threshold=8, limit=7)
        self.assertNotIn("presentation_create", [x["function"]["name"] for x in restricted])

    def test_export_preserves_native_text_and_rejects_overflow(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            args = {"slides": [{"title": "团队工作流程", "body": "明确目标与验收条件\n保留来源，验证结果"}], "output_name": "../../unsafe"}
            result = create_presentation(args, root)
            prs = Presentation(str(root / result["file"]))
            self.assertEqual(len(prs.slides), 1)
            self.assertEqual(prs.slides[0].shapes[1].text, args["slides"][0]["body"])
            with patch.object(artifacts, "EXPORT_DIR", root):
                self.assertTrue(artifacts.inspect_presentation_artifact(result["file"])["valid"])
            for slides in ([], [{"title": "超长", "body": "正文" * 241}], [{"title": "行数", "body": "行\n" * 12}]):
                with self.assertRaises(ValueError):
                    create_presentation({"slides": slides}, root)
            self.assertEqual(len(list(root.iterdir())), 1)

    def test_harness_retry_offers_tool_and_generates_verified_artifact(self):
        class Model:
            context_tokens = 32000
            calls = 0

            async def chat_with_tools(self, messages, tools=None, **kwargs):
                self.calls += 1
                names = {t["function"]["name"] for t in tools or []}
                if self.calls in {1, 3}:
                    return {"role": "assistant", "content": "", "tool_calls": [{"id": f"plan{self.calls}", "type": "function", "function": {
                        "name": "update_plan", "arguments": json.dumps({"plan": [
                            {"step": "读取已载入正文", "status": "completed"},
                            {"step": "生成中文演示稿", "status": "in_progress" if self.calls == 1 else "completed"}]})}}]}
                if self.calls == 2:
                    if "presentation_create" not in names:
                        raise AssertionError(names)
                    return {"role": "assistant", "content": "", "tool_calls": [{"id": "ppt1", "type": "function", "function": {
                        "name": "presentation_create", "arguments": json.dumps({"slides": [{"title": "团队工作流程", "body": "准备上下文\n验证交付物"}], "confirm": True})}}]}
                return {"role": "assistant", "content": "已生成中文可编辑演示稿。", "tool_calls": []}

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            ctx = builtin_tools.BuiltinToolContext(root=root, enabled_tools=set(builtin_tools.PUBLIC_AGENT_SAFE_TOOLS), approval_policy="auto")
            with patch.object(builtin_tools, "EXPORT_DIR", root), patch.object(artifacts, "EXPORT_DIR", root), patch.object(builtin_tools, "enforce_content", new=AsyncMock()), patch("backend.runtime.orchestrator.guard_model_client", side_effect=lambda model, **kwargs: model):
                asyncio.run(run_harness(Model(), "", "重试上面任务", history=[{"role": "user", "content": "使用技能中文重新创建这个ppt"}], attachment_text="团队工作流程：准备上下文；验证交付物", builtin_context=ctx,
                    tool_policy={"profile": "small_model", "router": {"enabled": True, "activation_threshold": 8, "max_candidates": 7}}))
                self.assertEqual(len(ctx.artifacts), 1)
                self.assertEqual(builtin_tools.completion_artifact_issues(ctx), [])
                (root / ctx.artifacts[0]).unlink()
                self.assertTrue(builtin_tools.completion_artifact_issues(ctx))

    def test_execution_cannot_bypass_tool_assignment(self):
        ctx = builtin_tools.BuiltinToolContext(enabled_tools=set())
        result = json.loads(asyncio.run(builtin_tools.execute("presentation_create", {"confirm": True}, ctx)))
        self.assertFalse(result["ok"])
        self.assertIn("未授权", result["error"])

    def test_chat_retry_requires_real_ppt_and_respects_global_disable(self):
        from sqlalchemy import create_engine
        from sqlalchemy.orm import Session
        from backend.database import Base
        from backend.models import Agent, User
        from backend.api import chat
        from backend.runtime.contracts import TaskInput
        engine = create_engine("sqlite://")
        Base.metadata.create_all(engine)
        self.addCleanup(engine.dispose)
        snapshot = {"provider": {}, "agent": {"memory_enabled": False}, "harness": {},
                    "skills": [], "mcp_servers": [], "sub_agents": [], "builtin_tools": []}
        async def no_artifact(*args, **kwargs):
            ctx = kwargs["builtin_context"]
            self.assertIn("presentation", ctx.deferred_artifact_kinds)
            self.assertNotIn("presentation_create", ctx.enabled_tools)
            return "无法生成 PPT", ""
        with tempfile.TemporaryDirectory() as temp, Session(engine) as db:
            user = User(id=1, username="test", password_hash="x", role="root")
            agent = Agent(id=1, name="test", created_by=1, is_public=True)
            db.add_all([user, agent])
            db.commit()
            with patch.object(chat, "_client_from_execution_snapshot", return_value=object()), \
                 patch("backend.guardrail_policies.guard_model_client", side_effect=lambda model, **kwargs: model), \
                 patch.object(builtin_tools, "workspace_for_run", return_value=Path(temp)), \
                 patch.object(builtin_tools, "globally_enabled_names", return_value=set()), \
                 patch.object(chat, "run_harness", side_effect=no_artifact), \
                 patch.object(chat, "render_chat_templates", new=AsyncMock(return_value=[])):
                with self.assertRaises(chat.CompletionVerificationError):
                    asyncio.run(chat.execute_chat(db, agent, TaskInput(query="重试上面任务"), user_id=1,
                        execution_snapshot=snapshot,
                        history=[{"role": "user", "content": "使用技能中文重新创建这个ppt"}]))


if __name__ == "__main__":
    unittest.main()
