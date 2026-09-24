"""Harness 内置工具目录。

与 MCP 工具不同，这些工具由本地 Runtime 实现，覆盖工作区文件、受限命令、
Git、代码诊断、HTML/图片产物、轻量浏览器、Cron 和异步子智能体控制。

安全不变量：
- 文件路径必须位于 ``WORKSPACE_DIR``；
- Shell 只接受单条白名单命令，拒绝重定向/管道/命令替换和解释器内联代码；
- 写入、提交、生成计划等有副作用操作要求 ``confirm=true``；
- 每个工具返回稳定 JSON，错误不会泄露环境变量或密钥。
"""
from __future__ import annotations

import asyncio
import base64
import datetime as dt
import fnmatch
import html
import json
import mimetypes
import os
import re
import shlex
import shutil
import subprocess
import uuid
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Awaitable, Callable
from urllib.parse import parse_qs, quote_plus, unquote, urlencode, urljoin, urlparse, urlunparse

import httpx

from ..config import EXPORT_DIR, WORKSPACE_DIR, settings
from ..approvals import ApprovalRequired, consume as consume_approval
from ..approval_policy import (
    ASK as DEFAULT_APPROVAL_POLICY,
    builtin_risk,
    normalize as normalize_approval_policy,
)
from ..net_guard import validate_outbound_url
from .. import guardrails
from ..guardrail_policies import enforce_content, ContentBlocked
from .presentation_requirements import PresentationRequirements


class ToolError(RuntimeError):
    """可安全展示给模型的确定性工具错误。"""

    def __init__(self, message: str, *, code: str = "tool_error", retryable: bool = False):
        super().__init__(message)
        self.code = str(code or "tool_error")[:64]
        self.retryable = bool(retryable)


_SENSITIVE_PATH_NAMES = {
    ".env",
    ".git",
    ".ssh",
    ".aws",
    ".azure",
    "credentials",
    "credentials.json",
    "secrets.json",
    "id_rsa",
    "id_ed25519",
}


def workspace_for_run(
    user_id: int | None, run_id: str | None, agent_id: int | None = None
) -> Path:
    """返回不可跨用户/跨运行共享的工作区，并拒绝路径字符注入。"""
    user_part = re.sub(r"[^a-zA-Z0-9_-]", "_", str(user_id or "anonymous"))[:64]
    run_part = re.sub(r"[^a-zA-Z0-9_-]", "_", str(run_id or uuid.uuid4().hex))[:96]
    agent_part = re.sub(r"[^a-zA-Z0-9_-]", "_", str(agent_id or "default"))[:64]
    # 用户与智能体是稳定空间边界；run 子目录继续保证同一空间中的并发任务互不覆盖。
    root = (
        WORKSPACE_DIR / f"user_{user_part}" / f"agent_{agent_part}" / f"run_{run_part}"
    ).resolve()
    root.relative_to(WORKSPACE_DIR.resolve())
    root.mkdir(parents=True, exist_ok=True)
    return root


@dataclass
class BuiltinToolContext:
    root: Path = WORKSPACE_DIR
    user_id: int | None = None
    agent_id: int | None = None
    session_id: str | None = None
    llm: Any = None
    sub_agents: list[dict] = field(default_factory=list)
    runtime_event: Callable[[str, dict], Awaitable[None] | None] | None = None
    subagent_depth: int = 0
    run_id: str | None = None
    approval_tokens: list[str] = field(default_factory=list)
    # 每个 Turn 在入队时固化。它只决定已授权写操作是否需要交互批准，
    # 不扩大工具、工作区、网络、资源所有权或角色权限。
    approval_policy: str = DEFAULT_APPROVAL_POLICY
    # execution_id 标识当前执行主体；顶层等于 run_id，内联子智能体使用独立 ID。
    # parent_run_id 只用于父子审计，不改变审批仍绑定的持久化 Job run_id。
    execution_id: str | None = None
    parent_run_id: str | None = None
    artifacts: list[str] = field(default_factory=list)
    # Trusted API callback persists completed files before later plan/model failure.
    artifact_callback: Callable[[str], Awaitable[None]] | None = None
    attachment_images: list[Path] = field(default_factory=list)
    required_artifact_kinds: set[str] = field(default_factory=set)
    active_skill_names: set[str] = field(default_factory=set)
    artifact_failures: dict[str, str] = field(default_factory=dict)
    artifact_metadata: dict[str, dict] = field(default_factory=dict)
    # Harness 计划的可观测镜像，供回答完成后的模板/Artifact 后处理真实回写步骤终态。
    plan_steps: list[dict] = field(default_factory=list)
    plan_revision: int = 0
    deferred_artifact_kinds: set[str] = field(default_factory=set)
    # 由已授权附件与用户目标确定；模型工具参数不得覆盖来源页数或语言要求。
    presentation_requirements: PresentationRequirements | None = None
    # None 仅用于测试/内部直接调用，表示使用全部工具；真实 Chat Turn 总是传入
    # Root 配置与 Agent 绑定计算后的显式集合。
    enabled_tools: set[str] | None = None


@dataclass(frozen=True)
class BuiltinTool:
    name: str
    description: str
    parameters: dict
    handler: Callable[[dict, BuiltinToolContext], Awaitable[Any] | Any]
    group: str
    mutating: bool = False

    def spec(self) -> dict:
        return {
            "type": "function",
            "function": {
                "name": self.name,
                "description": self.description,
                "parameters": self.parameters,
            },
        }


def _object(properties: dict, required: list[str] | None = None) -> dict:
    return {
        "type": "object",
        "properties": properties,
        "required": required or [],
        "additionalProperties": False,
    }


def _json_result(*, ok: bool = True, **values) -> str:
    return json.dumps({"ok": ok, **values}, ensure_ascii=False, default=str)


_WEB_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0 Safari/537.36 HarnessAgent/2.0"
    ),
    "Accept": "text/html,application/xhtml+xml,application/json,application/xml;q=0.9,text/plain;q=0.8,*/*;q=0.5",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.7",
}
_SEARCH_CACHE: dict[tuple, tuple[float, dict]] = {}


def _plain_html(fragment: str) -> str:
    value = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", fragment or "")
    value = re.sub(r"(?s)<[^>]+>", " ", value)
    value = html.unescape(value)
    return re.sub(r"\s+", " ", value).strip()


def _normalize_result_url(raw: str) -> str:
    value = html.unescape(raw or "").strip()
    if value.startswith("//"):
        value = "https:" + value
    parsed = urlparse(value)
    if parsed.netloc.endswith("duckduckgo.com") and parsed.path.startswith("/l/"):
        value = unquote((parse_qs(parsed.query).get("uddg") or [value])[0])
        parsed = urlparse(value)
    if parsed.scheme not in ("http", "https"):
        return ""
    ignored = {"utm_source", "utm_medium", "utm_campaign", "utm_term", "utm_content"}
    query = [
        (key, item)
        for key, values in parse_qs(parsed.query, keep_blank_values=True).items()
        if key.lower() not in ignored
        for item in values
    ]
    return urlunparse((
        parsed.scheme.lower(), parsed.netloc.lower(), parsed.path or "/",
        "", urlencode(query), "",
    ))


def _parse_duckduckgo(text: str) -> list[dict]:
    links = list(re.finditer(
        r'(?is)<a[^>]*class="[^"]*result__a[^"]*"[^>]*href="([^"]+)"[^>]*>(.*?)</a>',
        text,
    ))
    snippets = list(re.finditer(
        r'(?is)<(?:a|div)[^>]*class="[^"]*result__snippet[^"]*"[^>]*>(.*?)</(?:a|div)>',
        text,
    ))
    rows = []
    for index, match in enumerate(links):
        url = _normalize_result_url(match.group(1))
        if not url:
            continue
        snippet = _plain_html(snippets[index].group(1)) if index < len(snippets) else ""
        rows.append({
            "title": _plain_html(match.group(2)),
            "url": url,
            "snippet": snippet,
            "source": "duckduckgo",
        })
    return rows


def _parse_bing(text: str) -> list[dict]:
    if text.lstrip().startswith("<?xml") or "<rss" in text[:500].lower():
        rows = []
        try:
            root = ET.fromstring(text)
        except ET.ParseError:
            return rows
        for item in root.findall(".//item"):
            url = _normalize_result_url(item.findtext("link") or "")
            if not url:
                continue
            rows.append({
                "title": (item.findtext("title") or url).strip(),
                "url": url,
                "snippet": _plain_html(item.findtext("description") or ""),
                "source": "bing",
            })
        return rows
    rows = []
    for block in re.findall(
        r'(?is)<li[^>]*class="[^"]*\bb_algo\b[^"]*"[^>]*>(.*?)</li>',
        text,
    ):
        link = re.search(r'(?is)<h2[^>]*>.*?<a[^>]*href="([^"]+)"[^>]*>(.*?)</a>', block)
        if not link:
            continue
        url = _normalize_result_url(link.group(1))
        if not url:
            continue
        paragraph = re.search(r"(?is)<p[^>]*>(.*?)</p>", block)
        rows.append({
            "title": _plain_html(link.group(2)),
            "url": url,
            "snippet": _plain_html(paragraph.group(1)) if paragraph else "",
            "source": "bing",
        })
    return rows


def _merge_search_results(groups: list[list[dict]], limit: int) -> list[dict]:
    """以倒数排名融合多引擎结果，避免简单拼接让首个引擎垄断结果。"""
    merged: dict[str, dict] = {}
    for rows in groups:
        for rank, row in enumerate(rows, 1):
            url = row.get("url") or ""
            if not url:
                continue
            current = merged.setdefault(url, {
                **row, "sources": [], "score": 0.0,
            })
            current["score"] += 1.0 / (60 + rank)
            source = row.get("source")
            if source and source not in current["sources"]:
                current["sources"].append(source)
            if not current.get("snippet") and row.get("snippet"):
                current["snippet"] = row["snippet"]
    values = sorted(
        merged.values(),
        key=lambda item: (
            _is_search_intermediary_url(item["url"]),
            -item["score"],
            item["url"],
        ),
    )[:limit]
    for item in values:
        item["score"] = round(item["score"], 6)
        item["source"] = "+".join(item.pop("sources"))
    return values


def _is_search_intermediary_url(url: str) -> bool:
    """识别不能直接提供正文、通常依赖浏览器脚本跳转的搜索中间页。"""
    parsed = urlparse(url or "")
    host = (parsed.hostname or "").lower()
    return bool(
        host == "news.google.com"
        and parsed.path.startswith(("/rss/articles/", "/articles/"))
    )


async def _searxng_search(
    query: str, count: int, language: str, time_range: str
) -> list[dict]:
    base = settings.SEARXNG_URL
    if not base:
        return []
    url = base + "/search"
    await asyncio.to_thread(validate_outbound_url, url)
    params = {"q": query, "format": "json", "language": language or "all"}
    if time_range:
        params["time_range"] = time_range
    async with httpx.AsyncClient(
        timeout=settings.WEB_SEARCH_TIMEOUT_SECONDS,
        headers=_WEB_HEADERS,
        follow_redirects=False,
    ) as client:
        response = await client.get(url, params=params)
        response.raise_for_status()
        data = response.json()
    rows = []
    for item in (data.get("results") or [])[: max(count * 3, count)]:
        result_url = _normalize_result_url(str(item.get("url") or ""))
        if not result_url:
            continue
        rows.append({
            "title": str(item.get("title") or result_url),
            "url": result_url,
            "snippet": _plain_html(str(item.get("content") or "")),
            "source": ",".join(item.get("engines") or ["searxng"]),
        })
    return rows[:count]


async def _fallback_search(
    query: str, count: int, language: str, time_range: str = ""
) -> list[dict]:
    encoded = quote_plus(query)
    # cn.bing.com 会把 RSS 请求 301 到不含 /search 的首页；直接使用 www 并显式传 mkt。
    bing_market = "zh-CN" if (language or "").lower().startswith("zh") else "en-US"
    ddg_time = {"day": "d", "month": "m", "year": "y"}.get(time_range, "")
    ddg_filter = f"&df={ddg_time}" if ddg_time else ""
    urls = [
        ("duckduckgo", f"https://html.duckduckgo.com/html/?q={encoded}{ddg_filter}"),
        (
            "bing",
            f"https://www.bing.com/search?format=rss&q={encoded}"
            f"&count={max(count, 10)}&mkt={bing_market}",
        ),
    ]

    async def fetch(engine: str, url: str) -> list[dict]:
        await asyncio.to_thread(validate_outbound_url, url)
        async with httpx.AsyncClient(
            timeout=settings.WEB_SEARCH_TIMEOUT_SECONDS,
            headers=_WEB_HEADERS,
            follow_redirects=False,
        ) as client:
            response = await client.get(url)
            response.raise_for_status()
        return _parse_duckduckgo(response.text) if engine == "duckduckgo" else _parse_bing(response.text)

    outcomes = await asyncio.gather(
        *(fetch(engine, url) for engine, url in urls),
        return_exceptions=True,
    )
    groups = [value for value in outcomes if isinstance(value, list)]
    return _merge_search_results(groups, count)


async def _google_news_search(
    query: str, count: int, language: str, time_range: str = ""
) -> list[dict]:
    """通过 Google News RSS 获取稳定的新闻标题与来源链接。"""
    period = {"day": "1d", "month": "30d", "year": "365d"}.get(time_range, "")
    search_query = query.strip()
    # 用户请求常含“整理前 50 条并提供链接”等交付指令；RSS 搜索只保留主题词。
    search_query = re.sub(
        r"[，,。；;]*(?:并|且|and)?\s*(?:提供|附上|with)\s*(?:来源|source)?\s*(?:链接|links?).*$",
        "",
        search_query,
        flags=re.I,
    )
    search_query = re.sub(r"(?:前\s*|top\s*)\d{1,3}\s*(?:条|个|篇|项|items?)?", " ", search_query, flags=re.I)
    search_query = re.sub(r"(?:请|帮我|整理|汇总|列出|搜索|查找|compile|list|find)", " ", search_query, flags=re.I)
    search_query = re.sub(r"(?:国内外|的)", " ", search_query)
    search_query = re.sub(r"\s+", " ", search_query).strip() or query.strip()
    if period and "when:" not in search_query.lower():
        search_query += f" when:{period}"
    zh = (language or "").lower().startswith("zh")
    locale = "hl=zh-CN&gl=CN&ceid=CN:zh-Hans" if zh else "hl=en-US&gl=US&ceid=US:en"
    url = f"https://news.google.com/rss/search?q={quote_plus(search_query)}&{locale}"
    await asyncio.to_thread(validate_outbound_url, url)
    async with httpx.AsyncClient(
        timeout=settings.WEB_SEARCH_TIMEOUT_SECONDS,
        headers=_WEB_HEADERS,
        follow_redirects=False,
    ) as client:
        response = await client.get(url)
        response.raise_for_status()
    rows = _parse_bing(response.text)
    for row in rows:
        row["source"] = "google_news"
    return rows[:count]


async def _web_search(args: dict, ctx: BuiltinToolContext):
    query = str(args.get("query") or "").strip()
    if not query:
        raise ToolError("query 不能为空")
    count = max(1, min(int(args.get("count") or 8), 20))
    language = str(args.get("language") or "zh-CN").strip()
    time_range = str(args.get("time_range") or "").strip()
    key = (query, count, language, time_range, settings.SEARXNG_URL)
    now = asyncio.get_running_loop().time()
    cached = _SEARCH_CACHE.get(key)
    if cached and now - cached[0] < 300:
        return {**cached[1], "cached": True}

    warnings: list[str] = []
    groups: list[list[dict]] = []
    providers: list[str] = []
    is_news = bool(re.search(r"(?:新闻|资讯|news|headline)", query, re.I))
    if settings.SEARXNG_URL:
        try:
            rows = await _searxng_search(query, count, language, time_range)
            if rows:
                groups.append(rows)
                providers.append("searxng")
            else:
                warnings.append("SearXNG 未返回结果")
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"SearXNG 不可用：{type(exc).__name__}")
    if is_news:
        try:
            rows = await _google_news_search(query, count, language, time_range)
            if rows:
                groups.append(rows)
                providers.append("google_news")
            else:
                warnings.append("Google News 未返回结果")
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"Google News 不可用：{type(exc).__name__}")
    # 新闻 RSS 的链接可能是依赖 JavaScript 的中间页，因此新闻任务即使 RSS 已有
    # 结果也同时聚合 DuckDuckGo/Bing，优先为 web_fetch 提供可直接读取的原文 URL。
    if settings.WEB_SEARCH_FALLBACK and (is_news or not groups):
        try:
            rows = await _fallback_search(query, count, language, time_range)
            if rows:
                groups.append(rows)
                for row in rows:
                    for source in str(row.get("source") or "").split("+"):
                        if source and source not in providers:
                            providers.append(source)
            else:
                warnings.append("DuckDuckGo/Bing 未返回结果")
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"免费聚合检索不可用：{type(exc).__name__}")
    results = _merge_search_results(groups, count)
    if not results:
        raise ToolError("联网搜索未返回结果；" + "；".join(warnings))
    value = {
        "query": query,
        "provider": "+".join(providers),
        "results": results,
        "warnings": warnings,
        "cached": False,
    }
    if len(_SEARCH_CACHE) >= 256:
        oldest = min(_SEARCH_CACHE, key=lambda item: _SEARCH_CACHE[item][0])
        _SEARCH_CACHE.pop(oldest, None)
    _SEARCH_CACHE[key] = (now, value)
    return value


class _ReadableHtml(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.skip = 0
        self.title_depth = 0
        self.title: list[str] = []
        self.text: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript", "svg", "canvas"}:
            self.skip += 1
        if tag == "title":
            self.title_depth += 1
        if not self.skip and tag in {"p", "div", "article", "section", "li", "h1", "h2", "h3", "br"}:
            self.text.append("\n")

    def handle_endtag(self, tag):
        if tag == "title" and self.title_depth:
            self.title_depth -= 1
        if tag in {"script", "style", "noscript", "svg", "canvas"} and self.skip:
            self.skip -= 1

    def handle_data(self, data):
        if self.skip:
            return
        if self.title_depth:
            self.title.append(data)
        else:
            self.text.append(data)


async def _safe_web_get(url: str) -> tuple[httpx.Response, str]:
    timeout = settings.WEB_SEARCH_TIMEOUT_SECONDS
    retries = max(0, min(int(settings.WEB_REQUEST_MAX_RETRIES), 3))
    last_error: Exception | None = None
    for attempt in range(retries + 1):
        current = url
        try:
            async with httpx.AsyncClient(
                timeout=timeout, headers=_WEB_HEADERS, follow_redirects=False
            ) as client:
                for _ in range(6):
                    # 每一次重定向都重新解析并检查目标，不能只校验初始 URL。
                    await asyncio.to_thread(validate_outbound_url, current)
                    async with client.stream("GET", current) as response:
                        if response.is_redirect:
                            target = response.headers.get("location")
                            if not target:
                                raise ToolError("网页重定向缺少 Location")
                            current = urljoin(current, target)
                            continue
                        if response.status_code == 429 or response.status_code >= 500:
                            response.raise_for_status()
                        if response.status_code >= 400:
                            raise ToolError(
                                f"web_fetch HTTP {response.status_code}",
                                code=f"http_{response.status_code}",
                                retryable=False,
                            )
                        content_type = response.headers.get("content-type", "").lower()
                        allowed = (
                            "text/", "application/json", "application/xml",
                            "application/xhtml", "application/rss+xml", "application/atom+xml",
                        )
                        if not any(value in content_type for value in allowed):
                            raise ToolError(f"web_fetch 不支持内容类型：{content_type or '未知'}")
                        chunks = bytearray()
                        async for chunk in response.aiter_bytes():
                            chunks.extend(chunk)
                            if len(chunks) > settings.WEB_FETCH_MAX_BYTES:
                                raise ToolError("网页内容超过读取上限")
                        encoding = response.encoding or "utf-8"
                        return response, bytes(chunks).decode(encoding, errors="replace")
                raise ToolError("网页重定向次数过多")
        except ToolError:
            raise
        except (httpx.TimeoutException, httpx.TransportError, httpx.HTTPStatusError) as exc:
            last_error = exc
            status = exc.response.status_code if isinstance(exc, httpx.HTTPStatusError) else 0
            retryable = not status or status == 429 or status >= 500
            if not retryable or attempt >= retries:
                code = f"http_{status}" if status else "network_error"
                raise ToolError(
                    f"web_fetch 请求失败：{type(exc).__name__}",
                    code=code,
                    retryable=retryable,
                ) from exc
            await asyncio.sleep(min(1.5, 0.25 * (2 ** attempt)))
    raise ToolError(
        f"web_fetch 请求失败：{type(last_error).__name__ if last_error else '未知错误'}",
        code="network_error",
        retryable=True,
    )


async def _web_fetch(args: dict, ctx: BuiltinToolContext):
    url = str(args.get("url") or "").strip()
    if not url:
        raise ToolError("url 不能为空")
    response, body = await _safe_web_get(url)
    content_type = response.headers.get("content-type", "").lower()
    title = ""
    if "html" in content_type:
        parser = _ReadableHtml()
        parser.feed(body)
        title = re.sub(r"\s+", " ", "".join(parser.title)).strip()
        text = "\n".join(
            line.strip() for line in "".join(parser.text).splitlines() if line.strip()
        )
    elif "json" in content_type:
        try:
            text = json.dumps(json.loads(body), ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            text = body
    else:
        text = body
    text = text.strip()
    if not text:
        raise ToolError(
            "网页未提取到可用正文；请改用搜索结果中的原始来源 URL",
            code="empty_content",
            retryable=False,
        )
    original_chars = len(text)
    text = text[: settings.WEB_FETCH_MAX_CHARS]
    return {
        "requested_url": url,
        "url": str(response.url),
        "title": title,
        "content_type": content_type.split(";")[0],
        "text": text,
        "chars": len(text),
        "truncated": original_chars > settings.WEB_FETCH_MAX_CHARS,
    }


def _limit(value: str, maximum: int = 20000) -> tuple[str, bool]:
    value = value or ""
    if len(value) <= maximum:
        return value, False
    return value[:maximum] + f"\n…（已截断至 {maximum} 字符）", True


def _root(ctx: BuiltinToolContext) -> Path:
    root = Path(ctx.root or WORKSPACE_DIR).resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def _resolve(ctx: BuiltinToolContext, raw: str | None, *, must_exist: bool = False) -> Path:
    root = _root(ctx)
    text = str(raw or ".").strip()
    candidate = Path(text)
    if not candidate.is_absolute():
        candidate = root / candidate
    resolved = candidate.resolve(strict=False)
    try:
        relative = resolved.relative_to(root)
    except ValueError as exc:
        raise ToolError(f"路径越界：{text}；只允许访问工作区 {root}") from exc
    if any(part.lower() in _SENSITIVE_PATH_NAMES for part in relative.parts):
        raise ToolError(f"拒绝访问敏感路径：{text}")
    if must_exist and not resolved.exists():
        raise ToolError(f"路径不存在：{text}")
    return resolved


def _relative(ctx: BuiltinToolContext, path: Path) -> str:
    return path.resolve().relative_to(_root(ctx)).as_posix() or "."


def _confirmed(args: dict, action: str) -> None:
    if args.get("confirm") is not True:
        raise ToolError(f"{action}需要显式设置 confirm=true")


def _iter_files(base: Path, recursive: bool = True):
    iterator = base.rglob("*") if recursive else base.glob("*")
    for path in iterator:
        if path.is_file() and ".git" not in path.parts:
            yield path


async def _workspace_ls(args: dict, ctx: BuiltinToolContext):
    path = _resolve(ctx, args.get("path"), must_exist=True)
    if not path.is_dir():
        raise ToolError("ls 目标必须是目录")
    limit = min(max(int(args.get("limit") or 200), 1), 2000)
    rows = []
    for item in sorted(path.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower()))[:limit]:
        stat = item.stat()
        rows.append({
            "path": _relative(ctx, item),
            "type": "directory" if item.is_dir() else "file",
            "size": stat.st_size if item.is_file() else None,
            "modified_at": dt.datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(),
        })
    return _json_result(path=_relative(ctx, path), entries=rows, truncated=len(rows) >= limit)


async def _workspace_glob(args: dict, ctx: BuiltinToolContext):
    pattern = str(args.get("pattern") or "").strip()
    if not pattern:
        raise ToolError("pattern 不能为空")
    base = _resolve(ctx, args.get("path") or ".", must_exist=True)
    if not base.is_dir():
        raise ToolError("glob 起始路径必须是目录")
    limit = min(max(int(args.get("limit") or 500), 1), 5000)
    matches = []
    for path in base.glob(pattern):
        resolved = _resolve(ctx, str(path))
        matches.append({
            "path": _relative(ctx, resolved),
            "type": "directory" if resolved.is_dir() else "file",
        })
        if len(matches) >= limit:
            break
    return _json_result(pattern=pattern, matches=matches, truncated=len(matches) >= limit)


async def _workspace_grep(args: dict, ctx: BuiltinToolContext):
    query = str(args.get("query") or "")
    if not query:
        raise ToolError("query 不能为空")
    base = _resolve(ctx, args.get("path") or ".", must_exist=True)
    regex = bool(args.get("regex"))
    case_sensitive = bool(args.get("case_sensitive"))
    flags = 0 if case_sensitive else re.IGNORECASE
    try:
        matcher = re.compile(query if regex else re.escape(query), flags)
    except re.error as exc:
        raise ToolError(f"正则表达式无效：{exc}") from exc
    include = str(args.get("include") or "*")
    max_results = min(max(int(args.get("max_results") or 100), 1), 1000)
    paths = [base] if base.is_file() else _iter_files(base, bool(args.get("recursive", True)))
    results = []
    skipped_binary = 0
    for path in paths:
        if not fnmatch.fnmatch(path.name, include):
            continue
        try:
            raw = path.read_bytes()
            if b"\x00" in raw[:4096]:
                skipped_binary += 1
                continue
            text = raw.decode("utf-8", errors="replace")
        except OSError:
            continue
        for number, line in enumerate(text.splitlines(), 1):
            if matcher.search(line):
                results.append({
                    "path": _relative(ctx, path),
                    "line": number,
                    "text": line[:1000],
                })
                if len(results) >= max_results:
                    return _json_result(
                        query=query, matches=results, truncated=True,
                        skipped_binary=skipped_binary,
                    )
    return _json_result(
        query=query, matches=results, truncated=False, skipped_binary=skipped_binary
    )


def _read_one(path: Path, ctx: BuiltinToolContext, start_line: int, end_line: int | None) -> dict:
    raw = path.read_bytes()
    if b"\x00" in raw[:4096]:
        raise ToolError(f"二进制文件不能按文本读取：{_relative(ctx, path)}")
    text = raw.decode("utf-8", errors="replace")
    lines = text.splitlines()
    start = max(start_line, 1)
    end = min(end_line or len(lines), len(lines), start + 1999)
    selected = "\n".join(f"{idx}: {lines[idx - 1]}" for idx in range(start, end + 1))
    selected, truncated = _limit(selected)
    return {
        "path": _relative(ctx, path),
        "content": selected,
        "start_line": start,
        "end_line": end,
        "total_lines": len(lines),
        "truncated": truncated or (end_line is not None and end < end_line),
    }


async def _workspace_read(args: dict, ctx: BuiltinToolContext):
    path = _resolve(ctx, args.get("path"), must_exist=True)
    if not path.is_file():
        raise ToolError("read 目标必须是文件")
    return _json_result(**_read_one(
        path, ctx, int(args.get("start_line") or 1),
        int(args["end_line"]) if args.get("end_line") is not None else None,
    ))


async def _workspace_read_many(args: dict, ctx: BuiltinToolContext):
    paths = args.get("paths") or []
    if not isinstance(paths, list) or not paths:
        raise ToolError("paths 必须是非空数组")
    if len(paths) > 20:
        raise ToolError("单次最多读取 20 个文件")
    rows = []
    errors = []
    for raw in paths:
        try:
            path = _resolve(ctx, raw, must_exist=True)
            if not path.is_file():
                raise ToolError("目标不是文件")
            rows.append(_read_one(path, ctx, 1, 500))
        except Exception as exc:  # noqa: BLE001 - 聚合返回单文件错误
            errors.append({"path": str(raw), "error": str(exc)})
    return _json_result(files=rows, errors=errors)


def _atomic_write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    temp.write_text(content, encoding="utf-8", newline="")
    temp.replace(path)


async def _workspace_write(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "写入文件")
    path = _resolve(ctx, args.get("path"))
    overwrite = bool(args.get("overwrite"))
    if path.exists() and not overwrite:
        raise ToolError("文件已存在；若确需覆盖，请设置 overwrite=true")
    content = str(args.get("content") or "")
    _atomic_write(path, content)
    return _json_result(path=_relative(ctx, path), bytes=len(content.encode("utf-8")))


def _replace_exact(text: str, old: str, new: str, expected: int) -> tuple[str, int]:
    if old == "":
        raise ToolError("old 不能为空")
    count = text.count(old)
    if count != expected:
        raise ToolError(f"精确替换数量不匹配：期望 {expected}，实际 {count}")
    return text.replace(old, new), count


async def _workspace_edit(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "编辑文件")
    path = _resolve(ctx, args.get("path"), must_exist=True)
    if not path.is_file():
        raise ToolError("edit 目标必须是文件")
    text = path.read_text(encoding="utf-8")
    changed, count = _replace_exact(
        text, str(args.get("old") or ""), str(args.get("new") or ""),
        int(args.get("expected_replacements") or 1),
    )
    _atomic_write(path, changed)
    return _json_result(path=_relative(ctx, path), replacements=count)


async def _workspace_multi_edit(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "批量编辑")
    edits = args.get("edits") or []
    if not isinstance(edits, list) or not edits:
        raise ToolError("edits 必须是非空数组")
    if len(edits) > 50:
        raise ToolError("单次最多 50 项编辑")
    staged: dict[Path, str] = {}
    counts = []
    for index, edit in enumerate(edits):
        if not isinstance(edit, dict):
            raise ToolError(f"第 {index + 1} 项不是对象")
        path = _resolve(ctx, edit.get("path"), must_exist=True)
        text = staged.get(path)
        if text is None:
            text = path.read_text(encoding="utf-8")
        text, count = _replace_exact(
            text, str(edit.get("old") or ""), str(edit.get("new") or ""),
            int(edit.get("expected_replacements") or 1),
        )
        staged[path] = text
        counts.append({"path": _relative(ctx, path), "replacements": count})
    for path, content in staged.items():
        _atomic_write(path, content)
    return _json_result(files=len(staged), edits=counts)


async def _workspace_apply_patch(args: dict, ctx: BuiltinToolContext):
    """应用结构化补丁；先在内存完成全部校验，再原子写入，避免半成功。"""
    _confirmed(args, "应用补丁")
    files = args.get("files") or []
    if not isinstance(files, list) or not files:
        raise ToolError("files 必须是非空数组")
    staged: dict[Path, str] = {}
    summary = []
    for item in files:
        if not isinstance(item, dict):
            raise ToolError("files 中每项必须是对象")
        action = str(item.get("action") or "update")
        path = _resolve(ctx, item.get("path"), must_exist=action == "update")
        if action == "add":
            if path.exists():
                raise ToolError(f"新增文件已存在：{_relative(ctx, path)}")
            staged[path] = str(item.get("content") or "")
        elif action == "update":
            text = path.read_text(encoding="utf-8")
            operations = item.get("operations") or []
            if not operations:
                raise ToolError(f"更新文件缺少 operations：{_relative(ctx, path)}")
            for operation in operations:
                text, _ = _replace_exact(
                    text, str(operation.get("old") or ""), str(operation.get("new") or ""),
                    int(operation.get("expected_replacements") or 1),
                )
            staged[path] = text
        elif action == "delete":
            raise ToolError("结构化补丁不允许删除文件；请由用户在管理界面执行删除")
        else:
            raise ToolError(f"未知补丁动作：{action}")
        summary.append({"path": _relative(ctx, path), "action": action})
    for path, content in staged.items():
        _atomic_write(path, content)
    return _json_result(files=summary)


_SHELL_META = re.compile(r"[;&|><`$]|[\r\n]|\$\(|\$\{")
_SHELL_ALLOWED = {
    "powershell": {
        "get-childitem", "get-content", "select-string", "test-path",
        "get-command", "where-object", "sort-object", "measure-object",
        "python", "python.exe", "pytest", "git", "rg", "node", "npm", "npx",
    },
    "bash": {
        "ls", "pwd", "find", "grep", "sed", "head", "tail", "wc",
        "python", "python3", "pytest", "git", "rg", "node", "npm", "npx",
    },
}
_INLINE_FLAGS = {"-c", "-command", "-encodedcommand", "-e", "--eval", "-file"}


def _shell_executable(shell: str) -> list[str]:
    if shell == "powershell":
        exe = shutil.which("pwsh") or shutil.which("powershell")
        if not exe:
            raise ToolError("未安装 PowerShell")
        return [exe, "-NoLogo", "-NoProfile", "-NonInteractive", "-Command"]
    if shell == "bash":
        exe = shutil.which("bash")
        if not exe:
            git_bash = Path(r"C:\Program Files\Git\bin\bash.exe")
            exe = str(git_bash) if git_bash.exists() else ""
        if not exe:
            raise ToolError("未安装 Bash")
        return [exe, "--noprofile", "--norc", "-c"]
    raise ToolError("shell 只能是 powershell 或 bash")


def _validate_shell_script(shell: str, script: str) -> None:
    if not settings.SHELL_TOOL_ENABLED:
        raise ToolError("Shell 工具已由系统管理员关闭")
    if not settings.SHELL_SANDBOX_COMMAND:
        raise ToolError("Shell 缺少 OS 级低权限沙箱包装器，拒绝执行")
    if not script.strip():
        raise ToolError("script 不能为空")
    if _SHELL_META.search(script):
        raise ToolError("受限 Shell 不允许管道、重定向、命令连接、变量展开或多行脚本")
    try:
        tokens = shlex.split(script, posix=shell == "bash")
    except ValueError as exc:
        raise ToolError(f"命令解析失败：{exc}") from exc
    if not tokens:
        raise ToolError("script 不能为空")
    executable_token = tokens[0].strip("\"'")
    if "/" in executable_token or "\\" in executable_token:
        raise ToolError("受限 Shell 的命令必须使用白名单名称，不能指定可执行文件路径")
    command = Path(executable_token).name.lower()
    if command not in _SHELL_ALLOWED[shell]:
        raise ToolError(f"命令不在 {shell} 白名单中：{command}")
    if command in {"python", "python.exe", "python3", "node", "powershell", "pwsh"}:
        lowered = {token.strip("\"'").lower() for token in tokens[1:]}
        if lowered & _INLINE_FLAGS:
            raise ToolError("受限 Shell 禁止解释器内联代码或脚本文件参数")
    for token in tokens[1:]:
        clean = token.strip("\"'")
        if re.match(r"^[a-zA-Z]:[\\/]", clean) or clean.startswith(("/", "\\\\")):
            raise ToolError("受限 Shell 不允许绝对路径参数")
        if ".." in Path(clean).parts:
            raise ToolError("受限 Shell 不允许父目录路径")


async def _shell_run(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "执行 Shell")
    shell = str(args.get("shell") or "powershell").lower()
    script = str(args.get("script") or "")
    _validate_shell_script(shell, script)
    timeout = min(
        max(int(args.get("timeout_seconds") or settings.SHELL_TIMEOUT_SECONDS), 1),
        settings.SHELL_TIMEOUT_SECONDS,
    )
    temp_dir = _root(ctx) / ".runtime-tmp"
    temp_dir.mkdir(parents=True, exist_ok=True)
    env = {
        "PATH": os.environ.get("PATH", ""),
        "SystemRoot": os.environ.get("SystemRoot", r"C:\Windows"),
        "TEMP": str(temp_dir),
        "TMP": str(temp_dir),
        "PYTHONIOENCODING": "utf-8",
        "NO_COLOR": "1",
    }
    try:
        proc = await asyncio.to_thread(
            subprocess.run,
            [
                *shlex.split(settings.SHELL_SANDBOX_COMMAND, posix=os.name != "nt"),
                *_shell_executable(shell),
                script,
            ],
            cwd=str(_root(ctx)),
            env=env,
            capture_output=True,
            timeout=timeout,
            check=False,
        )
    except subprocess.TimeoutExpired:
        raise ToolError(f"命令超过 {timeout} 秒，已终止")
    out, out_truncated = _limit(
        proc.stdout.decode("utf-8", errors="replace"), settings.SHELL_OUTPUT_CHARS
    )
    err, err_truncated = _limit(
        proc.stderr.decode("utf-8", errors="replace"), settings.SHELL_OUTPUT_CHARS
    )
    return _json_result(
        ok=proc.returncode == 0, exit_code=proc.returncode,
        stdout=out, stderr=err, truncated=out_truncated or err_truncated,
    )


async def _run_process(
    ctx: BuiltinToolContext, command: list[str], timeout: int = 60
) -> tuple[int, str, str]:
    try:
        proc = await asyncio.to_thread(
            subprocess.run,
            command,
            cwd=str(_root(ctx)),
            env={
                "PATH": os.environ.get("PATH", ""),
                "SystemRoot": os.environ.get("SystemRoot", r"C:\Windows"),
                "TEMP": os.environ.get("TEMP", str(EXPORT_DIR)),
                "TMP": os.environ.get("TMP", str(EXPORT_DIR)),
                "GIT_TERMINAL_PROMPT": "0",
                "NO_COLOR": "1",
            },
            capture_output=True,
            timeout=timeout,
            check=False,
        )
    except subprocess.TimeoutExpired:
        raise ToolError(f"命令超过 {timeout} 秒，已终止")
    return (
        proc.returncode,
        _limit(proc.stdout.decode("utf-8", errors="replace"))[0],
        _limit(proc.stderr.decode("utf-8", errors="replace"))[0],
    )


async def _git_status(args: dict, ctx: BuiltinToolContext):
    code, out, err = await _run_process(ctx, ["git", "status", "--short", "--branch"])
    return _json_result(ok=code == 0, exit_code=code, status=out, stderr=err)


async def _git_diff(args: dict, ctx: BuiltinToolContext):
    command = ["git", "diff"]
    if args.get("staged"):
        command.append("--staged")
    if args.get("path"):
        path = _resolve(ctx, args["path"], must_exist=True)
        command.extend(["--", _relative(ctx, path)])
    code, out, err = await _run_process(ctx, command)
    return _json_result(ok=code == 0, exit_code=code, diff=out, stderr=err)


async def _git_commit(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "Git 提交")
    message = str(args.get("message") or "").strip()
    if not message or len(message) > 200:
        raise ToolError("提交说明必须为 1～200 字符")
    paths = args.get("paths") or []
    if not paths:
        raise ToolError("必须明确指定要暂存的 paths；禁止隐式提交全部变更")
    rels = []
    for raw in paths[:100]:
        # 删除项在工作树中已不存在，但仍需要交给 git add 记录删除。
        path = _resolve(ctx, raw, must_exist=False)
        rels.append(_relative(ctx, path))
    code, out, err = await _run_process(ctx, ["git", "add", "--", *rels])
    if code != 0:
        return _json_result(ok=False, stage_stdout=out, stderr=err)
    code, out, err = await _run_process(ctx, ["git", "commit", "-m", message])
    return _json_result(ok=code == 0, exit_code=code, stdout=out, stderr=err)


async def _lsp(args: dict, ctx: BuiltinToolContext):
    """确定性代码诊断入口；优先使用已安装语言工具，缺失时给出明确建议。"""
    action = str(args.get("action") or "diagnostics")
    path = _resolve(ctx, args.get("path") or ".", must_exist=True)
    if action == "diagnostics":
        if path.is_file() and path.suffix == ".py":
            command = [os.sys.executable, "-m", "py_compile", str(path)]
        elif path.is_dir() and any(path.rglob("*.py")):
            command = [os.sys.executable, "-m", "compileall", "-q", str(path)]
        elif shutil.which("npx") and (
            (path / "tsconfig.json").exists() if path.is_dir() else path.suffix in {".ts", ".tsx"}
        ):
            command = ["npx", "tsc", "--noEmit", "--pretty", "false"]
        else:
            raise ToolError("未找到可用诊断器；Python 使用 py_compile，TypeScript 需要 npx/tsc")
        code, out, err = await _run_process(ctx, command)
        return _json_result(ok=code == 0, action=action, stdout=out, diagnostics=err)
    symbol = str(args.get("symbol") or "").strip()
    if action in {"definition", "references", "symbols"}:
        if not symbol and action != "symbols":
            raise ToolError(f"{action} 需要 symbol")
        query = symbol or r"^(class|def|function|const|let|var)\s+"
        return await _workspace_grep({
            "query": query,
            "path": _relative(ctx, path),
            "regex": True,
            "case_sensitive": True,
            "recursive": True,
            "include": "*",
            "max_results": 200,
        }, ctx)
    raise ToolError("action 只能是 diagnostics、definition、references 或 symbols")


async def _html_artifact(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "生成 HTML 产物")
    title = str(args.get("title") or "artifact")[:80]
    body = str(args.get("html") or "")
    if not body.strip():
        raise ToolError("html 不能为空")
    safe_title = re.sub(r'[\\/:*?"<>|]', "_", title) or "artifact"
    name = f"{safe_title}_{uuid.uuid4().hex[:8]}.html"
    target = EXPORT_DIR / name
    if "<html" not in body.lower():
        body = (
            "<!doctype html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\">"
            f"<meta name=\"viewport\" content=\"width=device-width,initial-scale=1\"><title>{html.escape(title)}</title>"
            "<meta http-equiv=\"Content-Security-Policy\" "
            "content=\"default-src 'none'; style-src 'unsafe-inline'; img-src data: https:;\">"
            f"</head><body>{body}</body></html>"
        )
    _atomic_write(target, body)
    return _json_result(file=name, path=str(target), download_url=f"/api/v1/exports/{name}")


def _docx_text_signature(path: Path) -> tuple[str, int]:
    """提取 DOCX 各正文 XML 的文字/换行序列，用于证明格式化没有改内容。"""
    values: list[tuple[str, list[str]]] = []
    chars = 0
    try:
        with zipfile.ZipFile(path) as archive:
            for member in sorted(
                name for name in archive.namelist()
                if name.startswith("word/") and name.endswith(".xml")
            ):
                try:
                    root = ET.fromstring(archive.read(member))
                except ET.ParseError:
                    continue
                tokens: list[str] = []
                for node in root.iter():
                    local = node.tag.rsplit("}", 1)[-1]
                    if local in {"t", "instrText", "delText"}:
                        text_value = node.text or ""
                        tokens.append(f"{local}:{text_value}")
                        chars += len(text_value)
                    elif local in {"tab", "br", "cr", "noBreakHyphen", "softHyphen"}:
                        tokens.append(f"<{local}>")
                if tokens:
                    values.append((member, tokens))
    except (OSError, zipfile.BadZipFile) as exc:
        raise ToolError(f"Word 文件无效：{exc}") from exc
    digest = __import__("hashlib").sha256(
        json.dumps(values, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    return digest, chars


def _docx_paragraphs(document):
    """遍历正文、嵌套表格、页眉和页脚中的段落并按 XML 节点去重。"""
    seen: set[int] = set()

    def emit(paragraphs):
        for paragraph in paragraphs:
            key = id(paragraph._p)
            if key not in seen:
                seen.add(key)
                yield paragraph

    def table_paragraphs(table):
        for row in table.rows:
            for cell in row.cells:
                yield from emit(cell.paragraphs)
                for nested in cell.tables:
                    yield from table_paragraphs(nested)

    yield from emit(document.paragraphs)
    for table in document.tables:
        yield from table_paragraphs(table)
    for section in document.sections:
        for area in (section.header, section.footer):
            yield from emit(area.paragraphs)
            for table in area.tables:
                yield from table_paragraphs(table)


def _set_run_font(run, east_asia: str, latin: str, size_pt: float, bold=None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bool(bold)


def _append_inline_markdown(paragraph, value: str) -> None:
    """把常见行内 Markdown 转为真实 Word run 格式，禁止语法标记泄漏到 DOCX。"""
    text_value = str(value or "")
    pattern = re.compile(
        r"(\*\*(?=\S)(.+?)(?<=\S)\*\*|`([^`]+)`|\*(?=\S)(.+?)(?<=\S)\*)"
    )
    position = 0
    for match in pattern.finditer(text_value):
        if match.start() > position:
            paragraph.add_run(text_value[position:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(match.group(2) or "")
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(match.group(3) or "")
            run.font.name = "Consolas"
        else:
            run = paragraph.add_run(match.group(4) or "")
            run.italic = True
        position = match.end()
    if position < len(text_value):
        paragraph.add_run(text_value[position:])


def _set_exact_table_geometry(table, total_width_dxa: int) -> None:
    """固定 tblW/tblInd/tblGrid/tcW，避免自动调整导致跨页列宽漂移。"""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Twips

    column_count = len(table.columns)
    if column_count < 1:
        return
    total_width = max(column_count, int(total_width_dxa))
    widths = [total_width // column_count] * column_count
    widths[-1] += total_width - sum(widths)

    def ensure(parent, tag):
        child = parent.find(qn(tag))
        if child is None:
            child = OxmlElement(tag)
            parent.append(child)
        return child

    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr
    for tag, width in (("w:tblW", total_width), ("w:tblInd", 120)):
        node = ensure(properties, tag)
        node.set(qn("w:type"), "dxa")
        node.set(qn("w:w"), str(width))
    layout = ensure(properties, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            row_properties = row._tr.get_or_add_trPr()
            repeat = ensure(row_properties, "w:tblHeader")
            repeat.set(qn("w:val"), "true")
        for column_index, cell in enumerate(row.cells):
            width = widths[column_index]
            cell.width = Twips(width)
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = ensure(cell_properties, "w:tcW")
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(width))
            margins = ensure(cell_properties, "w:tcMar")
            # Word's default table look keeps vertical padding compact.  Adding
            # 80 twips above and below every cell noticeably inflates long OCR
            # tables and can create otherwise-empty trailing pages.
            for side, margin_width in (("top", 0), ("bottom", 0), ("start", 120), ("end", 120)):
                margin = ensure(margins, f"w:{side}")
                margin.set(qn("w:type"), "dxa")
                margin.set(qn("w:w"), str(margin_width))


async def _document_inspect(args: dict, ctx: BuiltinToolContext):
    path = _resolve(ctx, args.get("path"), must_exist=True)
    if not path.is_file() or path.suffix.lower() != ".docx":
        raise ToolError("document_inspect 仅支持工作区中的 .docx 文件")
    try:
        from docx import Document
        document = Document(str(path))
    except Exception as exc:
        raise ToolError(f"无法打开 Word 文件：{exc}") from exc
    signature, chars = _docx_text_signature(path)
    paragraphs = list(_docx_paragraphs(document))
    return _json_result(
        path=_relative(ctx, path),
        paragraphs=len(paragraphs),
        nonempty_paragraphs=sum(bool(item.text.strip()) for item in paragraphs),
        tables=len(document.tables),
        sections=len(document.sections),
        text_chars=chars,
        content_sha256=signature,
        styles=sorted({str(item.style.name) for item in paragraphs if item.style})[:50],
    )


async def _presentation_create(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "创建 PowerPoint 演示文稿")
    from ..capabilities.presentations import create_presentation
    try:
        result = await asyncio.to_thread(
            create_presentation, args, EXPORT_DIR, requirements=ctx.presentation_requirements,
        )
    except (ValueError, ImportError) as exc:
        raise ToolError(str(exc)) from exc
    return json.dumps(result, ensure_ascii=False)


async def _document_create(args: dict, ctx: BuiltinToolContext):
    """从模型已经核验的文字和表格数据创建一个可下载的 DOCX。"""
    _confirmed(args, "创建 Word 文档")
    title = str(args.get("title") or "").strip()[:300]
    content = str(args.get("content") or "").strip()
    raw_tables = args.get("tables") or []
    if len(content) > 200_000:
        raise ToolError("Word 正文超过 200000 字符安全上限")
    if not isinstance(raw_tables, list):
        raise ToolError("tables 必须是表格数组")
    if not content and not title and not raw_tables:
        raise ToolError("Word 内容为空；请提供 title、content 或 tables")

    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt
    except ImportError as exc:
        raise ToolError("缺少 python-docx，无法创建 Word 文件") from exc

    document = Document()
    for section in document.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.6)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)

    if title:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _append_inline_markdown(paragraph, title)
        for run in paragraph.runs:
            _set_run_font(run, "黑体", "Arial", 18, True)

    paragraph_count = 0
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            paragraph = document.add_heading("", level=len(heading.group(1)))
            _append_inline_markdown(paragraph, heading.group(2).strip())
        elif re.match(r"^[-*]\s+", line):
            paragraph = document.add_paragraph(style="List Bullet")
            _append_inline_markdown(paragraph, re.sub(r"^[-*]\s+", "", line))
        elif re.match(r"^\d+[.)、]\s*", line):
            paragraph = document.add_paragraph(style="List Number")
            _append_inline_markdown(paragraph, re.sub(r"^\d+[.)、]\s*", "", line))
        else:
            paragraph = document.add_paragraph()
            _append_inline_markdown(paragraph, line)
        paragraph_count += 1
        for run in paragraph.runs:
            style_name = str(paragraph.style.name if paragraph.style else "").lower()
            is_heading = style_name.startswith("heading") or "标题" in style_name
            _set_run_font(
                run,
                "黑体" if is_heading else "宋体",
                "Arial" if is_heading else "Times New Roman",
                14 if is_heading else 12,
                True if is_heading else None,
            )

    table_count = 0
    table_row_count = 0
    section = document.sections[0]
    table_width_dxa = max(
        600,
        int(
            section.page_width.twips
            - section.left_margin.twips
            - section.right_margin.twips
            - 120
        ),
    )
    for index, raw_table in enumerate(raw_tables[:20], 1):
        if not isinstance(raw_table, dict):
            raise ToolError(f"第 {index} 个表格必须是对象")
        headers = [str(value or "")[:20_000] for value in (raw_table.get("headers") or [])]
        rows = raw_table.get("rows") or []
        if not isinstance(rows, list):
            raise ToolError(f"第 {index} 个表格的 rows 必须是数组")
        normalized_rows = [
            [str(value or "")[:20_000] for value in row]
            for row in rows[:2_000]
            if isinstance(row, list)
        ]
        column_count = max(
            [len(headers), *(len(row) for row in normalized_rows)],
            default=0,
        )
        if column_count < 1:
            continue
        table = document.add_table(rows=0, cols=column_count)
        table.style = "Table Grid"
        if headers:
            cells = table.add_row().cells
            for column in range(column_count):
                cells[column].text = headers[column] if column < len(headers) else ""
                for run in cells[column].paragraphs[0].runs:
                    _set_run_font(run, "黑体", "Arial", 10.5, True)
        for row in normalized_rows:
            cells = table.add_row().cells
            for column in range(column_count):
                cells[column].text = row[column] if column < len(row) else ""
                for run in cells[column].paragraphs[0].runs:
                    _set_run_font(run, "宋体", "Times New Roman", 10.5)
        _set_exact_table_geometry(table, table_width_dxa)
        table_count += 1
        table_row_count += len(normalized_rows)

    requested_name = Path(str(args.get("output_name") or "识别结果.docx")).name
    if not requested_name.lower().endswith(".docx"):
        requested_name += ".docx"
    safe_name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", requested_name).strip(" .")
    name = f"{Path(safe_name).stem[:120]}_{uuid.uuid4().hex[:8]}.docx"
    target = (EXPORT_DIR / name).resolve()
    target.relative_to(EXPORT_DIR.resolve())
    document.save(str(target))
    signature, chars = _docx_text_signature(target)
    if chars < 1:
        target.unlink(missing_ok=True)
        raise ToolError("Word 文档没有可验证的文字内容，已拒绝导出")
    return _json_result(
        file=name,
        download_url=f"/api/v1/exports/{name}",
        paragraphs=paragraph_count,
        tables=table_count,
        table_rows=table_row_count,
        text_chars=chars,
        content_sha256=signature,
    )


async def _document_format(args: dict, ctx: BuiltinToolContext):
    """只修改 DOCX 样式属性；保存前后以 XML 文字签名建立内容不变门禁。"""
    _confirmed(args, "规范 Word 格式")
    source = _resolve(ctx, args.get("path"), must_exist=True)
    if not source.is_file() or source.suffix.lower() != ".docx":
        raise ToolError("document_format 仅支持工作区中的 .docx 文件")
    east_asia = str(args.get("east_asia_font") or "宋体").strip()[:80]
    latin = str(args.get("latin_font") or "Times New Roman").strip()[:80]
    body_size = min(72.0, max(6.0, float(args.get("body_size_pt") or 12.0)))
    table_size = min(72.0, max(6.0, float(args.get("table_size_pt") or 10.5)))
    line_spacing = min(3.0, max(1.0, float(args.get("line_spacing") or 1.5)))
    first_indent_chars = min(8.0, max(0.0, float(args.get("first_line_indent_chars") or 2.0)))
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Pt

        document = Document(str(source))
    except Exception as exc:
        raise ToolError(f"无法打开 Word 文件：{exc}") from exc

    before_signature, before_chars = _docx_text_signature(source)
    for style_name in ("Normal", "正文"):
        if style_name not in document.styles:
            continue
        style = document.styles[style_name]
        style.font.name = latin
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
        style.font.size = Pt(body_size)

    paragraph_count = 0
    run_count = 0
    for paragraph in _docx_paragraphs(document):
        if not paragraph.text.strip():
            continue
        paragraph_count += 1
        style_name = str(paragraph.style.name if paragraph.style else "").lower()
        is_heading = style_name.startswith("heading") or "标题" in style_name
        in_table = paragraph._p.getparent().tag.rsplit("}", 1)[-1] == "tc"
        size = table_size if in_table else body_size
        heading_level = 0
        match = re.search(r"(?:heading|标题)\s*([1-9])", style_name)
        if is_heading:
            heading_level = int(match.group(1)) if match else 1
            size = max(body_size, body_size + max(1, 4 - heading_level) * 2)
        fmt = paragraph.paragraph_format
        fmt.line_spacing = line_spacing
        fmt.space_after = Pt(0)
        fmt.first_line_indent = None if is_heading or in_table else Pt(body_size * first_indent_chars)
        for run in paragraph.runs:
            _set_run_font(run, east_asia, latin, size, True if is_heading else None)
            run_count += 1

    requested_name = Path(str(args.get("output_name") or "")).name
    if not requested_name:
        requested_name = f"{source.stem}_格式规范化.docx"
    if not requested_name.lower().endswith(".docx"):
        requested_name += ".docx"
    safe_name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", requested_name).strip(" .")
    name = f"{Path(safe_name).stem[:120]}_{uuid.uuid4().hex[:8]}.docx"
    target = (EXPORT_DIR / name).resolve()
    target.relative_to(EXPORT_DIR.resolve())
    document.save(str(target))
    after_signature, after_chars = _docx_text_signature(target)
    if after_signature != before_signature or after_chars != before_chars:
        target.unlink(missing_ok=True)
        raise ToolError("内容不变校验失败，已拒绝导出并删除临时结果")
    return _json_result(
        file=name,
        download_url=f"/api/v1/exports/{name}",
        source=_relative(ctx, source),
        paragraphs_formatted=paragraph_count,
        runs_formatted=run_count,
        content_unchanged=True,
        content_sha256=after_signature,
    )


def _local_editorial_available(ctx: BuiltinToolContext) -> bool:
    if not (
        "photo-abstract-editorial" in ctx.active_skill_names
        and bool(ctx.attachment_images)
        and bool(ctx.llm)
        and callable(getattr(ctx.llm, "vision", None))
    ):
        return False
    # LLMClient 始终暴露 vision() 方法；只有 vision_model 非空时才会真正
    # 发送图片。不能把“方法存在”误判成“提供商已启用图片输入”。其他实现
    # （例如 ChatGPTClient 或测试桩）没有 vision_model 属性，仍由其自身负责能力。
    vision_fallback = getattr(ctx.llm, "vision_fallback", None)
    if (
        hasattr(ctx.llm, "vision_model")
        and not str(getattr(ctx.llm, "vision_model", "") or "").strip()
        and vision_fallback is None
    ):
        return False
    return (
        vision_fallback is not None
        or
        not hasattr(ctx.llm, "model_input")
        or "image" in set(getattr(ctx.llm, "model_input", []) or [])
    )


def _editorial_panel_provider_available(ctx: BuiltinToolContext) -> bool:
    """图片编辑服务必须显式配置，不能从文生图配置猜测能力。"""
    return bool(
        _local_editorial_available(ctx)
        and settings.IMAGE_EDIT_API_BASE_URL
        and settings.IMAGE_EDIT_API_KEY
        and settings.IMAGE_EDIT_MODEL
    )


async def _generate_editorial_panel(
    spec: dict[str, Any],
    source: Path,
    user_prompt: str,
) -> dict[str, Any]:
    """让图片编辑模型只生成下半部抽象面板，原照片仍由本地精确合成。"""
    url = settings.IMAGE_EDIT_API_BASE_URL.rstrip("/") + "/images/edits"
    await asyncio.to_thread(validate_outbound_url, url)
    content_type = mimetypes.guess_type(source.name)[0] or "image/png"
    direction = (
        "Using the supplied photograph only as visual reference, create ONLY a landscape abstract "
        "editorial panel on a warm ivory ground. Do not reproduce the photograph and do not add text, "
        "letters, logos, borders or a second photo. Translate its mass hierarchy, negative space, "
        "occlusion, asymmetry and source colors into refined filled shapes with restrained continuous "
        "tonal transitions. Avoid generic icons and repeated outline symbols. "
        f"Layout specification: {json.dumps(spec, ensure_ascii=False)}. "
        f"User direction: {str(user_prompt or '')[:1200]}"
    )
    async with httpx.AsyncClient(timeout=180, follow_redirects=False) as client:
        response = await client.post(
            url,
            headers={"Authorization": f"Bearer {settings.IMAGE_EDIT_API_KEY}"},
            data={
                "model": settings.IMAGE_EDIT_MODEL,
                "prompt": direction[:4000],
                "size": "1536x1024",
                "response_format": "b64_json",
            },
            files={"image": (source.name, source.read_bytes(), content_type)},
        )
    if response.status_code >= 400:
        raise RuntimeError(f"图片编辑服务返回 HTTP {response.status_code}")
    payload = response.json()
    item = (payload.get("data") or [{}])[0]
    encoded = item.get("b64_json")
    remote_url = item.get("url")
    if encoded:
        try:
            image_bytes = base64.b64decode(str(encoded), validate=True)
        except (ValueError, TypeError) as exc:
            raise RuntimeError("图片编辑服务返回了无效 Base64") from exc
    elif remote_url:
        await asyncio.to_thread(validate_outbound_url, remote_url)
        async with httpx.AsyncClient(timeout=60, follow_redirects=False) as client:
            image_response = await client.get(remote_url)
            image_response.raise_for_status()
        image_bytes = image_response.content
    else:
        raise RuntimeError("图片编辑服务没有返回 b64_json 或 URL")
    return {"data": image_bytes, "model": settings.IMAGE_EDIT_MODEL}


def _image_preflight(ctx: BuiltinToolContext) -> None:
    if settings.IMAGE_API_BASE_URL and settings.IMAGE_API_KEY and settings.IMAGE_MODEL:
        return
    if _local_editorial_available(ctx):
        return
    raise ToolError(
        "图片服务未配置；请设置 IMAGE_API_BASE_URL、IMAGE_API_KEY、IMAGE_MODEL，"
        "或启用支持本地成品渲染的图片 Skill",
        code="image_provider_not_configured",
    )


async def _image_generate(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "生成图片")
    _image_preflight(ctx)
    prompt = str(args.get("prompt") or "").strip()
    if not prompt:
        raise ToolError("prompt 不能为空")
    size = str(args.get("size") or "1024x1024")
    action = str(args.get("action") or "auto")
    use_input_images = bool(args.get("use_input_images", True))
    input_images = list(ctx.attachment_images or []) if use_input_images else []
    if _local_editorial_available(ctx):
        from ..artifacts import save_raster_artifact
        from ..photo_editorial import render_photo_abstract_editorial

        try:
            generated = await render_photo_abstract_editorial(
                ctx.llm,
                input_images[0],
                prompt,
                panel_generator=(
                    _generate_editorial_panel
                    if _editorial_panel_provider_available(ctx) else None
                ),
            )
            artifact = save_raster_artifact(generated["data"], prefix="editorial")
        except Exception as exc:  # noqa: BLE001 - 模型/渲染错误转换为稳定工具错误
            raise ToolError(
                f"照片抽象编辑本地成品渲染失败：{exc}",
                code="editorial_render_failed",
            ) from exc
        return _json_result(
            **artifact,
            download_url=f"/api/v1/exports/{artifact['file']}",
            source=generated["source"],
            model=generated.get("model"),
            action=generated["action"],
            input_images=generated["input_images"],
            renderer_version=generated["renderer_version"],
            design_source=generated["design_source"],
            panel_renderer=generated["panel_renderer"],
            panel_model=generated["panel_model"],
            panel_provider_error=generated["panel_provider_error"],
            quality_checks=generated["quality_checks"],
            design=generated["design"],
        )

    if input_images or action == "edit":
        raise ToolError(
            "独立 IMAGE_* 图片服务只支持文生图，不能处理上传图片；"
            "请启用支持图片编辑的 Skill",
            code="image_edit_capability_unavailable",
        )
    url = settings.IMAGE_API_BASE_URL.rstrip("/") + "/images/generations"
    await asyncio.to_thread(validate_outbound_url, url)
    async with httpx.AsyncClient(timeout=120, follow_redirects=False) as client:
        response = await client.post(
            url,
            headers={"Authorization": f"Bearer {settings.IMAGE_API_KEY}"},
            json={
                "model": settings.IMAGE_MODEL,
                "prompt": prompt[:4000],
                "size": size,
                "response_format": "b64_json",
            },
        )
    if response.status_code >= 400:
        raise ToolError(f"图片生成失败：HTTP {response.status_code}")
    payload = response.json()
    item = (payload.get("data") or [{}])[0]
    encoded = item.get("b64_json")
    remote_url = item.get("url")
    if encoded:
        try:
            image_bytes = base64.b64decode(str(encoded), validate=True)
        except (ValueError, TypeError) as exc:
            raise ToolError("图片接口返回了无效 Base64") from exc
    elif remote_url:
        await asyncio.to_thread(validate_outbound_url, remote_url)
        async with httpx.AsyncClient(timeout=60, follow_redirects=False) as client:
            image_response = await client.get(remote_url)
            image_response.raise_for_status()
        image_bytes = image_response.content
    else:
        raise ToolError("图片接口没有返回 b64_json 或 URL")
    from ..artifacts import save_raster_artifact

    artifact = save_raster_artifact(image_bytes, prefix="image")
    return _json_result(
        **artifact,
        download_url=f"/api/v1/exports/{artifact['file']}",
        source="environment_images_api",
        model=settings.IMAGE_MODEL,
        action="generate",
        input_images=0,
    )


async def _image_render(args: dict, ctx: BuiltinToolContext):
    """无外部模型依赖地保存安全 SVG 图片产物。"""
    _confirmed(args, "生成 SVG 图片")
    svg = str(args.get("svg") or "").strip()
    if not re.search(r"<svg\b", svg, re.IGNORECASE):
        raise ToolError("svg 必须包含 <svg> 根元素")
    forbidden = [
        r"<script\b", r"<foreignObject\b", r"\son[a-z]+\s*=",
        r"\b(?:href|src)\s*=\s*[\"'](?:https?:|file:|javascript:|data:)",
    ]
    if any(re.search(pattern, svg, re.IGNORECASE) for pattern in forbidden):
        raise ToolError("SVG 包含脚本、事件、外部资源或嵌入数据，已拒绝")
    title = re.sub(r'[\\/:*?"<>|]', "_", str(args.get("title") or "image"))[:60]
    name = f"{title or 'image'}_{uuid.uuid4().hex[:8]}.svg"
    target = EXPORT_DIR / name
    _atomic_write(target, svg)
    return _json_result(file=name, path=str(target), download_url=f"/api/v1/exports/{name}")


_BROWSER_SESSIONS: dict[str, dict] = {}
_LINK_RE = re.compile(
    r"<a\b[^>]*href=[\"']([^\"']+)[\"'][^>]*>(.*?)</a>",
    re.IGNORECASE | re.DOTALL,
)
_TAG_RE = re.compile(r"<[^>]+>")


def _browser_snapshot(session: dict) -> dict:
    content = session.get("content", "")
    links = []
    for index, (url, label) in enumerate(_LINK_RE.findall(content)[:100]):
        links.append({
            "id": index,
            "url": url,
            "label": html.unescape(_TAG_RE.sub("", label)).strip()[:200],
        })
    text = html.unescape(_TAG_RE.sub(" ", content))
    text = re.sub(r"\s+", " ", text).strip()
    text, truncated = _limit(text, 16000)
    return {
        "session_id": session["id"],
        "url": session["url"],
        "status_code": session.get("status_code"),
        "title": session.get("title", ""),
        "text": text,
        "links": links,
        "truncated": truncated,
    }


async def _browser_fetch(url: str) -> dict:
    await asyncio.to_thread(validate_outbound_url, url)
    async with httpx.AsyncClient(
        timeout=settings.BROWSER_TIMEOUT_SECONDS,
        follow_redirects=True,
        headers={"User-Agent": "HarnessAgent/2.0"},
    ) as client:
        response = await client.get(url)
    content_type = response.headers.get("content-type", "")
    if "text/" not in content_type and "json" not in content_type and "xml" not in content_type:
        raise ToolError(f"浏览器只支持文本页面，收到 {content_type or '未知类型'}")
    content = response.text
    title_match = re.search(r"<title[^>]*>(.*?)</title>", content, re.I | re.S)
    return {
        "id": uuid.uuid4().hex[:16],
        "url": str(response.url),
        "status_code": response.status_code,
        "title": html.unescape(_TAG_RE.sub("", title_match.group(1))).strip()
        if title_match else "",
        "content": content,
    }


async def _browser_open(args: dict, ctx: BuiltinToolContext):
    from . import browser_cdp
    try:
        session = await browser_cdp.open_page(
            str(args.get("url") or ""),
            f"{ctx.user_id}:{ctx.run_id}",
        )
        return _json_result(**(await browser_cdp.snapshot(session)))
    except browser_cdp.BrowserError as exc:
        raise ToolError(str(exc)) from exc


async def _browser_snapshot_tool(args: dict, ctx: BuiltinToolContext):
    from . import browser_cdp
    try:
        session = browser_cdp.get(
            str(args.get("session_id") or ""), f"{ctx.user_id}:{ctx.run_id}"
        )
        return _json_result(**(await browser_cdp.snapshot(session)))
    except browser_cdp.BrowserError as exc:
        raise ToolError(str(exc)) from exc


async def _browser_click(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "浏览器点击")
    from . import browser_cdp
    try:
        session = browser_cdp.get(
            str(args.get("session_id") or ""), f"{ctx.user_id}:{ctx.run_id}"
        )
        result = await browser_cdp.click(session, str(args.get("element_id") or ""))
        return _json_result(**result)
    except browser_cdp.BrowserError as exc:
        raise ToolError(str(exc)) from exc


async def _browser_type(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "浏览器输入")
    from . import browser_cdp
    try:
        session = browser_cdp.get(
            str(args.get("session_id") or ""), f"{ctx.user_id}:{ctx.run_id}"
        )
        result = await browser_cdp.type_text(
            session,
            str(args.get("element_id") or ""),
            str(args.get("text") or ""),
            bool(args.get("submit")),
        )
        return _json_result(**result)
    except browser_cdp.BrowserError as exc:
        raise ToolError(str(exc)) from exc


async def _browser_screenshot(args: dict, ctx: BuiltinToolContext):
    from . import browser_cdp
    try:
        session = browser_cdp.get(
            str(args.get("session_id") or ""), f"{ctx.user_id}:{ctx.run_id}"
        )
        return _json_result(**(await browser_cdp.screenshot(session)))
    except browser_cdp.BrowserError as exc:
        raise ToolError(str(exc)) from exc


async def _browser_close(args: dict, ctx: BuiltinToolContext):
    from . import browser_cdp
    try:
        session = browser_cdp.get(
            str(args.get("session_id") or ""), f"{ctx.user_id}:{ctx.run_id}"
        )
        await browser_cdp.close(session)
        return _json_result(closed=True, session_id=session.id)
    except browser_cdp.BrowserError as exc:
        raise ToolError(str(exc)) from exc


def _require_identity(ctx: BuiltinToolContext) -> tuple[int, int]:
    if not ctx.user_id or not ctx.agent_id:
        raise ToolError("当前运行缺少用户或智能体身份，不能管理持久化任务")
    return int(ctx.user_id), int(ctx.agent_id)


async def _cron_create(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "创建 Cron 任务")
    owner_id, agent_id = _require_identity(ctx)
    await asyncio.to_thread(_require_schedule_permission, owner_id)
    from .. import scheduler
    try:
        task = await asyncio.to_thread(
            scheduler.create,
            owner_id,
            agent_id,
            str(args.get("name") or "定时任务"),
            str(args.get("cron") or ""),
            str(args.get("timezone") or "Asia/Shanghai"),
            str(args.get("query") or ""),
            ctx.session_id,
            True,
        )
    except ValueError as exc:
        raise ToolError(str(exc)) from exc
    return _json_result(task=task)


async def _cron_list(args: dict, ctx: BuiltinToolContext):
    owner_id, agent_id = _require_identity(ctx)
    await asyncio.to_thread(_require_schedule_permission, owner_id)
    from .. import scheduler
    tasks = await asyncio.to_thread(scheduler.list_tasks, owner_id, agent_id)
    return _json_result(tasks=tasks)


async def _cron_delete(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "删除 Cron 任务")
    owner_id, _ = _require_identity(ctx)
    await asyncio.to_thread(_require_schedule_permission, owner_id)
    from .. import scheduler
    deleted = await asyncio.to_thread(
        scheduler.delete, owner_id, str(args.get("task_id") or "")
    )
    if not deleted:
        raise ToolError("定时任务不存在或不属于当前用户")
    return _json_result(deleted=True, task_id=args.get("task_id"))


async def _cron_set_enabled(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "修改 Cron 任务")
    owner_id, _ = _require_identity(ctx)
    await asyncio.to_thread(_require_schedule_permission, owner_id)
    from .. import scheduler
    task = await asyncio.to_thread(
        scheduler.set_enabled,
        owner_id,
        str(args.get("task_id") or ""),
        bool(args.get("enabled")),
    )
    if task is None:
        raise ToolError("定时任务不存在或不属于当前用户")
    return _json_result(task=task)


def _subagent_payload(
    user_id: int, target_agent_id: int, query: str, session_id: str, depth: int,
    approval_policy: str = DEFAULT_APPROVAL_POLICY,
) -> dict:
    from ..database import SessionLocal
    from ..models import Agent
    from ..api.chat import build_execution_snapshot
    db = SessionLocal()
    try:
        agent = db.get(Agent, target_agent_id)
        if agent is None or not agent.enabled:
            raise ToolError("目标子智能体不存在或已停用")
        version = agent.active_version
        execution_snapshot = build_execution_snapshot(
            db, agent, query, approval_policy=approval_policy
        )
    finally:
        db.close()
    return {
        "user_id": user_id,
        "agent_id": target_agent_id,
        "harness_version": version,
        "inputs": {
            "query": query,
            "project_name": "", "city_name": "", "project_address": "",
            "project_info": "", "industry_structure": "", "electricity_trading": "",
            "image_scale": "", "satellite_images": [], "drawing_images": [],
            "bill_files": [], "documents": [], "custom_vars": {},
            "custom_files": {}, "custom_var_labels": {},
        },
        "template_ids": [], "dataset_ids": [], "skill_ids": [], "mcp_ids": [],
        "invoked_agent_ids": [], "provider_id": None,
        "attachment_images": [], "attachment_docs": [],
        "session_id": session_id,
        "source": "subagent",
        "subagent_depth": depth,
        "approval_policy": normalize_approval_policy(approval_policy),
        "execution_snapshot": execution_snapshot,
        "delegation_authorized": True,
    }


def _require_schedule_permission(user_id: int) -> None:
    from ..database import SessionLocal
    from ..models import User
    from ..security import has_module_access

    db = SessionLocal()
    try:
        user = db.get(User, user_id)
        if user is None or not user.is_active or not has_module_access(user, "schedules"):
            raise ToolError("当前用户没有定时任务模块权限")
    finally:
        db.close()


def _allowed_subagent_ids(ctx: BuiltinToolContext) -> set[int]:
    values = {
        int(item["id"]) for item in (ctx.sub_agents or [])
        if isinstance(item, dict) and item.get("id") is not None
    }
    # 同一智能体可以作为隔离的后台 Worker 执行子任务；工具预算和任务队列负责限流。
    if ctx.agent_id:
        values.add(int(ctx.agent_id))
    return values


async def _spawn_agent(args: dict, ctx: BuiltinToolContext):
    owner_id, current_agent_id = _require_identity(ctx)
    target = int(args.get("agent_id") or current_agent_id)
    if target not in _allowed_subagent_ids(ctx):
        raise ToolError("只能调用当前智能体或已显式绑定的子智能体")
    query = str(args.get("query") or "").strip()
    if not query:
        raise ToolError("query 不能为空")
    from .. import jobs
    if ctx.subagent_depth >= 3:
        raise ToolError("子智能体嵌套深度已达到上限 3")
    session_id = uuid.uuid4().hex
    payload = await asyncio.to_thread(
        _subagent_payload,
        owner_id,
        target,
        query,
        session_id,
        ctx.subagent_depth + 1,
        ctx.approval_policy,
    )
    payload["parent_agent_id"] = current_agent_id
    payload["parent_run_id"] = str(ctx.run_id or "") or None
    job_id = await asyncio.to_thread(jobs.enqueue, owner_id, target, "chat", payload)
    return _json_result(
        task_id=job_id, status=jobs.PENDING, agent_id=target, session_id=session_id
    )


def _job_summary(view) -> dict:
    result = view.result or {}
    completion_status = (
        str(result.get("completion_status") or "completed")
        if view.status == "done" else ""
    )
    raw_issues = result.get("completion_issues")
    raw_issues = raw_issues if isinstance(raw_issues, (list, tuple)) else []
    raw_summary = result.get("plan_summary")
    raw_summary = raw_summary if isinstance(raw_summary, dict) else {}
    return {
        "task_id": view.id,
        "agent_id": view.agent_id,
        "status": view.status,
        "completion_status": completion_status,
        "completion_issues": (
            [str(value)[:400] for value in raw_issues[:20]]
            if view.status == "done" else []
        ),
        "plan_summary": (
            dict(raw_summary)
            if view.status == "done" else {}
        ),
        "progress": view.progress,
        "answer": result.get("answer") if view.status == "done" else None,
        "error": view.error,
        "error_class": view.error_class,
        "session_id": view.payload.get("session_id"),
    }


async def _wait_agent(args: dict, ctx: BuiltinToolContext):
    owner_id, _ = _require_identity(ctx)
    from .. import jobs
    task_id = str(args.get("task_id") or "")
    parent_scope = str(ctx.run_id or "") or None
    timeout = min(max(float(args.get("timeout_seconds") or 10), 0), 30)
    deadline = asyncio.get_running_loop().time() + timeout
    while True:
        view = await asyncio.to_thread(jobs.view, task_id, owner_id)
        if (
            view is None
            or view.payload.get("source") != "subagent"
            or view.parent_job_id != parent_scope
        ):
            raise ToolError("子智能体任务不存在")
        if view.status in {jobs.DONE, jobs.FAILED, jobs.CANCELLED, jobs.DEAD_LETTER}:
            return _json_result(task=_job_summary(view))
        if asyncio.get_running_loop().time() >= deadline:
            return _json_result(task=_job_summary(view), timed_out=True)
        await asyncio.sleep(.25)


async def _list_agents(args: dict, ctx: BuiltinToolContext):
    owner_id, _ = _require_identity(ctx)
    from ..database import SessionLocal
    from ..models import Job
    from .. import jobs
    db = SessionLocal()
    try:
        rows = (
            db.query(Job)
            .filter(
                Job.owner_id == owner_id,
                Job.kind == "chat",
                Job.parent_job_id == (str(ctx.run_id or "") or None),
            )
            .order_by(Job.created_at.desc())
            .limit(min(max(int(args.get("limit") or 20), 1), 100))
            .all()
        )
        summaries = []
        for row in rows:
            view = jobs._view(row)  # 同模块的稳定只读转换；避免泄露 ORM 对象
            if view.payload.get("source") == "subagent":
                summaries.append(_job_summary(view))
    finally:
        db.close()
    return _json_result(tasks=summaries)


async def _resume_agent(args: dict, ctx: BuiltinToolContext):
    owner_id, current_agent_id = _require_identity(ctx)
    from .. import jobs
    previous = await asyncio.to_thread(
        jobs.view, str(args.get("task_id") or ""), owner_id
    )
    if (
        previous is None
        or previous.payload.get("source") != "subagent"
        or previous.parent_job_id != (str(ctx.run_id or "") or None)
    ):
        raise ToolError("子智能体任务不存在")
    target = int(previous.agent_id or current_agent_id)
    if target not in _allowed_subagent_ids(ctx):
        raise ToolError("该子智能体不再允许调用")
    query = str(args.get("query") or "").strip()
    if not query:
        raise ToolError("query 不能为空")
    session_id = previous.payload.get("session_id") or uuid.uuid4().hex
    payload = await asyncio.to_thread(
        _subagent_payload,
        owner_id,
        target,
        query,
        session_id,
        int(previous.payload.get("subagent_depth") or 1),
    )
    payload["parent_agent_id"] = current_agent_id
    payload["parent_run_id"] = str(ctx.run_id or "") or None
    payload["resumed_from"] = previous.id
    job_id = await asyncio.to_thread(jobs.enqueue, owner_id, target, "chat", payload)
    return _json_result(
        task_id=job_id, resumed_from=previous.id, status=jobs.PENDING,
        agent_id=target, session_id=session_id,
    )


async def _interrupt_agent(args: dict, ctx: BuiltinToolContext):
    _confirmed(args, "中断子智能体")
    owner_id, _ = _require_identity(ctx)
    from .. import jobs
    task_id = str(args.get("task_id") or "")
    view = await asyncio.to_thread(jobs.view, task_id, owner_id)
    if (
        view is None
        or view.payload.get("source") != "subagent"
        or view.parent_job_id != (str(ctx.run_id or "") or None)
    ):
        raise ToolError("子智能体任务不存在")
    cancelled = await jobs.request_cancel(task_id, owner_id)
    return _json_result(task_id=task_id, interrupted=cancelled)


async def _close_agent(args: dict, ctx: BuiltinToolContext):
    # 持久任务不删除审计记录；close 的语义是停止尚未结束的执行并关闭生命周期。
    result = await _interrupt_agent(args, ctx)
    parsed = json.loads(result)
    parsed["closed"] = parsed.pop("interrupted", False)
    return json.dumps(parsed, ensure_ascii=False)


def _specs() -> list[BuiltinTool]:
    path = {"type": "string", "description": "相对工作区路径"}
    confirm = {"type": "boolean", "description": "确认执行有副作用操作，必须为 true"}
    edit_item = _object({
        "path": path,
        "old": {"type": "string"},
        "new": {"type": "string"},
        "expected_replacements": {"type": "integer", "default": 1, "minimum": 1},
    }, ["path", "old", "new"])
    patch_operation = _object({
        "old": {"type": "string"},
        "new": {"type": "string"},
        "expected_replacements": {"type": "integer", "default": 1, "minimum": 1},
    }, ["old", "new"])
    return [
        BuiltinTool("ls", "列出工作区目录内容。", _object({
            "path": path, "limit": {"type": "integer", "default": 200},
        }), _workspace_ls, "filesystem"),
        BuiltinTool("glob", "按 glob 模式发现工作区文件。", _object({
            "pattern": {"type": "string"}, "path": path,
            "limit": {"type": "integer", "default": 500},
        }, ["pattern"]), _workspace_glob, "filesystem"),
        BuiltinTool("grep", "在工作区文本文件中搜索字符串或正则表达式。", _object({
            "query": {"type": "string"}, "path": path,
            "regex": {"type": "boolean", "default": False},
            "case_sensitive": {"type": "boolean", "default": False},
            "recursive": {"type": "boolean", "default": True},
            "include": {"type": "string", "default": "*"},
            "max_results": {"type": "integer", "default": 100},
        }, ["query"]), _workspace_grep, "filesystem"),
        BuiltinTool("read", "读取工作区文本文件，可指定行范围。", _object({
            "path": path, "start_line": {"type": "integer", "default": 1},
            "end_line": {"type": "integer"},
        }, ["path"]), _workspace_read, "filesystem"),
        BuiltinTool("read_many", "一次读取多个工作区文本文件。", _object({
            "paths": {"type": "array", "items": path, "maxItems": 20},
        }, ["paths"]), _workspace_read_many, "filesystem"),
        BuiltinTool("write", "创建全新文本文件；覆盖时必须显式声明。", _object({
            "path": path, "content": {"type": "string"},
            "overwrite": {"type": "boolean", "default": False}, "confirm": confirm,
        }, ["path", "content", "confirm"]), _workspace_write, "filesystem", True),
        BuiltinTool("edit", "对单个文件执行可验证的精确字符串替换。", _object({
            "path": path, "old": {"type": "string"}, "new": {"type": "string"},
            "expected_replacements": {"type": "integer", "default": 1},
            "confirm": confirm,
        }, ["path", "old", "new", "confirm"]), _workspace_edit, "filesystem", True),
        BuiltinTool("multi_edit", "原子执行多项精确字符串替换；任一失败则不写入。", _object({
            "edits": {"type": "array", "items": edit_item, "maxItems": 50},
            "confirm": confirm,
        }, ["edits", "confirm"]), _workspace_multi_edit, "filesystem", True),
        BuiltinTool("apply_patch", "应用结构化多文件补丁，支持 add/update，不允许删除。", _object({
            "files": {
                "type": "array",
                "items": _object({
                    "path": path,
                    "action": {"type": "string", "enum": ["add", "update"]},
                    "content": {"type": "string"},
                    "operations": {"type": "array", "items": patch_operation},
                }, ["path", "action"]),
            },
            "confirm": confirm,
        }, ["files", "confirm"]), _workspace_apply_patch, "filesystem", True),
        BuiltinTool("shell", "在工作区运行一条受限 Bash 或 PowerShell 白名单命令。", _object({
            "shell": {"type": "string", "enum": ["powershell", "bash"]},
            "script": {"type": "string"},
            "timeout_seconds": {"type": "integer", "minimum": 1, "maximum": 60},
            "confirm": confirm,
        }, ["shell", "script", "confirm"]), _shell_run, "shell", True),
        BuiltinTool("git_status", "查看当前工作区 Git 分支和变更状态。", _object({}), _git_status, "git"),
        BuiltinTool("git_diff", "查看 Git 工作区或暂存区差异。", _object({
            "staged": {"type": "boolean", "default": False}, "path": path,
        }), _git_diff, "git"),
        BuiltinTool("git_commit", "仅暂存明确 paths 并创建本地 Git 提交；不推送远程。", _object({
            "paths": {"type": "array", "items": path},
            "message": {"type": "string"}, "confirm": confirm,
        }, ["paths", "message", "confirm"]), _git_commit, "git", True),
        BuiltinTool("lsp", "执行代码诊断或查询定义、引用和符号；缺少语言服务器时确定性降级。", _object({
            "action": {"type": "string", "enum": [
                "diagnostics", "definition", "references", "symbols"
            ]},
            "path": path, "symbol": {"type": "string"},
        }, ["action"]), _lsp, "lsp"),
        BuiltinTool(
            "web_search",
            "免费联网搜索。优先使用自托管 SearXNG；不可用时聚合 DuckDuckGo 与 Bing 并去重排序。",
            _object({
                "query": {"type": "string"},
                "count": {"type": "integer", "minimum": 1, "maximum": 20, "default": 8},
                "language": {"type": "string", "default": "zh-CN"},
                "time_range": {
                    "type": "string",
                    "enum": ["", "day", "month", "year"],
                    "default": "",
                },
            }, ["query"]),
            _web_search,
            "web",
        ),
        BuiltinTool(
            "web_fetch",
            "读取指定公开网页正文；逐次校验重定向以阻止 SSRF，不执行页面脚本。",
            _object({"url": {"type": "string"}}, ["url"]),
            _web_fetch,
            "web",
        ),
        BuiltinTool("html_generate", "生成可下载的独立 HTML 产物。", _object({
            "title": {"type": "string"}, "html": {"type": "string"}, "confirm": confirm,
        }, ["title", "html", "confirm"]), _html_artifact, "artifact", True),
        BuiltinTool(
            "document_inspect",
            "检查工作区中的 Word DOCX 原文件结构、样式和内容指纹；处理附件、Word、排版任务时使用。",
            _object({"path": path}, ["path"]),
            _document_inspect,
            "document",
        ),
        BuiltinTool(
            "document_create",
            "把已识别、整理或生成的文字与结构化表格创建为新的可下载 Word DOCX；图片 OCR 转 Word、从文本生成 Word 时必须使用。",
            _object({
                "title": {"type": "string"},
                "content": {"type": "string", "description": "正文；支持 # 标题、- 项目符号和数字列表"},
                "tables": {
                    "type": "array",
                    "maxItems": 20,
                    "items": _object({
                        "headers": {"type": "array", "items": {"type": "string"}, "maxItems": 50},
                        "rows": {
                            "type": "array",
                            "maxItems": 2000,
                            "items": {"type": "array", "items": {"type": "string"}, "maxItems": 50},
                        },
                    }),
                },
                "output_name": {"type": "string", "default": "识别结果.docx"},
                "confirm": confirm,
            }, ["confirm"]),
            _document_create,
            "document",
            True,
        ),
        BuiltinTool(
            "document_format",
            "在严格保持文字内容不变的前提下规范 Word DOCX 的字体、字号、段落、行距和表格文字，并导出可下载文件。",
            _object({
                "path": path,
                "output_name": {"type": "string"},
                "east_asia_font": {"type": "string", "default": "宋体"},
                "latin_font": {"type": "string", "default": "Times New Roman"},
                "body_size_pt": {"type": "number", "default": 12},
                "table_size_pt": {"type": "number", "default": 10.5},
                "line_spacing": {"type": "number", "default": 1.5},
                "first_line_indent_chars": {"type": "number", "default": 2},
                "confirm": confirm,
            }, ["path", "confirm"]),
            _document_format,
            "document",
            True,
        ),
        BuiltinTool("image_generate", "优先使用当前模型提供商的原生图片能力生成或编辑图片，并保留独立 IMAGE_* 回退。", _object({
            "prompt": {"type": "string"},
            "size": {"type": "string", "enum": ["1024x1024", "1536x1024", "1024x1536"]},
            "action": {"type": "string", "enum": ["auto", "generate", "edit"], "default": "auto"},
            "use_input_images": {"type": "boolean", "default": True},
            "confirm": confirm,
        }, ["prompt", "confirm"]), _image_generate, "artifact", True),
        BuiltinTool("image_render", "将安全 SVG 源码保存为可下载图片，无需外部图片模型。", _object({
            "title": {"type": "string"}, "svg": {"type": "string"}, "confirm": confirm,
        }, ["title", "svg", "confirm"]), _image_render, "artifact", True),
        BuiltinTool("browser_open", "在隔离的无痕 Chrome/Edge 中打开网页并返回可交互元素。", _object({
            "url": {"type": "string"},
        }, ["url"]), _browser_open, "browser"),
        BuiltinTool("browser_snapshot", "读取浏览器当前页面正文与交互元素快照。", _object({
            "session_id": {"type": "string"},
        }, ["session_id"]), _browser_snapshot_tool, "browser"),
        BuiltinTool("browser_click", "按 snapshot 返回的 element_id 点击网页元素。", _object({
            "session_id": {"type": "string"}, "element_id": {"type": "string"},
            "confirm": confirm,
        }, ["session_id", "element_id", "confirm"]), _browser_click, "browser", True),
        BuiltinTool("browser_type", "向网页输入控件写入文本，可选择提交表单。", _object({
            "session_id": {"type": "string"}, "element_id": {"type": "string"},
            "text": {"type": "string"}, "submit": {"type": "boolean", "default": False},
            "confirm": confirm,
        }, ["session_id", "element_id", "text", "confirm"]), _browser_type, "browser", True),
        BuiltinTool("browser_screenshot", "保存浏览器当前视口截图为图片产物。", _object({
            "session_id": {"type": "string"},
        }, ["session_id"]), _browser_screenshot, "browser"),
        BuiltinTool("browser_close", "关闭隔离浏览器会话并清理临时资料。", _object({
            "session_id": {"type": "string"},
        }, ["session_id"]), _browser_close, "browser"),
        BuiltinTool("CronCreate", "创建当前智能体的持久化 Cron 定时任务。", _object({
            "name": {"type": "string"}, "cron": {"type": "string"},
            "timezone": {"type": "string", "default": "Asia/Shanghai"},
            "query": {"type": "string"}, "confirm": confirm,
        }, ["name", "cron", "query", "confirm"]), _cron_create, "scheduler", True),
        BuiltinTool("CronList", "列出当前用户为此智能体创建的 Cron 任务。", _object({}),
                    _cron_list, "scheduler"),
        BuiltinTool("CronDelete", "删除当前用户创建的 Cron 任务。", _object({
            "task_id": {"type": "string"}, "confirm": confirm,
        }, ["task_id", "confirm"]), _cron_delete, "scheduler", True),
        BuiltinTool("CronSetEnabled", "暂停或恢复 Cron 任务。", _object({
            "task_id": {"type": "string"}, "enabled": {"type": "boolean"},
            "confirm": confirm,
        }, ["task_id", "enabled", "confirm"]), _cron_set_enabled, "scheduler", True),
        BuiltinTool("spawn_agent", "启动当前智能体或已绑定子智能体的异步子任务。", _object({
            "agent_id": {"type": "integer"}, "query": {"type": "string"},
        }, ["query"]), _spawn_agent, "delegation", True),
        BuiltinTool("resume_agent", "在已有子智能体会话中追加任务并恢复执行。", _object({
            "task_id": {"type": "string"}, "query": {"type": "string"},
        }, ["task_id", "query"]), _resume_agent, "delegation", True),
        BuiltinTool("wait_agent", "等待子智能体任务终态，最长等待 30 秒。", _object({
            "task_id": {"type": "string"},
            "timeout_seconds": {"type": "number", "minimum": 0, "maximum": 30},
        }, ["task_id"]), _wait_agent, "delegation"),
        BuiltinTool("list_agents", "列出当前用户最近的子智能体任务和状态。", _object({
            "limit": {"type": "integer", "default": 20, "minimum": 1, "maximum": 100},
        }), _list_agents, "delegation"),
        BuiltinTool("interrupt_agent", "中断正在执行的子智能体任务。", _object({
            "task_id": {"type": "string"}, "confirm": confirm,
        }, ["task_id", "confirm"]), _interrupt_agent, "delegation", True),
        BuiltinTool("close_agent", "关闭子智能体任务；保留结果和审计记录。", _object({
            "task_id": {"type": "string"}, "confirm": confirm,
        }, ["task_id", "confirm"]), _close_agent, "delegation", True),
    ]


async def _service_list(args, ctx):
    from ..custom_services import service_definitions
    owner_id, agent_id = _require_identity(ctx)
    return _json_result(services=await asyncio.to_thread(service_definitions, owner_id, agent_id))


async def _service_call(args, ctx):
    from ..custom_services import execute_service
    owner_id, agent_id = _require_identity(ctx)
    if args.get("confirm") is not True:
        raise ToolError("运行服务需要 confirm=true")
    return _json_result(result=await execute_service(
        int(args["service_id"]), args.get("input", {}), owner_id, agent_id,
        run_id=ctx.run_id, approval_policy=ctx.approval_policy,
        approval_tokens=ctx.approval_tokens, runtime_event=ctx.runtime_event,
        provider_id=getattr(ctx.llm, "provider_id", None),
    ))


TOOLS: dict[str, BuiltinTool] = {item.name: item for item in _specs()}
TOOLS["presentation_create"] = BuiltinTool(
    "presentation_create", "将逐页标题与正文导出为原生文字可编辑 PPTX；无需 Shell。不复刻源图片或复杂版式。",
    _object({"title": {"type": "string"}, "output_name": {"type": "string"},
             "slides": {"type": "array", "minItems": 1, "maxItems": 80, "items": _object({
                 "title": {"type": "string", "maxLength": 48},
                 "body": {"type": "string", "maxLength": 480, "description": "最多10行，过长必须拆页"},
                 "source_pages": {"type": "array", "items": {"type": "integer", "minimum": 1},
                                  "description": "逐页重制时每页仅填写对应来源页号，例如第3页为[3]；须覆盖服务端指定的全部来源页"},
             }, ["title", "body"])}, "confirm": {"type": "boolean"}}, ["slides", "confirm"]),
    _presentation_create, "document", True,
)
TOOLS["service_list"] = BuiltinTool("service_list", "列出当前用户与智能体获准调用的自定义网络、编程服务和输入字段。", _object({}), _service_list, "services")
TOOLS["service_call"] = BuiltinTool("service_call", "按 service_list 返回的服务ID与字段运行已分配服务；可能修改外部数据，需批准。", _object({"service_id":{"type":"integer"},"input":{"type":"object"},"confirm":{"type":"boolean"}},["service_id","input","confirm"]), _service_call, "services", True)

# 公共智能体面对所有普通用户，不能继承创建者/服务账号对本机工作区的权限。
# 即使旧数据库仍绑定这些工具，运行时也会强制剔除，避免配置漂移重新打开攻击面。
PUBLIC_AGENT_DENIED_TOOLS = {
    name
    for name, tool in TOOLS.items()
    if tool.group in {"filesystem", "shell", "git"}
} | {
    "browser_click",
    "browser_type",
}
PUBLIC_AGENT_SAFE_TOOLS = set(TOOLS) - PUBLIC_AGENT_DENIED_TOOLS


def tool_specs(
    *, groups: set[str] | None = None, include_mutating: bool = True,
    enabled_names: set[str] | None = None,
) -> list[dict]:
    if not settings.BUILTIN_TOOLS_ENABLED:
        return []
    return [
        tool.spec()
        for tool in TOOLS.values()
        if (not groups or tool.group in groups)
        and (include_mutating or not tool.mutating)
        and (enabled_names is None or tool.name in enabled_names)
    ]


async def execute(name: str, args: dict, ctx: BuiltinToolContext) -> str:
    tool = TOOLS.get(name)
    if tool is None:
        raise ToolError(f"未知内置工具：{name}")
    if ctx.enabled_tools is not None and name not in ctx.enabled_tools:
        return _json_result(ok=False, error=f"内置工具未授权：{name}", tool=name)
    try:
        await enforce_content(
            "tool_input", args, user_id=ctx.user_id, agent_id=ctx.agent_id,
            provider_id=getattr(ctx.llm, "provider_id", None), runtime_event=ctx.runtime_event,
        )
        decision = await asyncio.to_thread(
            guardrails.runtime_decision,
            kind="builtin", tool_name=name, arguments=args,
            policy=ctx.approval_policy, mutating=tool.mutating,
        )
        if decision["decision"] == "block" or decision["guardrail_requires_approval"]:
            if ctx.runtime_event:
                event_value = ctx.runtime_event("guardrail.evaluated", guardrails.event_payload(
                    decision, agent_id=ctx.agent_id, execution_id=ctx.execution_id or ctx.run_id,
                ))
                if asyncio.iscoroutine(event_value):
                    await event_value
        from ..guardrail_reviews import review_tool
        decision = await review_tool(decision, ctx.user_id, ctx.agent_id, ctx.runtime_event)
        if decision["decision"] == "block":
            raise ToolError(decision["reason"], code="guardrail_blocked")
        if decision["guardrail_requires_approval"] and not ctx.run_id:
            raise ToolError("护栏要求单次审批，当前调用缺少可审批任务上下文", code="guardrail_approval_context_required")
        if name in {"image_generate", "image_render"}:
            ctx.required_artifact_kinds.add("image")
        if name in {"document_create", "document_format"}:
            ctx.required_artifact_kinds.add("document")
        if name == "presentation_create":
            ctx.required_artifact_kinds.add("presentation")
        # 能力缺失应在申请批准之前暴露，避免用户批准后整个 Agent Loop 重跑才失败。
        if name == "image_generate":
            _image_preflight(ctx)
        if tool.mutating and ctx.run_id and name != "service_call":
            selected_policy = normalize_approval_policy(ctx.approval_policy)
            risk = builtin_risk(name, args)
            approved_once = consume_approval(
                ctx.approval_tokens,
                run_id=str(ctx.run_id or ""),
                user_id=ctx.user_id,
                agent_id=ctx.agent_id,
                scope=name,
            )
            if not approved_once and decision["requires_approval"]:
                raise ApprovalRequired(
                    name,
                    tool.description,
                    agent_id=ctx.agent_id,
                    execution_context={
                        "execution_id": ctx.execution_id or ctx.run_id,
                        "parent_run_id": ctx.parent_run_id,
                        "agent_id": ctx.agent_id,
                        "subagent_depth": ctx.subagent_depth,
                        "approval_policy": selected_policy,
                        "risk": risk,
                    },
                )
            if not approved_once and ctx.runtime_event:
                event_value = ctx.runtime_event("approval.auto_approved", {
                    "scope": name,
                    "description": tool.description,
                    "policy": selected_policy,
                    "risk": risk,
                    "agent_id": ctx.agent_id,
                    "execution_id": ctx.execution_id or ctx.run_id,
                    "parent_run_id": ctx.parent_run_id,
                    "subagent_depth": ctx.subagent_depth,
                })
                if asyncio.iscoroutine(event_value):
                    await event_value
        value = tool.handler(args, ctx)
        if asyncio.iscoroutine(value):
            value = await value
        await enforce_content(
            "tool_output", value, user_id=ctx.user_id, agent_id=ctx.agent_id,
            provider_id=getattr(ctx.llm, "provider_id", None), runtime_event=ctx.runtime_event,
        )
        artifact_value = None
        if isinstance(value, str):
            try:
                artifact_value = json.loads(value)
            except json.JSONDecodeError:
                artifact_value = None
        elif isinstance(value, dict):
            artifact_value = value
        if isinstance(artifact_value, dict) and artifact_value.get("file"):
            filename = Path(str(artifact_value["file"])).name
            if ctx.artifact_callback:
                await ctx.artifact_callback(filename)
            if filename not in ctx.artifacts:
                ctx.artifacts.append(filename)
            if name in {"image_generate", "image_render"}:
                ctx.artifact_failures.pop("image", None)
                ctx.artifact_metadata["image"] = dict(artifact_value)
            if name in {"document_create", "document_format"}:
                ctx.artifact_failures.pop("document", None)
                ctx.artifact_metadata["document"] = dict(artifact_value)
        if isinstance(value, str):
            return value
        return _json_result(result=value)
    except ApprovalRequired:
        raise
    except ContentBlocked as exc:
        return _json_result(ok=False, error=str(exc), code=exc.code, retryable=False, tool=name)
    except ToolError as exc:
        if name in {"image_generate", "image_render"}:
            ctx.artifact_failures["image"] = f"{exc.code}：{exc}"
        if name in {"document_create", "document_format"}:
            ctx.artifact_failures["document"] = f"{exc.code}：{exc}"
        return _json_result(
            ok=False,
            error=str(exc),
            code=exc.code,
            retryable=exc.retryable,
            tool=name,
        )
    except Exception as exc:  # noqa: BLE001 - 对模型返回稳定错误，详细栈留给日志
        if name in {"image_generate", "image_render"}:
            ctx.artifact_failures["image"] = f"{type(exc).__name__}：{exc}"
        if name in {"document_create", "document_format"}:
            ctx.artifact_failures["document"] = f"{type(exc).__name__}：{exc}"
        return _json_result(ok=False, error=f"{type(exc).__name__}: {exc}", tool=name)


def completion_artifact_issues(ctx: BuiltinToolContext) -> list[str]:
    """把工具成功声明收紧为真实、可解码的产物证据。"""
    issues: list[str] = []
    if "image" in ctx.required_artifact_kinds:
        from ..artifacts import valid_image_artifact

        if any(valid_image_artifact(filename) for filename in ctx.artifacts):
            if "photo-abstract-editorial" in ctx.active_skill_names:
                metadata = ctx.artifact_metadata.get("image") or {}
                if metadata.get("design_source") != "vision_model":
                    issues.append("缺少任务所需图片产物：编辑作品没有经过有效视觉模型设计")
                elif metadata.get("renderer_version") != "photo_editorial_v3":
                    issues.append("缺少任务所需图片产物：编辑作品未使用当前质量渲染器")
                else:
                    checks = metadata.get("quality_checks") or {}
                    failed = [name for name, passed in checks.items() if not passed]
                    if failed or not checks:
                        issues.append(
                            "缺少任务所需图片产物：编辑作品质量检查未通过"
                            + ("（" + "、".join(failed) + "）" if failed else "")
                        )
        else:
            reason = str(ctx.artifact_failures.get("image") or "").strip()
            issues.append(
                f"缺少任务所需图片产物：图片工具失败（{reason}）"
                if reason else
                "缺少任务所需图片产物：图片工具已调用，但没有可解码的图片 Artifact"
            )

    if "document" in ctx.required_artifact_kinds:
        valid_document = False
        for filename in ctx.artifacts:
            path = (EXPORT_DIR / Path(str(filename)).name).resolve()
            try:
                path.relative_to(EXPORT_DIR.resolve())
                valid_document = (
                    path.is_file()
                    and path.suffix.lower() == ".docx"
                    and zipfile.is_zipfile(path)
                )
                if valid_document:
                    with zipfile.ZipFile(path) as archive:
                        valid_document = "word/document.xml" in archive.namelist()
            except (OSError, ValueError, zipfile.BadZipFile):
                valid_document = False
            if valid_document:
                break
        if not valid_document:
            reason = str(ctx.artifact_failures.get("document") or "").strip()
            issues.append(
                f"缺少任务所需 Word 产物：Word 工具失败（{reason}）"
                if reason else
                "缺少任务所需 Word 产物：没有生成可下载且有效的 DOCX Artifact"
            )
    if "presentation" in ctx.required_artifact_kinds:
        from ..artifacts import inspect_presentation_artifact
        from .presentation_requirements import inspect_presentation_requirements
        valid_presentation = False
        coverage_issues = []
        for name in ctx.artifacts:
            if not name.lower().endswith(".pptx") or not inspect_presentation_artifact(name).get("valid"):
                continue
            report = inspect_presentation_requirements(name, ctx.presentation_requirements, EXPORT_DIR)
            if report["valid"]:
                valid_presentation = True
                break
            coverage_issues.extend(report["issues"])
        if not valid_presentation:
            issues.append("缺少任务所需 PPT 产物：没有生成可下载且逐页有效的 PPTX Artifact")
            issues.extend(coverage_issues[:12])
    return issues


CAPABILITY_SETTING_KEY = "builtin_capability_states"


def _state_overrides(db) -> dict[str, bool]:
    if db is None:
        return {}
    from ..models import AppSetting
    row = db.get(AppSetting, CAPABILITY_SETTING_KEY)
    try:
        value = json.loads(row.value or "{}") if row else {}
    except (json.JSONDecodeError, TypeError):
        value = {}
    return {
        str(name): bool(enabled)
        for name, enabled in value.items()
        if name in TOOLS
    } if isinstance(value, dict) else {}


def globally_enabled_names(db) -> set[str]:
    if not settings.BUILTIN_TOOLS_ENABLED:
        return set()
    overrides = _state_overrides(db)
    return {
        name for name in TOOLS
        if overrides.get(name, True)
    }


def set_global_enabled(db, name: str, enabled: bool) -> None:
    if name not in TOOLS:
        raise KeyError(name)
    from ..models import AppSetting
    states = _state_overrides(db)
    states[name] = bool(enabled)
    row = db.get(AppSetting, CAPABILITY_SETTING_KEY)
    if row is None:
        row = AppSetting(key=CAPABILITY_SETTING_KEY, value="{}")
        db.add(row)
    row.value = json.dumps(states, ensure_ascii=False, sort_keys=True)


def agent_assigned_names(agent) -> set[str]:
    try:
        values = json.loads(getattr(agent, "builtin_tools", "[]") or "[]")
    except (json.JSONDecodeError, TypeError):
        values = []
    return (
        {str(value) for value in values if str(value) in TOOLS}
        if isinstance(values, list) else set()
    )


def effective_tool_names(db, agent) -> set[str]:
    names = globally_enabled_names(db) & agent_assigned_names(agent)
    if bool(getattr(agent, "is_public", False)):
        names &= PUBLIC_AGENT_SAFE_TOOLS
    return names


def capability_catalog(db=None, agent=None) -> list[dict]:
    global_names = globally_enabled_names(db) if db is not None else set(TOOLS)
    assigned = agent_assigned_names(agent) if agent is not None else None
    return [
        {
            "name": tool.name,
            "group": tool.group,
            "mutating": tool.mutating,
            "description": tool.description,
            "enabled": tool.name in global_names,
            "assigned": tool.name in assigned if assigned is not None else None,
        }
        for tool in TOOLS.values()
    ]
