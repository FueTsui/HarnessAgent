"""统一管理运行产物的登记、授权和文件生命周期。"""
from __future__ import annotations

import datetime
import io
import logging
import mimetypes
import re
import uuid
import zipfile
from pathlib import Path

from .config import EXPORT_DIR, settings
from .database import SessionLocal
from .models import Artifact


logger = logging.getLogger(__name__)


_RASTER_FORMATS = {
    "PNG": ("image/png", ".png"),
    "JPEG": ("image/jpeg", ".jpg"),
    "WEBP": ("image/webp", ".webp"),
}
_RASTER_SUFFIXES = {value[1] for value in _RASTER_FORMATS.values()} | {".jpeg"}


def _now() -> datetime.datetime:
    return datetime.datetime.now(datetime.timezone.utc)


def inspect_raster_bytes(data: bytes) -> dict:
    """在落盘前验证图片签名、尺寸与资源上限，不信任文件名或上游 MIME。"""
    from PIL import Image, UnidentifiedImageError

    if not data:
        raise ValueError("图片接口返回空文件")
    max_bytes = max(1, int(settings.IMAGE_OUTPUT_MAX_MB)) * 1024 * 1024
    if len(data) > max_bytes:
        raise ValueError(
            f"图片产物超过 {settings.IMAGE_OUTPUT_MAX_MB} MB 安全上限"
        )
    try:
        with Image.open(io.BytesIO(data)) as image:
            image_format = str(image.format or "").upper()
            width, height = image.size
            image.verify()
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise ValueError("图片产物无法解码或文件已损坏") from exc
    if image_format not in _RASTER_FORMATS:
        raise ValueError(f"不支持的图片产物格式：{image_format or '未知'}")
    pixels = int(width) * int(height)
    if width < 1 or height < 1 or pixels > max(1, int(settings.IMAGE_MAX_PIXELS)):
        raise ValueError(
            f"图片尺寸不安全：{width}x{height}，像素上限 {settings.IMAGE_MAX_PIXELS}"
        )
    media_type, suffix = _RASTER_FORMATS[image_format]
    return {
        "format": image_format,
        "media_type": media_type,
        "suffix": suffix,
        "width": int(width),
        "height": int(height),
        "size_bytes": len(data),
    }


def validate_raster_file(path: Path) -> dict:
    return inspect_raster_bytes(Path(path).read_bytes())


def save_raster_artifact(data: bytes, *, prefix: str = "image") -> dict:
    """验证并原子保存上游生成的位图，返回安全的公开元数据。"""
    metadata = inspect_raster_bytes(data)
    safe_prefix = "".join(
        char for char in str(prefix or "image") if char.isalnum() or char in "-_"
    )[:40] or "image"
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = f"{safe_prefix}_{uuid.uuid4().hex[:10]}{metadata['suffix']}"
    target = (EXPORT_DIR / filename).resolve()
    if target.parent != EXPORT_DIR.resolve():
        raise ValueError("图片产物路径越界")
    temporary = target.with_name(f".{target.name}.{uuid.uuid4().hex}.tmp")
    try:
        temporary.write_bytes(data)
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)
    return {"file": filename, **metadata}


def valid_image_artifact(filename: str) -> bool:
    path = (EXPORT_DIR / Path(str(filename or "")).name).resolve()
    if path.parent != EXPORT_DIR.resolve() or not path.is_file():
        return False
    suffix = path.suffix.lower()
    if suffix in _RASTER_SUFFIXES:
        try:
            validate_raster_file(path)
            return True
        except ValueError:
            return False
    if suffix == ".svg":
        try:
            head = path.read_text(encoding="utf-8", errors="strict")[:4096]
        except (OSError, UnicodeError):
            return False
        return "<svg" in head.lower()
    return False


def valid_presentation_artifact(filename: str) -> bool:
    """验证导出的 PPTX 是结构完整、至少包含一页幻灯片的 OPC 包。"""
    path = (EXPORT_DIR / Path(str(filename or "")).name).resolve()
    if (
        path.parent != EXPORT_DIR.resolve()
        or path.suffix.lower() != ".pptx"
        or not path.is_file()
        or not zipfile.is_zipfile(path)
    ):
        return False
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            slides = [
                name for name in names
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            ]
            return (
                "[Content_Types].xml" in names
                and "ppt/presentation.xml" in names
                and bool(slides)
                and all(bool(archive.read(name)) for name in slides)
            )
    except (OSError, KeyError, zipfile.BadZipFile):
        return False


_PRESENTATION_MARKUP_RE = re.compile(r"\*\*|(?<!\w)__(?!\w)|`")
_PRESENTATION_PROMPT_RE = re.compile(
    r"(?:单击此处|双击此处|请输入|点击添加|Click to add|Lorem ipsum)",
    re.IGNORECASE,
)


def inspect_presentation_artifact(filename: str) -> dict:
    """逐页检查 PPTX 的结构、可见文本和未清理模板痕迹。"""
    if not valid_presentation_artifact(filename):
        return {
            "valid": False,
            "slide_count": 0,
            "issues": ["PPTX 文件结构损坏或不含幻灯片"],
        }
    path = (EXPORT_DIR / Path(str(filename)).name).resolve()
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

        presentation = Presentation(str(path))
    except (ImportError, OSError, ValueError, KeyError) as exc:
        return {
            "valid": False,
            "slide_count": 0,
            "issues": [f"PPTX 无法逐页解析：{type(exc).__name__}"],
        }

    content_placeholder_types = {
        PP_PLACEHOLDER.TITLE,
        PP_PLACEHOLDER.CENTER_TITLE,
        PP_PLACEHOLDER.SUBTITLE,
        PP_PLACEHOLDER.BODY,
        PP_PLACEHOLDER.OBJECT,
    }
    issues: list[str] = []
    checked: list[dict] = []

    def iter_shapes(shapes):
        for shape in shapes:
            yield shape
            nested = getattr(shape, "shapes", None)
            if nested is not None:
                yield from iter_shapes(nested)

    for index, slide in enumerate(presentation.slides, 1):
        local_texts = [
            str(getattr(shape, "text", "") or "").strip()
            for shape in iter_shapes(slide.shapes)
            if hasattr(shape, "text") and str(getattr(shape, "text", "") or "").strip()
        ]
        layout_texts = [
            str(getattr(shape, "text", "") or "").strip()
            for shape in iter_shapes(slide.slide_layout.shapes)
            if hasattr(shape, "text") and str(getattr(shape, "text", "") or "").strip()
        ]
        effective_text = "\n".join([*local_texts, *layout_texts]).strip()
        has_picture = any(
            getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE
            for shape in [
                *iter_shapes(slide.shapes),
                *iter_shapes(slide.slide_layout.shapes),
            ]
        )
        slide_issues: list[str] = []
        if not effective_text and not has_picture:
            slide_issues.append("没有可见文本或图片内容")
        if _PRESENTATION_MARKUP_RE.search(effective_text):
            slide_issues.append("残留 Markdown 标记")
        if _PRESENTATION_PROMPT_RE.search(effective_text):
            slide_issues.append("残留模板提示文字")
        empty_placeholders = []
        for shape in iter_shapes(slide.shapes):
            if not getattr(shape, "is_placeholder", False) or not hasattr(shape, "text"):
                continue
            try:
                placeholder_type = shape.placeholder_format.type
            except (AttributeError, ValueError):
                continue
            if (
                placeholder_type in content_placeholder_types
                and not str(getattr(shape, "text", "") or "").strip()
            ):
                empty_placeholders.append(str(placeholder_type))
        if empty_placeholders:
            slide_issues.append("存在未填内容占位符")
        checked.append({
            "slide": index,
            "text_chars": len(effective_text),
            "issues": list(slide_issues),
        })
        issues.extend(f"第{index}页：{value}" for value in slide_issues)
    return {
        "valid": not issues,
        "slide_count": len(presentation.slides),
        "issues": issues,
        "slides": checked,
    }


def register_many(
    db,
    *,
    owner_id: int,
    run_id: str | None,
    turn_id: str | None,
    filenames: list[str],
) -> list[Artifact]:
    rows = []
    for raw in dict.fromkeys(filenames or []):
        filename = Path(str(raw)).name
        path = (EXPORT_DIR / filename).resolve()
        if path.parent != EXPORT_DIR.resolve() or not path.is_file():
            continue
        media_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"
        if path.suffix.lower() in _RASTER_SUFFIXES:
            # 工具成功但落盘内容损坏时必须让 Job 失败，不能登记一个不可下载的成品。
            media_type = validate_raster_file(path)["media_type"]
        if path.suffix.lower() == ".pptx" and not valid_presentation_artifact(filename):
            raise ValueError("PPTX 产物结构损坏或不含幻灯片")
        existing = db.query(Artifact).filter(
            Artifact.run_id == run_id,
            Artifact.filename == filename,
        ).first()
        if existing:
            rows.append(existing)
            continue
        row = Artifact(
            id=uuid.uuid4().hex,
            run_id=run_id,
            turn_id=turn_id,
            owner_id=owner_id,
            filename=filename,
            media_type=media_type,
            size_bytes=path.stat().st_size,
            expires_at=_now() + datetime.timedelta(
                seconds=settings.ARTIFACT_TTL_SECONDS
            ),
        )
        db.add(row)
        rows.append(row)
    return rows


def register_generated(*, owner_id: int, run_id: str, filename: str) -> None:
    """Persist an intermediate output independently of final-answer verification."""
    from .models import Job, Turn

    with SessionLocal() as db:
        job = db.get(Job, run_id)
        turn = db.get(Turn, run_id)
        if job is None or turn is None or job.owner_id != owner_id or turn.owner_id != owner_id:
            raise ValueError("产物登记缺少有效的运行归属")
        rows = register_many(
            db, owner_id=owner_id, run_id=run_id, turn_id=run_id,
            filenames=[filename],
        )
        if not rows:
            raise ValueError("产物登记失败：生成文件不存在")
        db.commit()


def delete_for_owner(db, owner_id: int) -> list[str]:
    """在当前事务中删除账号的产物元数据，并返回待清理的文件名。

    文件必须等调用方成功提交账号删除事务后再清理；这样若仍有隐藏外键导致
    事务回滚，不会留下“账号尚在但产物文件已丢失”的半完成状态。
    """
    rows = db.query(Artifact).filter(Artifact.owner_id == int(owner_id)).all()
    filenames = list(dict.fromkeys(row.filename for row in rows))
    if rows:
        db.query(Artifact).filter(Artifact.owner_id == int(owner_id)).delete(
            synchronize_session=False
        )
    return filenames


def purge_unreferenced_files(db, filenames: list[str]) -> int:
    """删除已无 Artifact 元数据引用的导出文件，保留同名共享引用。"""
    root = EXPORT_DIR.resolve()
    deleted = 0
    for raw in dict.fromkeys(filenames or []):
        filename = Path(str(raw)).name
        if not filename:
            continue
        if db.query(Artifact.id).filter(Artifact.filename == filename).first() is not None:
            continue
        path = (EXPORT_DIR / filename).resolve()
        if path.parent != root:
            continue
        try:
            existed = path.is_file()
            path.unlink(missing_ok=True)
            deleted += int(existed)
        except OSError:
            # 数据库归属已经安全移除；文件系统的瞬时占用不应把成功的账号删除
            # 伪装成失败。保留告警，便于运维后续清理无元数据引用的文件。
            logger.warning("删除账号运行产物文件失败：%s", path, exc_info=True)
    return deleted


def cleanup(db) -> int:
    now = _now()
    rows = db.query(Artifact).filter(Artifact.expires_at <= now).all()
    deleted = 0
    for row in rows:
        path = (EXPORT_DIR / Path(row.filename).name).resolve()
        if path.parent == EXPORT_DIR.resolve():
            path.unlink(missing_ok=True)
        db.delete(row)
        deleted += 1
    return deleted


def cleanup_expired() -> int:
    db = SessionLocal()
    try:
        deleted = cleanup(db)
        if deleted:
            db.commit()
        return deleted
    finally:
        db.close()
