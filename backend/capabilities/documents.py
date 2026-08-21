"""附件文档文本提取能力。

支持 TXT/MD/CSV 原生解析；PDF/DOCX/XLSX 在安装可选依赖（pypdf/python-docx/openpyxl）时解析。
"""
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

TEXT_SUFFIXES = {".txt", ".md", ".csv"}


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"[工作表 {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_documents(paths: list[Path]) -> str:
    """逐个文件提取文本，单文件失败不影响整体（返回提示而非中断）。"""
    sections = []
    for path in paths:
        suffix = path.suffix.lower()
        try:
            if suffix in TEXT_SUFFIXES:
                text = path.read_text(encoding="utf-8", errors="ignore")
            elif suffix == ".pdf":
                text = _extract_pdf(path)
            elif suffix == ".docx":
                text = _extract_docx(path)
            elif suffix == ".xlsx":
                text = _extract_xlsx(path)
            else:
                text = f"[暂不支持的文件类型 {suffix}]"
        except ImportError as exc:
            text = f"[缺少解析依赖，无法读取 {path.name}: {exc.name}]"
        except Exception as exc:  # noqa: BLE001
            logger.warning("文档提取失败 %s: %s", path.name, exc)
            text = f"[文件解析失败 {path.name}]"
        sections.append(f"### 文件：{path.name}\n{text.strip()}")
    return "\n\n".join(sections)
