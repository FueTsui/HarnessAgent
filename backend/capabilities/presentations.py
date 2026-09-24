"""Bounded native-text PowerPoint export, without user code or filesystem access."""
from pathlib import Path
import re
import uuid


def create_presentation(args: dict, export_dir: Path, *, requirements=None) -> dict:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    slides = args.get("slides")
    if not isinstance(slides, list) or not 1 <= len(slides) <= 80:
        raise ValueError("slides 必须包含 1–80 页")
    if requirements is not None:
        issues = requirements.validate_slides(slides)
        if issues:
            raise ValueError("；".join(issues[:12]))
    # Reject overflow rather than truncate source content or shrink individual fonts.
    for index, spec in enumerate(slides, 1):
        if not isinstance(spec, dict):
            raise ValueError(f"第 {index} 页必须是对象")
        title, body = spec.get("title"), spec.get("body")
        if not isinstance(title, str) or not title.strip() or len(title) > 48:
            raise ValueError(f"第 {index} 页标题必须为 1–48 字符")
        if not isinstance(body, str) or len(body) > 480:
            raise ValueError(f"第 {index} 页正文最多 480 字符，请拆页")
        units = lambda line: sum(2 if ord(c) > 255 else 1 for c in line)
        lines = sum(max(1, (units(line) + 69) // 70) for line in body.splitlines())
        if lines > 10:
            raise ValueError(f"第 {index} 页正文超过 10 行，请拆页")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = str(args.get("title") or "演示文稿")[:300]
    for index, spec in enumerate(slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if requirements is not None:
            source = requirements.pages[index - 1]
            slide.notes_slide.notes_text_frame.text = (
                f"来源页 {source.index}/{requirements.page_count}；"
                f"{source.source_name} 原第 {source.source_page}/{source.source_count} 页"
            )
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F6F8FC")
        for text, top, height, size, bold, color in [
            (spec["title"], .5, 1.2, 32, True, "183153"),
            (spec["body"], 1.9, 4.8, 22, False, "25364A"),
            (str(index), 6.95, .3, 11, False, "64748B"),
        ]:
            box = slide.shapes.add_textbox(Inches(.7), Inches(top), Inches(11.9), Inches(height))
            frame = box.text_frame
            frame.word_wrap = True
            frame.margin_left = frame.margin_right = 0
            for n, line in enumerate(text.split("\n")):
                p = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
                p.text = line
                p.font.name, p.font.size, p.font.bold = "Microsoft YaHei", Pt(size), bold
                p.font.color.rgb = RGBColor.from_string(color)
                p.space_after = Pt(6)
    stem = re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', str(args.get("output_name") or "演示文稿"))
    stem = Path(stem).stem.strip(" .")[:100] or "演示文稿"
    name = f"{stem}_{uuid.uuid4().hex[:12]}.pptx"
    export_dir.mkdir(parents=True, exist_ok=True)
    target = export_dir / name
    prs.save(str(target))
    reopened = Presentation(str(target))
    if len(reopened.slides) != len(slides):
        target.unlink(missing_ok=True)
        raise ValueError("PPTX 页数校验失败")
    if requirements is not None:
        from ..runtime.presentation_requirements import inspect_presentation_requirements
        report = inspect_presentation_requirements(name, requirements, export_dir)
        if not report["valid"]:
            target.unlink(missing_ok=True)
            raise ValueError("；".join(report["issues"][:12]))
    return {"ok": True, "file": name, "download_url": f"/api/v1/exports/{name}",
            "slide_count": len(slides), "editable_text": True,
            "layout": "native_text", "limitations": "文字重建，不含源文件图片或原版式复刻"}
